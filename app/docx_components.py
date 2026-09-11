from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

from app.models import Run


RATING_FONT_COLORS = {
    "critical": RGBColor(192, 0, 0),
    "high": RGBColor(237, 125, 49),
    "medium": RGBColor(255, 192, 0),
    "low": RGBColor(112, 173, 71),
    "informational": RGBColor(81, 81, 81),
}
RATING_COLOR_TOKENS = {
    "rating",
    "vuln_severity",
    "likelihood-rating",
    "impact-rating",
    "severity-rating",
}


class ComponentCompositionError(ValueError):
    """Raised when a DOCX component cannot be safely merged."""


@dataclass(frozen=True)
class DocxComponent:
    path: Path
    values: dict[str, str] = field(default_factory=dict)


def clone_component_elements(
    document: DocumentType,
    component_path: Path,
    *,
    numbering_ids: dict[tuple[Path, int], int] | None = None,
) -> list:
    """Clone a component body into a target document's style and numbering context."""
    if not component_path.is_file():
        raise ComponentCompositionError(f"Component not found: {component_path}")
    component_document = Document(component_path)
    elements = [
        deepcopy(element)
        for element in component_document.element.body.iterchildren()
        if element.tag != qn("w:sectPr")
    ]
    if not elements:
        raise ComponentCompositionError(f"Component has no body content: {component_path}")
    _validate_component(elements, {style.style_id for style in document.styles}, component_path)
    _remap_numbering(
        document,
        component_document,
        elements,
        component_path,
        numbering_ids if numbering_ids is not None else {},
    )
    return elements


def replace_component_token(elements: list, token: str, value: str) -> None:
    """Replace a component token while retaining its source run formatting."""
    _replace_token(elements, token, value)


def replace_component_token_runs(elements: list, token: str, runs: list[Run]) -> None:
    """Replace a component token with rich runs layered over its source formatting."""
    pattern = re.compile(r"\{\{\s*" + re.escape(token) + r"\s*\}\}", re.IGNORECASE)
    for element in elements:
        for paragraph in element.iter(qn("w:p")):
            _replace_pattern_with_runs(paragraph, pattern, runs)


def compose_docx_components(
    main_template: Path,
    components: list[DocxComponent],
    *,
    anchor: str = "findings",
) -> bytes:
    """Insert component document bodies at a main-template placeholder."""
    if not main_template.is_file():
        raise ComponentCompositionError(f"Main template not found: {main_template}")
    if not components:
        raise ComponentCompositionError("At least one component is required")

    document = Document(main_template)
    anchor_paragraph = _find_anchor_paragraph(document, anchor)
    anchor_element = anchor_paragraph._p
    main_style_ids = {style.style_id for style in document.styles}

    for component in components:
        component_document = Document(component.path)
        body_elements = [
            deepcopy(element)
            for element in component_document.element.body.iterchildren()
            if element.tag != qn("w:sectPr")
        ]
        if not body_elements:
            raise ComponentCompositionError(f"Component has no body content: {component.path}")
        _validate_component(body_elements, main_style_ids, component.path)
        for token, value in component.values.items():
            _replace_token(body_elements, token, value)
        unresolved = _tokens(body_elements)
        if unresolved:
            raise ComponentCompositionError(
                f"Unresolved placeholders in {component.path.name}: {', '.join(unresolved)}"
            )
        first_paragraph = next((element for element in body_elements if element.tag == qn("w:p")), None)
        if first_paragraph is None:
            raise ComponentCompositionError(f"Component must begin with paragraph content: {component.path}")
        _set_page_break_before(first_paragraph)
        for element in body_elements:
            anchor_element.addprevious(element)

    anchor_element.getparent().remove(anchor_element)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def compose_docx_template(
    template: Path,
    *,
    values: dict[str, str],
    components: dict[str, list[DocxComponent]],
) -> bytes:
    """Populate a DOCX template and insert formatted components at its anchors."""
    if not template.is_file():
        raise ComponentCompositionError(f"Template not found: {template}")

    document = Document(template)
    body = document.element.body
    main_style_ids = {style.style_id for style in document.styles}
    for token, value in values.items():
        _replace_token([body], token, value)

    for anchor, anchor_components in components.items():
        anchor_paragraph = _find_anchor_paragraph(document, anchor)
        anchor_element = anchor_paragraph._p
        numbering_ids: dict[tuple[Path, int], int] = {}
        for component in anchor_components:
            component_document = Document(component.path)
            body_elements = [
                deepcopy(element)
                for element in component_document.element.body.iterchildren()
                if element.tag != qn("w:sectPr")
            ]
            if not body_elements:
                raise ComponentCompositionError(f"Component has no body content: {component.path}")
            _validate_component(body_elements, main_style_ids, component.path)
            _remap_numbering(document, component_document, body_elements, component.path, numbering_ids)
            for token, value in component.values.items():
                _replace_token(body_elements, token, value)
            unresolved = _tokens(body_elements)
            if unresolved:
                raise ComponentCompositionError(
                    f"Unresolved placeholders in {component.path.name}: {', '.join(unresolved)}"
                )
            for element in body_elements:
                anchor_element.addprevious(element)
        anchor_element.getparent().remove(anchor_element)

    unresolved = _tokens([body])
    if unresolved:
        raise ComponentCompositionError(f"Unresolved placeholders in {template.name}: {', '.join(unresolved)}")

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _find_anchor_paragraph(document, anchor: str):
    pattern = re.compile(r"\{\{\s*" + re.escape(anchor) + r"\s*\}\}", re.IGNORECASE)
    matches = [
        paragraph
        for paragraph in document.paragraphs
        if pattern.search(paragraph.text)
    ]
    if len(matches) != 1:
        raise ComponentCompositionError(
            f"Expected exactly one {{{{{anchor}}}}} paragraph; found {len(matches)}"
        )
    return matches[0]


def _validate_component(elements: list, main_style_ids: set[str], path: Path) -> None:
    referenced_styles = {
        style.get(qn("w:val"))
        for element in elements
        for style in element.iter(qn("w:pStyle"))
        if style.get(qn("w:val"))
    }
    missing_styles = sorted(referenced_styles - main_style_ids)
    if missing_styles:
        raise ComponentCompositionError(
            f"Component {path.name} uses styles missing from the main template: {', '.join(missing_styles)}"
        )
    relationship_attributes = {
        qn("r:id"),
        qn("r:embed"),
        qn("r:link"),
    }
    if any(
        attribute in node.attrib
        for element in elements
        for node in element.iter()
        for attribute in relationship_attributes
    ):
        raise ComponentCompositionError(
            f"Component {path.name} contains body relationships; relationship remapping is required"
        )


def _remap_numbering(
    document,
    component_document,
    elements: list,
    component_path: Path,
    numbering_ids: dict[tuple[Path, int], int],
) -> None:
    source_numbering = component_document.part.numbering_part.element
    target_numbering = document.part.numbering_part.element
    for numbering_id_element in (
        numbering_id
        for element in elements
        for numbering_id in element.iter(qn("w:numId"))
    ):
        source_numbering_id = int(numbering_id_element.get(qn("w:val")))
        key = (component_path.resolve(), source_numbering_id)
        target_numbering_id = numbering_ids.get(key)
        if target_numbering_id is None:
            source_number = next(
                (
                    number
                    for number in source_numbering.iterchildren(qn("w:num"))
                    if number.get(qn("w:numId")) == str(source_numbering_id)
                ),
                None,
            )
            if source_number is None:
                raise ComponentCompositionError(
                    f"Component {component_path.name} references missing numbering ID {source_numbering_id}"
                )
            source_abstract_id = source_number.find(qn("w:abstractNumId")).get(qn("w:val"))
            source_abstract = next(
                (
                    abstract
                    for abstract in source_numbering.iterchildren(qn("w:abstractNum"))
                    if abstract.get(qn("w:abstractNumId")) == source_abstract_id
                ),
                None,
            )
            if source_abstract is None:
                raise ComponentCompositionError(
                    f"Component {component_path.name} references missing abstract numbering ID {source_abstract_id}"
                )

            target_abstract_id = _next_numbering_id(target_numbering, "w:abstractNum", "w:abstractNumId")
            target_numbering_id = _next_numbering_id(target_numbering, "w:num", "w:numId")
            abstract = deepcopy(source_abstract)
            abstract.set(qn("w:abstractNumId"), str(target_abstract_id))
            _assign_unique_numbering_identity(target_numbering, abstract)
            number = deepcopy(source_number)
            number.set(qn("w:numId"), str(target_numbering_id))
            number.find(qn("w:abstractNumId")).set(qn("w:val"), str(target_abstract_id))
            _restart_numbering_levels(number, source_abstract, elements, source_numbering_id)
            first_number = next(target_numbering.iterchildren(qn("w:num")), None)
            if first_number is None:
                target_numbering.append(abstract)
            else:
                first_number.addprevious(abstract)
            target_numbering.append(number)
            numbering_ids[key] = target_numbering_id
        numbering_id_element.set(qn("w:val"), str(target_numbering_id))


def _next_numbering_id(numbering, element_tag: str, id_attribute: str) -> int:
    identifiers = [
        int(element.get(qn(id_attribute)))
        for element in numbering.iterchildren(qn(element_tag))
    ]
    return max(identifiers, default=0) + 1


def _assign_unique_numbering_identity(numbering, abstract) -> None:
    existing = {
        identity.get(qn("w:val")).upper()
        for item in numbering.iterchildren(qn("w:abstractNum"))
        if (identity := item.find(qn("w:nsid"))) is not None
        and identity.get(qn("w:val"))
    }
    candidate = (max((int(value, 16) for value in existing), default=0) + 1) & 0xFFFFFFFF
    while f"{candidate:08X}" in existing:
        candidate = (candidate + 1) & 0xFFFFFFFF
    identity = abstract.find(qn("w:nsid"))
    if identity is None:
        identity = OxmlElement("w:nsid")
        abstract.insert(0, identity)
    identity.set(qn("w:val"), f"{candidate:08X}")


def _restart_numbering_levels(number, abstract, elements: list, numbering_id: int) -> None:
    levels = {
        int(level.get(qn("w:val"))) if level is not None else 0
        for element in elements
        for properties in element.iter(qn("w:numPr"))
        if (identifier := properties.find(qn("w:numId"))) is not None
        and identifier.get(qn("w:val")) == str(numbering_id)
        for level in [properties.find(qn("w:ilvl"))]
    }
    for level_index in levels:
        level = next(
            (
                item
                for item in abstract.iterchildren(qn("w:lvl"))
                if item.get(qn("w:ilvl")) == str(level_index)
            ),
            None,
        )
        if level is None:
            continue
        number_format = level.find(qn("w:numFmt"))
        if number_format is not None and number_format.get(qn("w:val")) in {"bullet", "none"}:
            continue
        start = level.find(qn("w:start"))
        start_value = start.get(qn("w:val")) if start is not None else "1"
        override = next(
            (
                item
                for item in number.iterchildren(qn("w:lvlOverride"))
                if item.get(qn("w:ilvl")) == str(level_index)
            ),
            None,
        )
        if override is None:
            override = OxmlElement("w:lvlOverride")
            override.set(qn("w:ilvl"), str(level_index))
            number.append(override)
        start_override = override.find(qn("w:startOverride"))
        if start_override is None:
            start_override = OxmlElement("w:startOverride")
            level_override = override.find(qn("w:lvl"))
            if level_override is None:
                override.append(start_override)
            else:
                level_override.addprevious(start_override)
        start_override.set(qn("w:val"), start_value)


def _replace_token(elements: list, token: str, value: str) -> None:
    pattern = re.compile(r"\{\{\s*" + re.escape(token) + r"\s*\}\}", re.IGNORECASE)
    font_color = (
        RATING_FONT_COLORS.get(value.strip().casefold())
        if token.strip().casefold() in RATING_COLOR_TOKENS
        else None
    )
    for element in elements:
        for paragraph in element.iter(qn("w:p")):
            replace_pattern_across_text_nodes(paragraph, pattern, value, font_color)


def replace_pattern_across_text_nodes(
    paragraph,
    pattern: re.Pattern,
    replacement: str,
    font_color: RGBColor | None = None,
    font_size_pt: int | None = None,
) -> None:
    """Replace split-run paragraph text while preserving its source formatting."""
    nodes = list(paragraph.iter(qn("w:t")))
    while nodes:
        text = "".join(node.text or "" for node in nodes)
        match = pattern.search(text)
        if not match:
            return
        positions = []
        offset = 0
        for node in nodes:
            length = len(node.text or "")
            positions.append((offset, offset + length))
            offset += length
        start_index = next(index for index, (_, end) in enumerate(positions) if match.start() < end)
        end_index = next(index for index, (_, end) in enumerate(positions) if match.end() <= end)
        start_offset = match.start() - positions[start_index][0]
        end_offset = match.end() - positions[end_index][0]
        prefix = (nodes[start_index].text or "")[:start_offset]
        suffix = (nodes[end_index].text or "")[end_offset:]
        nodes[start_index].text = prefix + replacement + (suffix if start_index == end_index else "")
        if font_color is not None:
            _set_text_node_color(nodes[start_index], font_color)
        if font_size_pt is not None:
            _set_text_node_size(nodes[start_index], font_size_pt)
        if start_index != end_index:
            for index in range(start_index + 1, end_index):
                nodes[index].text = ""
            nodes[end_index].text = suffix
        for node in (nodes[start_index], nodes[end_index]):
            _preserve_text_spaces(node)


def _set_text_node_color(text_node, font_color: RGBColor) -> None:
    run = text_node.getparent()
    while run is not None and run.tag != qn("w:r"):
        run = run.getparent()
    if run is None:
        return
    properties = run.find(qn("w:rPr"))
    if properties is None:
        properties = OxmlElement("w:rPr")
        run.insert(0, properties)
    color = properties.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        properties.append(color)
    color.set(qn("w:val"), str(font_color))
    for attribute in ("w:themeColor", "w:themeTint", "w:themeShade"):
        color.attrib.pop(qn(attribute), None)


def _set_text_node_size(text_node, font_size_pt: int) -> None:
    run = text_node.getparent()
    while run is not None and run.tag != qn("w:r"):
        run = run.getparent()
    if run is None:
        return
    properties = run.find(qn("w:rPr"))
    if properties is None:
        properties = OxmlElement("w:rPr")
        run.insert(0, properties)
    for tag in ("w:sz", "w:szCs"):
        size = properties.find(qn(tag))
        if size is None:
            size = OxmlElement(tag)
            properties.append(size)
        size.set(qn("w:val"), str(font_size_pt * 2))


def _replace_pattern_with_runs(paragraph, pattern: re.Pattern, replacement_runs: list[Run]) -> None:
    nodes = list(paragraph.iter(qn("w:t")))
    while nodes:
        text = "".join(node.text or "" for node in nodes)
        match = pattern.search(text)
        if not match:
            return
        positions = []
        offset = 0
        for node in nodes:
            length = len(node.text or "")
            positions.append((offset, offset + length))
            offset += length
        start_index = next(index for index, (_, end) in enumerate(positions) if match.start() < end)
        end_index = next(index for index, (_, end) in enumerate(positions) if match.end() <= end)
        start_node = nodes[start_index]
        end_node = nodes[end_index]
        start_run = _parent_run(start_node)
        end_run = _parent_run(end_node)
        if start_run is None or end_run is None:
            return

        prefix = (start_node.text or "")[:match.start() - positions[start_index][0]]
        suffix = (end_node.text or "")[match.end() - positions[end_index][0]:]
        base_run = deepcopy(start_run)
        for index in range(start_index, end_index + 1):
            nodes[index].text = ""

        source_runs = [source for source in replacement_runs if source.text]
        insertion_point = start_run
        if prefix:
            start_node.text = prefix
            for source in source_runs:
                generated_run = deepcopy(base_run)
                _set_run_text(generated_run, source.text)
                _apply_run_formatting(generated_run, source)
                insertion_point.addnext(generated_run)
                insertion_point = generated_run
        elif source_runs:
            first, *remaining = source_runs
            _set_run_text(start_run, first.text)
            _apply_run_formatting(start_run, first)
            for source in remaining:
                generated_run = deepcopy(base_run)
                _set_run_text(generated_run, source.text)
                _apply_run_formatting(generated_run, source)
                insertion_point.addnext(generated_run)
                insertion_point = generated_run

        if suffix:
            if end_run is not start_run:
                end_node.text = suffix
                _preserve_text_spaces(end_node)
            elif prefix and not source_runs:
                start_node.text = prefix + suffix
                _preserve_text_spaces(start_node)
            else:
                suffix_run = deepcopy(base_run)
                _set_run_text(suffix_run, suffix)
                insertion_point.addnext(suffix_run)
        nodes = list(paragraph.iter(qn("w:t")))


def _parent_run(node):
    parent = node.getparent()
    while parent is not None and parent.tag != qn("w:r"):
        parent = parent.getparent()
    return parent


def _set_run_text(run, text: str) -> None:
    for child in list(run):
        if child.tag != qn("w:rPr"):
            run.remove(child)
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    for index, line in enumerate(lines):
        if index:
            run.append(OxmlElement("w:br"))
        text_node = OxmlElement("w:t")
        text_node.text = line
        _preserve_text_spaces(text_node)
        run.append(text_node)


def _preserve_text_spaces(text_node) -> None:
    text = text_node.text or ""
    if text and (text[0].isspace() or text[-1].isspace()):
        text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _apply_run_formatting(run, source: Run) -> None:
    properties = run.find(qn("w:rPr"))
    if properties is None:
        properties = OxmlElement("w:rPr")
        run.insert(0, properties)
    for enabled, tag in ((source.bold, "w:b"), (source.italic, "w:i")):
        if enabled:
            setting = properties.find(qn(tag))
            if setting is None:
                setting = OxmlElement(tag)
                properties.append(setting)
            setting.attrib.pop(qn("w:val"), None)
    if source.underline:
        underline = properties.find(qn("w:u"))
        if underline is None:
            underline = OxmlElement("w:u")
            properties.append(underline)
        underline.set(qn("w:val"), "single")


def _tokens(elements: list) -> list[str]:
    text = "\n".join(
        "".join(node.text or "" for node in element.iter(qn("w:t")))
        for element in elements
    )
    return sorted(set(re.findall(r"\{\{\s*(.*?)\s*\}\}", text)))


def _set_page_break_before(paragraph) -> None:
    properties = paragraph.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        paragraph.insert(0, properties)
    if properties.find(qn("w:pageBreakBefore")) is None:
        properties.append(OxmlElement("w:pageBreakBefore"))