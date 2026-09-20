from __future__ import annotations

import re
from copy import deepcopy
from functools import lru_cache
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.parts.hdrftr import FooterPart, HeaderPart
from docx.shared import Mm, Pt, RGBColor
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps

from app.models import (
    CodeFragment,
    Content,
    ImageFragment,
    InstanceTitleFragment,
    ListFragment,
    NoteFragment,
    ParagraphFragment,
    Report,
    Run,
    TableFragment,
    Vulnerability,
)
from .docx_captions import add_native_image_captions, center_paragraph
from .docx_import import FINDING_HEADING_STYLE, INSTANCE_PREFIX
from .docx_components import (
    RATING_FONT_COLORS,
    clone_component_elements,
    replace_component_token,
    replace_component_token_runs,
    replace_pattern_across_text_nodes,
)
from .models import CHANNEL_LABELS, CHANNELS, COMPONENT_CHANNELS
from .report_service import REPORT_TYPE_LABELS, affected_environments, content_types_for_status, finding_is_complete, fragment_applies, is_default_status_conclusion, location_lines, setup_issues

SEVERITY_ORDER = ["critical", "high", "medium", "low", "informational"]
# Targets number from zero within each channel, so channel rank has to come first when ordering them.
CHANNEL_ORDER = list(CHANNELS)
STATUS_LABELS = {
    "open_new": "Open (New)",
    "open_previously_discovered": "Open (Previously Discovered)",
    "resolved": "Resolved",
}
# Belongs to the document, never to the draft, so the stored value stays bare digits.
SEVERITY_TICKET_PREFIX = "GRIMPEN-"
PLAIN_METADATA_TOKENS = {"app-name", "app-owner", "tester-name", "report-name"}
UNRESOLVED_MARKERS = (
    "-vuln",
    "-affected-location",
    "-fragments-here",
    "-step-1-here",
    "-step-2-here",
    "-images-and-caption-here",
    "brief-explanation-here",
)
PLACEHOLDER_TEXT = re.compile(
    r"\(\s*insert[^)]*\)|insert\s+(technology|version|eol\s+date|cves|latest)\s+\w*\s*here",
    re.IGNORECASE,
)
CELL_PLACEHOLDER = re.compile(r"\{\{\s*[^{}]+?\s*\}\}")
SEVERITY_COMPONENT_FILES = {
    "critical": "critical_severity.docx",
    "high": "high_severity.docx",
    "medium": "medium_severity.docx",
    "low": "low_severity.docx",
    "informational": "informatinal_severity.docx",
}
FRAGMENT_COMPONENT_FILES = {
    "paragraph": ("paragraph_fragment.docx", "paragraph-fragment"),
    "numbered_list": ("numbered_fragment.docx", "numbered-list-fragment"),
    "bulleted_list": ("bulleted_fragment.docx", "bullet-list-fragment"),
    "note": ("note_fragment.docx", "note-fragment"),
    "code_block": ("code_fragment.docx", "code-fragment"),
    "instance_title": ("title_fragment.docx", "instance-fragment"),
    "caption": ("caption_fragment.docx", "caption-fragment"),
}
IMAGE_BORDER_RASTER_DPI = 192
# Long endpoints are broken so they stay inside their column.
SCOPE_WRAP_CHARACTERS = 84
LOCATION_WRAP_CHARACTERS = 74


class ReportGenerationError(ValueError):
    """Raised when a report cannot be rendered safely into the template."""


def generation_issues(report: Report) -> list[str]:
    """Return report-completeness issues that should block final DOCX output."""
    issues = setup_issues(report)
    if not report.vulnerabilities:
        issues.append("at least one finding")
    for finding in report.vulnerabilities:
        label = finding.title or "Untitled finding"
        if not finding_is_complete(finding, report):
            issues.append(f"{label}: finding details or affected locations are incomplete")
        proof = next((content for content in finding.contents if content.type == "proof_of_concept"), None)
        # Last year's screenshot is not this year's proof, so only the proof of concept counts.
        images = [fragment for fragment in proof.fragments if isinstance(fragment, ImageFragment)] if proof else []
        environments = affected_environments(finding, report)
        for environment in environments:
            if not any(image.environment == environment and image.evidence_id for image in images):
                issues.append(f"{label}: {'Production' if environment == 'production' else 'Non-Production'} evidence image required")
        # Only Asia prints these, and a field the tester never sees must never block their report.
        if report.engagement.segment == "Asia":
            for field_label, value in (("CVSS Score", finding.cvss_score), ("CVSS Vector", finding.cvss_vector)):
                if not value.strip():
                    issues.append(f"{label}: {field_label} is required")
        printed = set(content_types_for_status(finding.status))
        for content in finding.contents:
            # A section this status does not print is carried for safekeeping, not for completing.
            if content.type not in printed:
                continue
            # Twin of the editor's requiresFragment rule: these carry the finding, so none is ever left empty.
            if content.type in {"description", "recommended_remediation", "in_conclusion"} and not content.fragments:
                issues.append(f"{label}: {content.type} needs at least one fragment")
            # The app writes this sentence itself, so leaving it is the same as leaving the section blank.
            if content.type == "in_conclusion" and any(
                isinstance(fragment, ParagraphFragment) and is_default_status_conclusion(fragment)
                for fragment in content.fragments
            ):
                issues.append(f"{label}: in_conclusion still holds the default sentence")
            for fragment in content.fragments:
                if isinstance(fragment, (ParagraphFragment, NoteFragment)) and not _runs_have_text(fragment.runs):
                    issues.append(f"{label}: {content.type} text is required")
                elif isinstance(fragment, ListFragment) and any(not _runs_have_text(item.runs) for item in fragment.items):
                    issues.append(f"{label}: {content.type} list item text is required")
                elif isinstance(fragment, ImageFragment):
                    if not fragment_applies(fragment, finding, report, content.type):
                        continue
                    missing = []
                    if not fragment.environment:
                        missing.append("environment")
                    if not fragment.evidence_id:
                        missing.append("image")
                    if not fragment.caption.strip():
                        missing.append("caption")
                    if missing:
                        issues.append(f"{label}: {'/'.join(missing)} required for image fragment")
                elif isinstance(fragment, CodeFragment) and not fragment.text.strip():
                    issues.append(f"{label}: {fragment.type} text is required")
                elif isinstance(fragment, TableFragment):
                    cells = [*fragment.header, *(cell for row in fragment.rows for cell in row)]
                    if not cells or any(not _runs_have_text(cell.runs) for cell in cells):
                        issues.append(f"{label}: every table cell is required")
                if PLACEHOLDER_TEXT.search(_fragment_text(fragment)):
                    issues.append(f"{label}: replace placeholder text in {content.type}")
    return list(dict.fromkeys(issues))


def main_template_path(report: Report, resources: Path) -> Path:
    """Pick the main template on both axes: Asia segment, and whether a component app type is covered.

    Sole owner of the choice, because `template_path.parent` is also the component fragment root."""
    component = any(channel in report.engagement.tested_channels for channel in COMPONENT_CHANNELS)
    asia = report.engagement.segment == "Asia"
    stem = "MAIN_THICK_MOBILE" if component else "MAIN"
    return resources / f"{stem}{'_ASIA' if asia else ''}.docx"


def render_report_docx(
    report: Report,
    template_path: Path,
    report_folder: Path,
    *,
    allow_incomplete: bool = False,
    validation_issues: list[str] | None = None,
) -> bytes:
    """Render a validated report into the supplied canonical DOCX template."""
    if not template_path.is_file():
        raise ReportGenerationError(f"Template not found: {template_path}")
    issues = generation_issues(report) if validation_issues is None else validation_issues
    if issues and not allow_incomplete:
        raise ReportGenerationError("; ".join(issues))

    document = Document(template_path)
    _normalize_page_numbering(document)
    _populate_scope_tables(document, report)
    _populate_summary_table(document, report)
    _replace_metadata(document, report)
    if not _has_exact_body_token(document, "findings"):
        raise ReportGenerationError(f"Template has no {{{{findings}}}} anchor paragraph: {template_path}")
    rendered = _populate_component_findings(
        document,
        report,
        report_folder,
        template_path.parent,
        allow_incomplete,
    )
    # After the findings so the headings exist to bookmark, before the caption pass so its
    # mark-every-field-dirty sweep reaches the references this writes.
    _populate_cvss_table(document, rendered)
    add_native_image_captions(document)
    unresolved = _unresolved_placeholders(document)
    if unresolved:
        raise ReportGenerationError(f"Unresolved template placeholders: {', '.join(unresolved)}")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _runs_have_text(runs: list[Run]) -> bool:
    return any(run.text.strip() for run in runs)


def _fragment_text(fragment) -> str:
    if isinstance(fragment, (ParagraphFragment, NoteFragment)):
        return "".join(run.text for run in fragment.runs)
    if isinstance(fragment, ListFragment):
        return " ".join(run.text for item in fragment.items for run in item.runs)
    if isinstance(fragment, TableFragment):
        cells = [*fragment.header, *(cell for row in fragment.rows for cell in row)]
        return " ".join(run.text for cell in cells for run in cell.runs)
    if isinstance(fragment, ImageFragment):
        return fragment.caption
    if isinstance(fragment, (CodeFragment, InstanceTitleFragment)):
        return fragment.text
    return ""


def _display_date(value) -> str:
    return f"{value:%B} {value.day}, {value.year}" if value else "N/A"


def _display_value(value: str | None) -> str:
    return value.strip() if value and value.strip() else "N/A"


def _target_values(report: Report, environment: str, channel: str) -> list[str]:
    return [
        target.value
        for target in sorted(report.scope_targets, key=lambda item: item.order)
        if target.environment == environment and target.channel == channel
    ]


def _component_rows(report: Report) -> list[tuple[str, str]]:
    """Every component in one list, production first. The table has no environment column, so the
    ordering is the only thing carrying that distinction to the reader."""
    return [
        (target.value, target.description)
        for environment in ("production", "non_production")
        for target in sorted(report.scope_targets, key=lambda item: item.order)
        if target.environment == environment and target.channel in COMPONENT_CHANNELS and target.value.strip()
    ]


def _component_channel_label(report: Report) -> str:
    """Names the one component app type in the caption. The pair is refused by setup_issues long
    before generation, so the CHANNELS-order tie-break only ever settles a scripted render."""
    covered = [channel for channel in COMPONENT_CHANNELS if channel in report.engagement.tested_channels]
    return CHANNEL_LABELS[covered[0]] if covered else "N/A"


def _metadata(report: Report) -> dict[str, str]:
    engagement = report.engagement
    # Windows are kept for an unchecked environment so re-checking restores them, but an
    # untested environment must not print dates in the report.
    tested = lambda environment: engagement.test_windows.get(environment) if environment in engagement.tested_environments else None
    production = tested("production")
    non_production = tested("non_production")
    accounts = engagement.test_accounts
    values = {
        "segment": engagement.segment or "N/A",
        "network": engagement.network,
        "app-name": _display_value(engagement.app_name),
        "report-name": _display_value(engagement.app_name),
        "test-type": REPORT_TYPE_LABELS.get(engagement.report_type or "", "N/A"),
        "app-owner": _display_value(engagement.app_owner),
        "report-date": _display_date(engagement.report_date),
        "tester-name": _display_value(engagement.tester),
        "non-prod-start": _display_date(non_production.start_date if non_production else None),
        "non-prod-end": _display_date(non_production.end_date if non_production else None),
        "non-prod-time": _display_value(non_production.test_time if non_production else None),
        "prod-start": _display_date(production.start_date if production else None),
        "prod-end": _display_date(production.end_date if production else None),
        "prod-time": _display_value(production.test_time if production else None),
        "prod-web": "; ".join(_target_values(report, "production", "web")) or "N/A",
        "non-prod-web": "; ".join(_target_values(report, "non_production", "web")) or "N/A",
        "prod-api": "; ".join(_target_values(report, "production", "api")) or "N/A",
        "non-prod-api": "; ".join(_target_values(report, "non_production", "api")) or "N/A",
        "limitation-set": _display_value(engagement.limitations),
    }
    # The two component templates spell the same token in opposite orders. Both keys are always
    # supplied: _replace_metadata no-ops on a token the template does not carry.
    values["mobile-thick"] = values["thick-mobile"] = _component_channel_label(report)
    for index in range(2):
        account = accounts[index] if index < len(accounts) else None
        values[f"role{index + 1}"] = _display_value(account.user_role if account else None)
        values[f"username{index + 1}"] = _display_value(account.username if account else None)
    return values


def _document_roots(document: DocumentType):
    yield document.element
    for part in document.part.package.parts:
        if isinstance(part, (HeaderPart, FooterPart)):
            yield part._element


def _replace_metadata(document: DocumentType, report: Report) -> None:
    for token, value in sorted(_metadata(report).items(), key=lambda item: len(item[0]), reverse=True):
        escaped = re.escape(token)
        patterns = [re.compile(r"\{\{\s*" + escaped + r"\s*\}\}", re.IGNORECASE)]
        if token in PLAIN_METADATA_TOKENS:
            patterns.append(re.compile(r"(?<![\w-])" + escaped + r"(?![\w-])", re.IGNORECASE))
        for root in _document_roots(document):
            for paragraph in root.iter(qn("w:p")):
                for pattern in patterns:
                    replace_pattern_across_text_nodes(paragraph, pattern, value)


def _replace_cell_placeholder(
    cell,
    value: str,
    *,
    font_color: RGBColor | None = None,
    font_size_pt: int | None = None,
) -> None:
    paragraphs = [
        paragraph
        for paragraph in cell._tc.iter(qn("w:p"))
        if CELL_PLACEHOLDER.search("".join(node.text or "" for node in paragraph.iter(qn("w:t"))))
    ]
    if len(paragraphs) != 1:
        raise ReportGenerationError(
            f"Expected exactly one placeholder in summary cell; found {len(paragraphs)}"
        )
    replace_pattern_across_text_nodes(
        paragraphs[0],
        CELL_PLACEHOLDER,
        value,
        font_color,
        font_size_pt,
    )


def _find_table(document: DocumentType, header: str) -> Table:
    for table in document.tables:
        if table.rows and table.rows[0].cells and table.rows[0].cells[0].text.strip().casefold() == header.casefold():
            return table
    raise ReportGenerationError(f"Template table not found: {header}")


def _optional_table(document: DocumentType, header: str) -> Table | None:
    """For a table only some of the four templates carry.

    Keyed on the table being present rather than on the report, so rendering stays a function of
    the template it was handed -- the scripts pass MAIN.docx to reports of every shape."""
    try:
        return _find_table(document, header)
    except ReportGenerationError:
        return None


def _append_prototype_row(table: Table, prototype):
    """Clone the template's data row so each entry keeps that row's own formatting.

    Stacking values as paragraphs inside one cell instead would misalign every row below the first
    value that wraps -- silently, in the delivered document."""
    row_element = deepcopy(prototype)
    table._tbl.append(row_element)
    return _Row(row_element, table)


def _strip_data_rows(table: Table):
    prototype = deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    return prototype


def _clear_paragraph(paragraph: Paragraph) -> None:
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def _wrap_long_value(value: str, limit: int) -> list[str]:
    """Break a long location so no line can widen its column.

    Count to the limit, walk back to the first character that is not a letter or
    digit, and break after it; counting then restarts from the break. A stretch
    offering no such character is cut at the limit rather than left to overflow.
    Characters inside "://" are skipped so a long host cannot strand the scheme on a
    line of its own.
    """
    scheme = value.find("://")
    offset = scheme + 3 if scheme != -1 else 0
    segments = []
    remaining = value
    while len(remaining) > limit:
        window = remaining[:limit]
        cut = next(
            (index + 1 for index in range(limit - 1, offset - 1, -1) if not window[index].isalnum()),
            limit,
        )
        segments.append(remaining[:cut])
        remaining = remaining[cut:]
        offset = 0
    segments.append(remaining)
    return segments


def _add_wrapped_run(paragraph, value: str, limit: int | None) -> None:
    """Add the value as one run, breaking inside it rather than starting a new bullet."""
    segments = _wrap_long_value(value, limit) if limit else [value]
    run = paragraph.add_run(segments[0])
    for segment in segments[1:]:
        run.add_break()
        run.add_text(segment)


def _set_cell_lines(cell, lines: list[str], *, wrap: int | None = None) -> None:
    paragraphs = cell.paragraphs
    first = paragraphs[0]
    paragraph_style = first.style
    for paragraph in paragraphs[1:]:
        cell._tc.remove(paragraph._p)
    _clear_paragraph(first)
    # An added paragraph inherits the style but not the template paragraph's direct
    # formatting, so every line after the first would lose its alignment and indent.
    template_properties = first._p.find(qn("w:pPr"))
    for index, value in enumerate(lines or ["N/A"]):
        if index == 0:
            paragraph = first
        else:
            paragraph = cell.add_paragraph(style=paragraph_style)
            if template_properties is not None:
                existing = paragraph._p.find(qn("w:pPr"))
                if existing is not None:
                    paragraph._p.remove(existing)
                paragraph._p.insert(0, deepcopy(template_properties))
        _add_wrapped_run(paragraph, value, wrap)


def _center_plain(cell) -> None:
    """The scope cells use a bullet-list style, so a placeholder needs the plain style to lose its bullet."""
    for paragraph in cell.paragraphs:
        paragraph.style = "Normal"
        properties = paragraph._p.find(qn("w:pPr"))
        if properties is not None:
            for tag in ("w:numPr", "w:ind"):
                for element in properties.findall(qn(tag)):
                    properties.remove(element)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _populate_scope_tables(document: DocumentType, report: Report) -> None:
    engagement = report.engagement
    web = _find_table(document, "URL(s) in Scope")
    _set_cell_lines(web.cell(2, 0), _target_values(report, "production", "web"), wrap=SCOPE_WRAP_CHARACTERS)
    _set_cell_lines(web.cell(4, 0), _target_values(report, "non_production", "web"), wrap=SCOPE_WRAP_CHARACTERS)
    api = _find_table(document, "API Routes")
    _set_cell_lines(api.cell(2, 0), _target_values(report, "production", "api"), wrap=SCOPE_WRAP_CHARACTERS)
    _set_cell_lines(api.cell(4, 0), _target_values(report, "non_production", "api"), wrap=SCOPE_WRAP_CHARACTERS)
    limitations = _find_table(document, "Limitations")
    limitation_text = _display_value(engagement.limitations)
    _set_cell_lines(limitations.cell(1, 0), [limitation_text])
    if limitation_text.casefold() in {"n/a", "na"}:
        _center_plain(limitations.cell(1, 0))

    accounts = _find_table(document, "User Roles")
    prototype = _strip_data_rows(accounts)
    for account in engagement.test_accounts or []:
        row = _append_prototype_row(accounts, prototype)
        _set_cell_lines(row.cells[0], [_display_value(account.user_role)])
        _set_cell_lines(row.cells[1], [_display_value(account.username)])
    if len(accounts.rows) == 1:
        row = _append_prototype_row(accounts, prototype)
        _set_cell_lines(row.cells[0], ["N/A"])
        _set_cell_lines(row.cells[1], ["N/A"])

    components = _optional_table(document, "Component")
    if components is not None:
        prototype = _strip_data_rows(components)
        for value, description in _component_rows(report) or [("N/A", "N/A")]:
            row = _append_prototype_row(components, prototype)
            _set_cell_lines(row.cells[0], [_display_value(value)], wrap=SCOPE_WRAP_CHARACTERS)
            _set_cell_lines(row.cells[1], [_display_value(description)])


def _populate_summary_table(document: DocumentType, report: Report) -> None:
    table = _find_table(document, "Findings")
    generic_prototype = len(table.rows) == 2 and "{{finding}}" in table.rows[1].cells[0].text.casefold()
    if generic_prototype:
        prototype = deepcopy(table.rows[1]._tr)
        prototypes = {severity: prototype for severity in SEVERITY_ORDER}
    else:
        prototypes = {
            severity: deepcopy(table.rows[index + 1]._tr)
            for index, severity in enumerate(SEVERITY_ORDER)
        }
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for finding in sorted(report.vulnerabilities, key=lambda item: (SEVERITY_ORDER.index(item.severity or "informational"), item.title.casefold())):
        severity = finding.severity or "informational"
        row = _append_prototype_row(table, prototypes[severity])
        likelihood = finding.likelihood or "informational"
        impact = finding.impact or "informational"
        values = [
            (finding.title, None),
            (likelihood.title(), RATING_FONT_COLORS[likelihood]),
            (impact.title(), RATING_FONT_COLORS[impact]),
            (severity.title(), RATING_FONT_COLORS[severity]),
            # The internal uid is not a finding number, so an unnumbered finding stays blank.
            (finding.display_id or "", None),
            (STATUS_LABELS[finding.status], None),
        ]
        for cell, (value, font_color) in zip(row.cells, values):
            _replace_cell_placeholder(
                cell,
                value,
                font_color=font_color,
                font_size_pt=12 if font_color is not None else None,
            )


def _element_text(element) -> str:
    return re.sub(r"\s+", " ", "".join(node.text or "" for node in element.iter(qn("w:t")))).strip()


def _has_exact_body_token(document: DocumentType, token: str) -> bool:
    pattern = re.compile(r"\{\{\s*" + re.escape(token) + r"\s*\}\}", re.IGNORECASE)
    return any(
        element.tag == qn("w:p") and pattern.fullmatch(_element_text(element))
        for element in document.element.body.iterchildren()
    )


def _find_body_element(document: DocumentType, text: str):
    for element in document.element.body.iterchildren():
        if _element_text(element).casefold() == text.casefold():
            return element
    raise ReportGenerationError(f"Template block not found: {text}")


def _normalize_page_numbering(document: DocumentType) -> None:
    body = document.element.body
    document_order = list(body.iter())
    revision_index = document_order.index(_find_body_element(document, "DOCUMENT REVISION HISTORY"))
    section_properties = [
        properties
        for properties in body.iter(qn("w:sectPr"))
        if not any(ancestor.tag == qn("w:sectPrChange") for ancestor in properties.iterancestors())
    ]
    revision_section = next(
        properties
        for properties in section_properties
        if document_order.index(properties) > revision_index
    )
    revision_reached = False
    for properties in section_properties:
        page_numbering = properties.find(qn("w:pgNumType"))
        if properties is revision_section:
            revision_reached = True
            if page_numbering is None:
                page_numbering = OxmlElement("w:pgNumType")
                properties.append(page_numbering)
            page_numbering.set(qn("w:start"), "1")
        elif revision_reached and page_numbering is not None:
            page_numbering.attrib.pop(qn("w:start"), None)
            if not page_numbering.attrib and not len(page_numbering):
                properties.remove(page_numbering)


def _paragraph_style_id(element) -> str | None:
    properties = element.find(qn("w:pPr"))
    style = properties.find(qn("w:pStyle")) if properties is not None else None
    return style.get(qn("w:val")) if style is not None else None


def _bookmark_paragraph(paragraph, name: str, bookmark_id: int) -> None:
    """Wrap a heading so a REF field elsewhere can quote the number Word gives it."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph.insert(0, start)
    paragraph.append(end)


def _reference_field(paragraph: Paragraph, bookmark: str) -> None:
    """Write " REF <bookmark> \\w \\h " as a real field, so Word reports the heading's own number.

    Counting headings in Python instead would silently disagree with what Word renders. \\w rather
    than \\r because the reference sits in a different section, where a relative number is short."""
    _clear_paragraph(paragraph)
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = f" REF {bookmark} \\w \\h "
    for kind, child in (("begin", None), (None, instruction), ("separate", None), ("end", None)):
        run = OxmlElement("w:r")
        if kind is not None:
            marker = OxmlElement("w:fldChar")
            marker.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                # docx_captions marks every begin dirty, which is what makes Word compute the value.
                marker.set(qn("w:dirty"), "true")
            run.append(marker)
        else:
            run.append(child)
        paragraph._p.append(run)


def _populate_cvss_table(document: DocumentType, rendered: list[tuple[Vulnerability, str]]) -> None:
    """Fill the Asia-only findings table. Absent from the other two templates, and from any report
    the scripts render against MAIN.docx, so the table's presence is what selects this."""
    table = _optional_table(document, "Section")
    if table is None:
        return
    prototype = _strip_data_rows(table)
    for finding, bookmark in rendered or []:
        row = _append_prototype_row(table, prototype)
        severity = finding.severity or "informational"
        _reference_field(row.cells[0].paragraphs[0], bookmark)
        _replace_cell_placeholder(row.cells[1], finding.title)
        _replace_cell_placeholder(row.cells[2], severity.title(), font_color=RATING_FONT_COLORS[severity])
        _replace_cell_placeholder(row.cells[3], finding.cvss_score)
        _replace_cell_placeholder(row.cells[4], finding.cvss_vector)
    if not rendered:
        _append_prototype_row(table, prototype)
        for index in range(5):
            _replace_cell_placeholder(table.rows[1].cells[index], "")


def _populate_component_findings(
    document: DocumentType,
    report: Report,
    report_folder: Path,
    component_root: Path,
    allow_incomplete: bool,
) -> list[tuple[Vulnerability, str]]:
    body_elements = list(document.element.body.iterchildren())
    anchor_index = _component_anchor_index(body_elements, "findings")
    anchor = body_elements[anchor_index]
    rendered_order: list[tuple[Vulnerability, str]] = []
    for severity in SEVERITY_ORDER:
        findings = sorted(
            (item for item in report.vulnerabilities if item.severity == severity),
            key=lambda item: item.title.casefold(),
        )
        if not findings:
            continue
        severity_elements = clone_component_elements(
            document,
            component_root / "severity_titles" / SEVERITY_COMPONENT_FILES[severity],
        )
        finding_anchor = _component_anchor_index(severity_elements, "finding_title")
        rendered_findings = []
        for index, finding in enumerate(findings):
            finding_elements = _render_finding_component(
                document,
                finding,
                report,
                report_folder,
                component_root,
                allow_incomplete,
            )
            title_paragraph = next(
                (element for element in finding_elements if element.tag == qn("w:p") and _paragraph_style_id(element) == FINDING_HEADING_STYLE),
                None,
            )
            if title_paragraph is not None:
                bookmark = f"vuln_{finding.uid}"
                _bookmark_paragraph(title_paragraph, bookmark, len(rendered_order) + 1)
                rendered_order.append((finding, bookmark))
            if index:
                first_paragraph = next(
                    (element for element in finding_elements if element.tag == qn("w:p")),
                    None,
                )
                if first_paragraph is not None:
                    _set_page_break_before(first_paragraph)
            rendered_findings.extend(finding_elements)
        severity_elements = [
            *severity_elements[:finding_anchor],
            *rendered_findings,
            *severity_elements[finding_anchor + 1:],
        ]
        heading = next(
            (element for element in severity_elements if element.tag == qn("w:p")),
            None,
        )
        if heading is not None:
            _set_page_break_before(heading)
        for element in severity_elements:
            anchor.addprevious(element)
    anchor.getparent().remove(anchor)
    return rendered_order


def _replace_token_with_bullets(
    document: DocumentType,
    elements: list,
    token: str,
    values: list[str],
    component_root: Path,
) -> None:
    """Swap a cell token for real bullet-list paragraphs cloned from the bullet fragment."""
    pattern = re.compile(r"\{\{\s*" + re.escape(token) + r"\s*\}\}", re.IGNORECASE)
    if not values:
        replace_component_token_runs(elements, token, [Run(text="N/A")])
        return
    filename, bullet_token = FRAGMENT_COMPONENT_FILES["bulleted_list"]
    for element in elements:
        for paragraph in list(element.iter(qn("w:p"))):
            if not pattern.search(_element_text(paragraph)):
                continue
            numbering_ids: dict[tuple[Path, int], int] = {}
            for value in values:
                bullet = clone_component_elements(
                    document,
                    component_root / "fragments" / filename,
                    numbering_ids=numbering_ids,
                )
                replace_component_token_runs(
                    bullet,
                    bullet_token,
                    [Run(text="\n".join(_wrap_long_value(value, LOCATION_WRAP_CHARACTERS)))],
                )
                for node in bullet:
                    for bullet_paragraph in node.iter(qn("w:p")):
                        _align_left(bullet_paragraph)
                    paragraph.addprevious(node)
            paragraph.getparent().remove(paragraph)
            return


def _align_left(paragraph_element) -> None:
    properties = paragraph_element.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        paragraph_element.insert(0, properties)
    for existing in properties.findall(qn("w:jc")):
        properties.remove(existing)
    alignment = OxmlElement("w:jc")
    alignment.set(qn("w:val"), "left")
    properties.append(alignment)


def _render_finding_component(
    document: DocumentType,
    finding: Vulnerability,
    report: Report,
    report_folder: Path,
    component_root: Path,
    allow_incomplete: bool,
) -> list:
    template_name = "new_finding.docx" if finding.status == "open_new" else "retest_finding.docx"
    elements = clone_component_elements(
        document,
        component_root / "finding_types" / template_name,
    )
    locations = _finding_locations(report, finding)
    values = {
        "finding_title": finding.title,
        "vuln_severity": (finding.severity or "informational").title(),
        "vuln_id": finding.display_id or "",
        "status": STATUS_LABELS[finding.status],
    }
    for token, value in values.items():
        replace_component_token(elements, token, value)
    # Runs, not the string path: only _set_run_text turns a newline into a real w:br. A no-op on
    # new_finding.docx, which carries no such token.
    tickets = [line.strip() for line in finding.severity_review_tickets.splitlines() if line.strip()]
    printed = "\n".join(f"{SEVERITY_TICKET_PREFIX}{ticket}" for ticket in tickets) or "N/A"
    replace_component_token_runs(elements, "severity-review-tickets", [Run(text=printed)])
    for token, location_values in (
        ("prod_affected_locations", locations["production"]),
        ("non-prod-affected-locations", locations["non_production"]),
    ):
        _replace_token_with_bullets(document, elements, token, location_values, component_root)

    contents = {content.type: content for content in finding.contents}
    anchors = {
        "description-fragments-here": contents.get("description"),
        "recommended-remediation-fragments-here": contents.get("recommended_remediation"),
        "poc-fragments-here": contents.get("proof_of_concept"),
    }
    if finding.status != "open_new":
        anchors.update({
            "prev-poc-fragments-here": contents.get("previous_proof_of_concept"),
            "conclusion-fragments-here": contents.get("in_conclusion"),
        })
    for anchor, content in anchors.items():
        rendered = _render_component_content(
            document,
            content,
            report,
            report_folder,
            component_root,
            allow_incomplete,
            finding=finding,
            label_images=anchor in {"poc-fragments-here", "prev-poc-fragments-here"},
        )
        # Before the splice: the anchor has to still be present, and re-located each pass because
        # an earlier replacement has already shifted the indices.
        _keep_anchor_lead_in(elements, anchor)
        elements = _replace_component_anchor(elements, anchor, rendered)
    return elements


def _environment_label(report: Report, environment: str) -> str:
    return "PROD:" if environment == "production" else f"{report.engagement.non_production_label.upper()}:"


def _render_environment_label(
    document: DocumentType,
    component_root: Path,
    report: Report,
    environment: str,
) -> list:
    elements = _render_text_component(
        document,
        component_root,
        "instance_title",
        [Run(text=_environment_label(report, environment))],
    )
    for element in elements:
        if element.tag == qn("w:p"):
            Paragraph(element, document._body).paragraph_format.space_after = Pt(0)
            _keep_with_next(element)
    return elements


def _render_instance_title(
    document: DocumentType,
    component_root: Path,
    fragment: InstanceTitleFragment,
    number: int,
) -> list:
    """The number is the report's and the title is the tester's, so an untitled instance still reads."""
    # Drafts written before the label was generated have it typed into the title by hand.
    title = INSTANCE_PREFIX.sub("", fragment.text.strip())
    # The label is the app's own and stays bold; the tester's words are content, not heading.
    runs = [Run(text=f"Instance {number}:", bold=True)]
    if title:
        runs.append(Run(text=f" {title}"))
    elements = _render_text_component(
        document,
        component_root,
        "instance_title",
        runs,
    )
    for element in elements:
        if element.tag == qn("w:p"):
            _keep_with_next(element)
            _unbold_runs_after_first(element)
    return elements


def _unbold_runs_after_first(paragraph) -> None:
    """Run formatting is additive, so the tester's words inherit the component's bold."""
    text_runs = [run for run in paragraph.findall(qn("w:r")) if run.find(qn("w:t")) is not None]
    for run in text_runs[1:]:
        properties = run.find(qn("w:rPr"))
        if properties is None:
            properties = OxmlElement("w:rPr")
            run.insert(0, properties)
        for tag in ("w:b", "w:bCs"):
            setting = properties.find(qn(tag))
            if setting is None:
                setting = OxmlElement(tag)
                properties.append(setting)
            # Off, not absent: the paragraph style may declare bold too.
            setting.set(qn("w:val"), "0")


def _render_component_content(
    document: DocumentType,
    content: Content | None,
    report: Report,
    report_folder: Path,
    component_root: Path,
    allow_incomplete: bool,
    *,
    finding: Vulnerability | None = None,
    label_images: bool = False,
) -> list:
    fragments = [] if content is None else [
        fragment
        for fragment in content.fragments
        if finding is None or fragment_applies(fragment, finding, report, content.type)
    ]
    if not fragments:
        rendered = _render_text_component(
            document,
            component_root,
            "paragraph",
            [Run(text="N/A")],
        )
    else:
        rendered = []
        labelled_environment = None
        # Declared here, in the body, and nowhere higher: this is what confines a continued chain to
        # one section. _render_component_content runs once per anchor per finding, so a numbered list
        # can never continue one from another section or another finding.
        numbering_carry: dict[tuple[Path, int], int] | None = None
        instance_number = 0
        for fragment in fragments:
            # Numbered here rather than in the fragment renderer, so the count is confined to this
            # section the same way the numbered-list chain above is.
            if isinstance(fragment, InstanceTitleFragment):
                instance_number += 1
                rendered.extend(_render_instance_title(document, component_root, fragment, instance_number))
                continue
            # One label per run of images.
            if (
                label_images
                and isinstance(fragment, ImageFragment)
                and fragment.environment
                and fragment.environment != labelled_environment
            ):
                labelled_environment = fragment.environment
                rendered.extend(_render_environment_label(document, component_root, report, fragment.environment))
            if isinstance(fragment, ListFragment) and fragment.type == "numbered_list":
                # Anything else between two lists leaves the carry alone, so a chain survives an
                # image or an instance title -- the shape this option exists for.
                if not (fragment.continue_numbering and numbering_carry is not None):
                    numbering_carry = {}
                fragment_numbering = numbering_carry
            else:
                fragment_numbering = None
            rendered.extend(
                _render_component_fragment(
                    document,
                    fragment,
                    report,
                    report_folder,
                    component_root,
                    allow_incomplete,
                    numbering_ids=fragment_numbering,
                )
            )
    rendered = _trim_trailing_empty_paragraphs(rendered)
    _keep_tables_with_lead_in(rendered)
    return rendered


def _render_component_fragment(
    document: DocumentType,
    fragment,
    report: Report,
    report_folder: Path,
    component_root: Path,
    allow_incomplete: bool,
    *,
    # Never `{}`: a mutable default is built once at import and would join every list in the report
    # into one chain. None means "your own cache", which is what makes each list restart.
    numbering_ids: dict[tuple[Path, int], int] | None = None,
) -> list:
    if isinstance(fragment, ParagraphFragment):
        return _render_multiline_text_component(document, component_root, "paragraph", fragment.runs)
    if isinstance(fragment, NoteFragment):
        return _render_multiline_text_component(document, component_root, "note", fragment.runs)
    if isinstance(fragment, ListFragment):
        component_type = "numbered_list" if fragment.type == "numbered_list" else "bulleted_list"
        filename, token = FRAGMENT_COMPONENT_FILES[component_type]
        # A cache hit inside _remap_numbering reuses the allocated numId and writes no startOverride,
        # which is exactly what Word reads as one continuing list. _restart_numbering_levels takes its
        # levels from the chain's first clone group; harmless while every clone is the same component.
        if numbering_ids is None:
            numbering_ids = {}
        rendered = []
        for item in fragment.items:
            elements = clone_component_elements(
                document,
                component_root / "fragments" / filename,
                numbering_ids=numbering_ids,
            )
            replace_component_token_runs(elements, token, item.runs)
            rendered.extend(elements)
        # Only the component's first element: it ships with a trailing blank that would make two.
        if fragment.type == "numbered_list":
            rendered.extend(_render_text_component(document, component_root, "paragraph", [])[:1])
        return rendered
    if isinstance(fragment, CodeFragment):
        rendered = _render_caption_component(document, component_root, fragment.caption)
        for element in rendered:
            if element.tag == qn("w:p"):
                _keep_with_next(element)
        rendered.extend(
            _render_text_component(
                document,
                component_root,
                "code_block",
                [Run(text=fragment.text)],
            )
        )
        return rendered
    if isinstance(fragment, InstanceTitleFragment):
        rendered = _render_text_component(
            document,
            component_root,
            "instance_title",
            [Run(text=fragment.text)],
        )
        for element in rendered:
            if element.tag == qn("w:p"):
                _keep_with_next(element)
        return rendered
    if isinstance(fragment, TableFragment):
        return _render_table_component(document, fragment, component_root)
    if isinstance(fragment, ImageFragment):
        return _render_image_component(
            document,
            fragment,
            report,
            report_folder,
            component_root,
            allow_incomplete,
        )
    return []


def _render_text_component(
    document: DocumentType,
    component_root: Path,
    component_type: str,
    runs: list[Run],
) -> list:
    filename, token = FRAGMENT_COMPONENT_FILES[component_type]
    elements = clone_component_elements(document, component_root / "fragments" / filename)
    replace_component_token_runs(elements, token, runs)
    _keep_broken_lines_together(elements)
    return elements


def _keep_broken_lines_together(elements: list) -> None:
    """A manual break means the lines belong together, as a lead-in does with the URL beneath it.
    keepNext cannot help: Word is splitting inside one paragraph, not between two."""
    for element in elements:
        if element.tag != qn("w:p"):
            continue
        if element.find(qn("w:r")) is None or not element.findall(".//" + qn("w:br")):
            continue
        _keep_lines_together(element)


def _split_runs_at_newlines(runs: list[Run]) -> list[list[Run]]:
    lines: list[list[Run]] = [[]]
    skip_leading_line_feed = False
    for source in runs:
        text = source.text
        index = 0
        if skip_leading_line_feed and text.startswith("\n"):
            index = 1
        skip_leading_line_feed = False
        start = index
        while index < len(text):
            if text[index] not in "\r\n":
                index += 1
                continue
            if start < index:
                lines[-1].append(source.model_copy(update={"text": text[start:index]}))
            if text[index] == "\r":
                if index + 1 < len(text) and text[index + 1] == "\n":
                    index += 1
                elif index + 1 == len(text):
                    skip_leading_line_feed = True
            lines.append([])
            index += 1
            start = index
        if start < len(text):
            lines[-1].append(source.model_copy(update={"text": text[start:]}))
    return lines


def _render_multiline_text_component(
    document: DocumentType,
    component_root: Path,
    component_type: str,
    runs: list[Run],
) -> list:
    filename, token = FRAGMENT_COMPONENT_FILES[component_type]
    elements = clone_component_elements(document, component_root / "fragments" / filename)
    pattern = re.compile(r"\{\{\s*" + re.escape(token) + r"\s*\}\}", re.IGNORECASE)
    matches = [
        (index, element)
        for index, element in enumerate(elements)
        if element.tag == qn("w:p") and pattern.search(_element_text(element))
    ]
    if len(matches) != 1:
        raise ReportGenerationError(f"Expected exactly one {token} paragraph in {filename}")

    paragraph_index, paragraph = matches[0]
    template = deepcopy(paragraph)
    lines = _split_runs_at_newlines(runs)
    replace_component_token_runs([paragraph], token, lines[0])
    continuation_paragraphs = []
    for line in lines[1:]:
        continuation = deepcopy(template)
        if component_type == "note":
            replace_pattern_across_text_nodes(
                continuation,
                re.compile(r"^\s*Note:\s*", re.IGNORECASE),
                "",
            )
        replace_component_token_runs([continuation], token, line)
        continuation_paragraphs.append(continuation)
    elements[paragraph_index + 1:paragraph_index + 1] = continuation_paragraphs
    return elements


def _trim_trailing_empty_paragraphs(elements: list) -> list:
    end = len(elements)
    while end and _is_empty_component_paragraph(elements[end - 1]):
        end -= 1
    return elements[:end]


def _keep_tables_with_lead_in(elements: list) -> None:
    """Tie each table to the paragraph that introduces it.

    Runs after the trim, so a table ending a section has already lost its spacer. A caption is not
    enough on its own: the caption component carries a trailing blank, and most tables have no
    caption at all, so the walk reaches back to the real lead-in sentence.
    """
    for index, element in enumerate(elements):
        if element.tag != qn("w:tbl"):
            continue
        for previous in reversed(elements[:index]):
            if previous.tag != qn("w:p"):
                break
            _keep_with_next(previous)
            if _element_text(previous).strip():
                break


def _is_empty_component_paragraph(element) -> bool:
    return (
        element.tag == qn("w:p")
        and not _element_text(element)
        and not any(True for _ in element.iter(qn("w:drawing")))
        and not any(True for _ in element.iter(qn("w:pict")))
        and not any(True for _ in element.iter(qn("w:br")))
    )


def _render_image_caption(
    document: DocumentType,
    component_root: Path,
    caption: str | None,
) -> list:
    """Centred body text; the post-processor finds it by that, since it carries no caption style."""
    if not caption:
        return []
    elements = _render_text_component(
        document,
        component_root,
        "paragraph",
        [Run(text=caption)],
    )
    for element in elements:
        if element.tag == qn("w:p"):
            center_paragraph(element)
    return elements


def _render_caption_component(
    document: DocumentType,
    component_root: Path,
    caption: str | None,
) -> list:
    if not caption:
        return []
    return _render_text_component(
        document,
        component_root,
        "caption",
        [Run(text=caption)],
    )


def _render_image_component(
    document: DocumentType,
    fragment: ImageFragment,
    report: Report,
    report_folder: Path,
    component_root: Path,
    allow_incomplete: bool,
) -> list:
    evidence = report.evidence.get(fragment.evidence_id or "")
    if evidence is None:
        if not allow_incomplete:
            raise ReportGenerationError(f"Missing evidence for image fragment {fragment.frag_id}")
        return _render_text_component(
            document,
            component_root,
            "paragraph",
            [Run(text=f"[Evidence image not attached: {_display_value(fragment.caption)}]")],
        )
    path = report_folder / evidence.file
    if not path.is_file():
        raise ReportGenerationError(f"Evidence file not found: {evidence.file}")
    image_elements = _render_text_component(document, component_root, "paragraph", [])
    paragraph_element = next(element for element in image_elements if element.tag == qn("w:p"))
    paragraph = Paragraph(paragraph_element, document._body)
    image_width_mm = min(fragment.width_mm or 155, 155)
    image_template = component_root / "fragments" / "image_fragment.docx"
    picture = paragraph.add_run().add_picture(
        _bordered_image_stream(path, image_width_mm, image_template),
        width=Mm(image_width_mm),
    )
    _apply_image_fragment_format(
        picture,
        paragraph,
        image_template,
    )
    image_elements = _trim_trailing_empty_paragraphs(image_elements)
    caption_elements = _render_image_caption(document, component_root, _display_value(fragment.caption))
    if caption_elements:
        _keep_with_next(paragraph_element)
    return [*image_elements, *caption_elements]


def _render_table_component(
    document: DocumentType,
    fragment: TableFragment,
    component_root: Path,
) -> list:
    rendered = _render_caption_component(document, component_root, fragment.caption)
    elements = clone_component_elements(
        document,
        component_root / "fragments" / "table_fragment.docx",
    )
    table_element = next(
        (element for element in elements if element.tag == qn("w:tbl")),
        None,
    )
    if table_element is None:
        raise ReportGenerationError("table_fragment.docx must contain one table")
    source_rows = list(table_element.iterchildren(qn("w:tr")))
    if len(source_rows) != 2:
        raise ReportGenerationError("table_fragment.docx must contain header and body prototype rows")
    source_widths = [
        int(column.get(qn("w:w")))
        for column in table_element.tblGrid.iterchildren(qn("w:gridCol"))
    ]
    if not source_widths:
        raise ReportGenerationError("table_fragment.docx must define a column grid")
    columns = max([len(fragment.header), *(len(row) for row in fragment.rows), 1])
    widths = source_widths if columns == len(source_widths) else _equal_component_widths(sum(source_widths), columns)
    for row in source_rows:
        table_element.remove(row)
    _set_component_table_grid(table_element, widths)
    if fragment.header:
        table_element.append(_resize_component_row(source_rows[0], widths))
    for _ in fragment.rows:
        table_element.append(_resize_component_row(source_rows[1], widths))
    table = Table(table_element, document._body)
    table.autofit = False
    row_offset = 0
    if fragment.header:
        for index, cell in enumerate(table.rows[0].cells):
            runs = fragment.header[index].runs if index < len(fragment.header) else []
            replace_component_token_runs([cell._tc], "table-header-cell", runs)
        row_offset = 1
    for row_index, row in enumerate(table.rows[row_offset:]):
        source_row = fragment.rows[row_index]
        for column_index, cell in enumerate(row.cells):
            runs = source_row[column_index].runs if column_index < len(source_row) else []
            replace_component_token_runs([cell._tc], "table-body-cell", runs)
    rendered.append(table_element)
    # A table fragment contributes only its w:tbl, so without this two in a row become one in Word.
    rendered.append(
        next(
            element
            for element in _render_text_component(document, component_root, "paragraph", [])
            if element.tag == qn("w:p")
        )
    )
    return rendered


def _equal_component_widths(total_width: int, columns: int) -> list[int]:
    width, remainder = divmod(total_width, columns)
    return [width + (1 if index < remainder else 0) for index in range(columns)]


def _set_component_table_grid(table, widths: list[int]) -> None:
    grid = table.tblGrid
    for column in list(grid.iterchildren(qn("w:gridCol"))):
        grid.remove(column)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    table_width = table.tblPr.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        table.tblPr.insert(0, table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(sum(widths)))
    layout = table.tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table.tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _resize_component_row(source_row, widths: list[int]):
    row = deepcopy(source_row)
    source_cells = list(source_row.iterchildren(qn("w:tc")))
    if not source_cells:
        raise ReportGenerationError("table_fragment.docx contains a prototype row without cells")
    for cell in list(row.iterchildren(qn("w:tc"))):
        row.remove(cell)
    preserve_columns = len(source_cells) == len(widths)
    for index, width in enumerate(widths):
        source_cell = source_cells[index] if preserve_columns else source_cells[min(index, len(source_cells) - 1)]
        cell = deepcopy(source_cell)
        properties = cell.find(qn("w:tcPr"))
        if properties is None:
            properties = OxmlElement("w:tcPr")
            cell.insert(0, properties)
        cell_width = properties.find(qn("w:tcW"))
        if cell_width is None:
            cell_width = OxmlElement("w:tcW")
            properties.insert(0, cell_width)
        cell_width.set(qn("w:type"), "dxa")
        cell_width.set(qn("w:w"), str(width))
        row.append(cell)
    return row


def _replace_component_anchor(elements: list, token: str, replacements: list) -> list:
    index = _component_anchor_index(elements, token)
    return [*elements[:index], *replacements, *elements[index + 1:]]


def _component_anchor_index(elements: list, token: str) -> int:
    pattern = re.compile(r"\{\{\s*" + re.escape(token) + r"\s*\}\}", re.IGNORECASE)
    matches = [
        index
        for index, element in enumerate(elements)
        if element.tag == qn("w:p") and pattern.fullmatch(_element_text(element))
    ]
    if len(matches) != 1:
        raise ReportGenerationError(
            f"Expected exactly one {{{{{token}}}}} component anchor; found {len(matches)}"
        )
    return matches[0]


def _keep_anchor_lead_in(elements: list, token: str) -> None:
    """Keep a section heading with the content that replaces its anchor.

    Walks by position rather than by heading text, so docx_import's SECTION_HEADINGS stays the only
    place those strings live. "Severity Review Ticket" has no anchor and is deliberately not covered.
    """
    index = _component_anchor_index(elements, token)
    for element in reversed(elements[:index]):
        if element.tag != qn("w:p") or not _element_text(element).strip():
            break
        _keep_with_next(element)


def _set_page_break_before(paragraph_element) -> None:
    paragraph_properties = paragraph_element.find(qn("w:pPr"))
    if paragraph_properties is None:
        paragraph_properties = OxmlElement("w:pPr")
        paragraph_element.insert(0, paragraph_properties)
    if paragraph_properties.find(qn("w:pageBreakBefore")) is None:
        paragraph_properties.append(OxmlElement("w:pageBreakBefore"))


def _keep_lines_together(paragraph_element) -> None:
    properties = paragraph_element.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        paragraph_element.insert(0, properties)
    if properties.find(qn("w:keepLines")) is not None:
        return
    keep = OxmlElement("w:keepLines")
    # CT_PPr is a sequence: w:keepLines sits after w:pStyle and w:keepNext, before everything else.
    anchor = properties.find(qn("w:keepNext")) or properties.find(qn("w:pStyle"))
    if anchor is None:
        properties.insert(0, keep)
    else:
        anchor.addnext(keep)


def _keep_with_next(paragraph_element) -> None:
    """Word decides automatic page breaks itself and never records them in the file, so keeping a
    heading beside the content it introduces has to be declared here rather than measured."""
    properties = paragraph_element.find(qn("w:pPr"))
    if properties is None:
        properties = OxmlElement("w:pPr")
        paragraph_element.insert(0, properties)
    if properties.find(qn("w:keepNext")) is not None:
        return
    keep = OxmlElement("w:keepNext")
    # CT_PPr is a sequence: w:keepNext has to follow w:pStyle and precede every other child.
    style = properties.find(qn("w:pStyle"))
    if style is None:
        properties.insert(0, keep)
    else:
        style.addnext(keep)


def _finding_locations(report: Report, finding: Vulnerability) -> dict[str, list[str]]:
    values = {"production": [], "non_production": []}
    targets = {target.target_id: target for target in report.scope_targets}
    for target_id in finding.scope.target_ids:
        target = targets.get(target_id)
        if target:
            values[target.environment].append(finding.scope.location_values.get(target_id, target.value))
    # Typed endpoints are labelled "additional", so they print after whatever was selected.
    for environment, by_channel in finding.scope.custom_locations.items():
        for channel in CHANNEL_ORDER:
            for value in location_lines(by_channel.get(channel, [])):
                if value not in values[environment]:
                    values[environment].append(value)
    return values


def _bordered_image_stream(path: Path, width_mm: float, template_path: Path) -> BytesIO:
    source_line, source_alignment = _cached_image_fragment_format(
        str(template_path.resolve()),
        template_path.stat().st_mtime_ns,
    )
    del source_alignment
    border_width_inches = int(source_line.get("w") or Pt(0.75)) / 914400
    display_width_inches = width_mm / 25.4
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
    minimum_width = round(display_width_inches * IMAGE_BORDER_RASTER_DPI)
    if image.width < minimum_width:
        image = image.resize(
            (minimum_width, max(1, round(image.height * minimum_width / image.width))),
            Image.Resampling.LANCZOS,
        )
    border_fraction = border_width_inches / display_width_inches
    border_pixels = max(1, round(image.width * border_fraction / (1 - 2 * border_fraction)))
    bordered = ImageOps.expand(image, border=border_pixels, fill="black")
    stream = BytesIO()
    bordered.save(stream, format="PNG")
    stream.seek(0)
    return stream


def _apply_image_fragment_format(picture, paragraph: Paragraph, template_path: Path) -> None:
    source_line, source_alignment = _cached_image_fragment_format(
        str(template_path.resolve()),
        template_path.stat().st_mtime_ns,
    )
    del source_line
    properties = picture._inline.graphic.graphicData.pic.spPr
    line = properties.find(qn("a:ln"))
    if line is not None:
        properties.remove(line)
    paragraph_properties = paragraph._p.get_or_add_pPr()
    alignment = paragraph_properties.find(qn("w:jc"))
    if alignment is not None:
        paragraph_properties.remove(alignment)
    paragraph_properties.append(deepcopy(source_alignment))


@lru_cache(maxsize=8)
def _cached_image_fragment_format(template_path: str, modified_at: int):
    del modified_at
    document = Document(template_path)
    if len(document.inline_shapes) != 1:
        raise ReportGenerationError("image_fragment.docx must contain exactly one inline image")
    source_picture = document.inline_shapes[0]._inline.graphic.graphicData.pic
    source_line = source_picture.spPr.find(qn("a:ln"))
    source_paragraph = next(
        (paragraph for paragraph in document.paragraphs if any(True for _ in paragraph._p.iter(qn("w:drawing")))),
        None,
    )
    source_alignment = (
        source_paragraph._p.pPr.find(qn("w:jc"))
        if source_paragraph is not None and source_paragraph._p.pPr is not None
        else None
    )
    if source_line is None or source_alignment is None:
        raise ReportGenerationError("image_fragment.docx must define an image border and paragraph alignment")
    return deepcopy(source_line), deepcopy(source_alignment)


def _unresolved_placeholders(document: DocumentType) -> list[str]:
    text = "\n".join(_element_text(root) for root in _document_roots(document))
    unresolved = set(re.findall(r"\{\{.*?\}\}", text))
    lower = text.casefold()
    unresolved.update(marker for marker in UNRESOLVED_MARKERS if marker in lower)
    return sorted(unresolved)