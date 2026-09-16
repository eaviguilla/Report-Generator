from __future__ import annotations

from copy import deepcopy
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm
from docx.text.paragraph import Paragraph
from PIL import Image, ImageDraw

from app.docx_captions import postprocess_image_captions
from app.docx_components import DocxComponent, clone_component_elements, compose_docx_components, compose_docx_template, replace_component_token_runs
from app.storage import atomic_write_bytes
from app.models import Run


ROOT = Path(__file__).resolve().parent.parent


def compose_severity_sample() -> Path:
    components = [
        DocxComponent(ROOT / "resources" / "severity_titles" / "critical_severity.docx", {"finding_title": "Critical SQL Injection"}),
        DocxComponent(ROOT / "resources" / "severity_titles" / "high_severity.docx", {"finding_title": "High Authorization Bypass"}),
        DocxComponent(ROOT / "resources" / "severity_titles" / "medium_severity.docx", {"finding_title": "Medium Concurrent Session Allowed"}),
    ]
    output = ROOT / "generated" / "MAIN-composed.docx"
    atomic_write_bytes(
        output,
        compose_docx_components(ROOT / "resources" / "MAIN.docx", components),
    )
    return output


def compose_retest_finding_sample() -> Path:
    fragments = ROOT / "resources" / "fragments"
    output = ROOT / "generated" / "retest-finding-composed.docx"
    atomic_write_bytes(
        output,
        compose_docx_template(
            ROOT / "resources" / "finding_types" / "retest_finding.docx",
            values={
                "finding_title": "Authorization Bypass Retest",
                "vuln_severity": "High",
                "vuln_id": "VULN-001",
                "status": "Open (Previously Discovered)",
                "prod_affected_locations": "https://production.example.test",
                "non-prod-affected-locations": "https://uat.example.test",
                "severity-review-tickets": "N/A",
            },
            components={
                "description-fragments-here": [
                    DocxComponent(fragments / "paragraph_fragment.docx", {"paragraph-fragment": "Authorization checks can be bypassed."}),
                    DocxComponent(fragments / "note_fragment.docx", {"note-fragment": "Retest both affected environments."}),
                ],
                "recommended-remediation-fragments-here": [
                    DocxComponent(fragments / "bulleted_fragment.docx", {"bullet-list-fragment": "Enforce authorization on the server."}),
                    DocxComponent(fragments / "bulleted_fragment.docx", {"bullet-list-fragment": "Add regression coverage."}),
                ],
                "prev-poc-fragments-here": [
                    DocxComponent(fragments / "title_fragment.docx", {"instance-fragment": "Instance 1: Production"}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Sign in as a standard user."}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Request another user's record."}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Change the record identifier."}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Submit the modified request."}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Confirm the unauthorized response data."}),
                ],
                "poc-fragments-here": [
                    DocxComponent(fragments / "title_fragment.docx", {"instance-fragment": "Instance 1: Production"}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Repeat the original request."}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Observe that access is still allowed."}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Apply the proposed authorization control."}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Resend the unauthorized request."}),
                    DocxComponent(fragments / "numbered_fragment.docx", {"numbered-list-fragment": "Confirm that the request is rejected."}),
                ],
                "conclusion-fragments-here": [
                    DocxComponent(fragments / "paragraph_fragment.docx", {"paragraph-fragment": "The finding remains reproducible."}),
                ],
            },
        ),
    )
    return output


def compose_image_caption_sample() -> Path:
    caption_template = ROOT / "resources" / "fragments" / "caption_fragment.docx"
    document = _caption_test_document(caption_template)
    caption_element = next(element for element in clone_component_elements(document, caption_template) if element.tag == qn("w:p"))
    document.element.body.sectPr.addprevious(caption_element)
    caption = Paragraph(caption_element, document._body)
    replace_component_token_runs(
        [document.element.body],
        "caption-fragment",
        [Run(text="Figure 1. Sample imported image with a styled caption.")],
    )

    image = Image.new("RGB", (1200, 675), "white")
    drawing = ImageDraw.Draw(image)
    drawing.rectangle((0, 0, 1200, 110), fill=(0, 176, 80))
    drawing.text((48, 42), "VulnReport image caption test", fill="white")
    drawing.rectangle((70, 180, 1130, 560), outline=(81, 81, 81), width=4)
    drawing.text((115, 250), "Sample evidence imported with python-docx", fill=(40, 40, 40))
    drawing.text((115, 315), "The styled caption appears directly below this image.", fill=(40, 40, 40))
    image_stream = BytesIO()
    image.save(image_stream, format="PNG")
    image_stream.seek(0)

    image_paragraph = document.add_paragraph()
    document.element.body.remove(image_paragraph._p)
    caption._p.addprevious(image_paragraph._p)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(image_stream, width=Mm(155))

    output_stream = BytesIO()
    document.save(output_stream)
    output = ROOT / "generated" / "image-caption-test.docx"
    atomic_write_bytes(output, output_stream.getvalue())
    return output


def compose_native_image_caption_sample() -> Path:
    caption_template = ROOT / "resources" / "fragments" / "caption_fragment.docx"
    document = _caption_test_document(caption_template)
    captions = (
        "Authentication response captured during testing.",
        "Authorization response for the modified account request.",
        "Verification response after remediation was applied.",
    )
    colors = ((0, 176, 80), (237, 125, 49), (81, 81, 81))
    for index, (caption_text, color) in enumerate(zip(captions, colors), 1):
        image = Image.new("RGB", (1200, 675), "white")
        drawing = ImageDraw.Draw(image)
        drawing.rectangle((0, 0, 1200, 110), fill=color)
        drawing.text((48, 42), f"Native Word caption test image {index}", fill="white")
        drawing.rectangle((70, 180, 1130, 560), outline=(81, 81, 81), width=4)
        drawing.text((115, 250), f"Sample evidence image {index} of 3", fill=(40, 40, 40))
        drawing.text((115, 315), caption_text, fill=(40, 40, 40))
        image_stream = BytesIO()
        image.save(image_stream, format="PNG")
        image_stream.seek(0)

        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        image_paragraph.add_run().add_picture(image_stream, width=Mm(155))
        caption_element = next(
            element
            for element in clone_component_elements(document, caption_template)
            if element.tag == qn("w:p")
        )
        replace_component_token_runs(
            [caption_element],
            "caption-fragment",
            [Run(text=caption_text)],
        )
        document.element.body.sectPr.addprevious(caption_element)

    source_stream = BytesIO()
    document.save(source_stream)
    source_path = ROOT / "generated" / "native-image-caption-source.docx"
    atomic_write_bytes(source_path, source_stream.getvalue())
    output = ROOT / "generated" / "native-image-caption-test.docx"
    postprocess_image_captions(source_path, output)
    return output


def _caption_test_document(caption_template: Path):
    source = Document(caption_template)
    document = Document()
    target_style_ids = {style.style_id for style in document.styles}
    for style in source.styles.element.iterchildren(qn("w:style")):
        style_id = style.get(qn("w:styleId"))
        if style_id in {"FiguresandTables", "FiguresandTablesChar"} and style_id not in target_style_ids:
            document.styles.element.append(deepcopy(style))
    for paragraph in list(document.paragraphs):
        document.element.body.remove(paragraph._p)
    return document


def main() -> None:
    for output in (
        compose_severity_sample(),
        compose_retest_finding_sample(),
        compose_image_caption_sample(),
        compose_native_image_caption_sample(),
    ):
        print(output)


if __name__ == "__main__":
    main()