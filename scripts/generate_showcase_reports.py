from __future__ import annotations

import argparse
import hashlib
import textwrap
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

from app.docx_captions import update_docx_bytes_with_word
from app.docx_report import generation_issues, render_report_docx
from app.library import Library
from app.models import (
    CodeFragment,
    Content,
    Engagement,
    EvidenceItem,
    ImageFragment,
    LibraryRef,
    ListFragment,
    ListItem,
    NoteFragment,
    ParagraphFragment,
    Report,
    Run,
    Scope,
    ScopeTarget,
    TableFragment,
    TestAccount,
    TestWindow,
    Vulnerability,
)
from app.report_service import affected_environments, assign_fresh_fragment_ids, report_export_filename
from app.storage import atomic_write_bytes
from app.workspace import Workspace


ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "resources" / "MAIN_TEST.docx"
LIBRARY = Library(ROOT / "resources" / "vuln_library.json")
MISSING_HEADER_LIBRARY_IDS = {f"VDB-{source_id:03d}" for source_id in range(36, 45)}
EXPECTED_FRAGMENT_TYPES = {
    "paragraph",
    "numbered_list",
    "bulleted_list",
    "table",
    "note",
    "image",
    "code_block",
}


@dataclass(frozen=True)
class FindingSpec:
    library_id: str
    header_name: str
    channels: tuple[str, ...]
    endpoint: str


@dataclass(frozen=True)
class ReportSpec:
    report_id: str
    app_name: str
    ci_number: str
    segment: str
    report_type: str
    non_production_label: str
    targets: dict[str, tuple[str, str]]
    findings: tuple[FindingSpec, ...]


REPORT_SPECS = (
    ReportSpec(
        report_id="r_showcase_atlas",
        app_name="Atlas Customer Portal",
        ci_number="CI-ATLAS-101",
        segment="JH",
        report_type="annual_pentest",
        non_production_label="UAT",
        targets={
            "web": ("https://portal.atlas.example.test", "https://uat.atlas.example.test"),
            "api": ("https://api.atlas.example.test", "https://api-uat.atlas.example.test"),
        },
        findings=(
            FindingSpec(
                library_id="VDB-036",
                header_name="Content-Security-Policy",
                channels=("api",),
                endpoint="GET /dashboard",
            ),
            FindingSpec(
                library_id="VDB-037",
                header_name="Cross-Origin-Embedder-Policy",
                channels=("web",),
                endpoint="GET /directory",
            ),
            FindingSpec(
                library_id="VDB-038",
                header_name="Cross-Origin-Resource-Policy",
                channels=("web",),
                endpoint="GET /assets/application.js",
            ),
        ),
    ),
    ReportSpec(
        report_id="r_showcase_meridian",
        app_name="Meridian Commerce API",
        ci_number="CI-MERIDIAN-202",
        segment="GWAM",
        report_type="deployment_pentest",
        non_production_label="TEST/MO",
        targets={
            "api": ("https://commerce-api.meridian.example.test", "https://commerce-api-test.meridian.example.test"),
        },
        findings=(
            FindingSpec(
                library_id="VDB-039",
                header_name="Permissions-Policy",
                channels=("api",),
                endpoint="GET /v2/orders",
            ),
            FindingSpec(
                library_id="VDB-040",
                header_name="Referrer-Policy",
                channels=("api",),
                endpoint="GET /v2/customers/summary?view=compact",
            ),
            FindingSpec(
                library_id="VDB-041",
                header_name="Strict-Transport-Security",
                channels=("api",),
                endpoint="GET /v2/session",
            ),
        ),
    ),
    ReportSpec(
        report_id="r_showcase_horizon",
        app_name="Horizon Workforce Suite",
        ci_number="CI-HORIZON-303",
        segment="Asia",
        report_type="retest",
        non_production_label="DEV",
        targets={
            "web": ("https://workforce.horizon.example.test", "https://workforce-dev.horizon.example.test"),
            "api": ("https://workforce-api.horizon.example.test", "https://workforce-api-dev.horizon.example.test"),
            "mobile": ("Horizon Mobile iOS and Android", "Horizon Mobile DEV build"),
        },
        findings=(
            FindingSpec(
                library_id="VDB-042",
                header_name="X-Content-Type-Options",
                channels=("api",),
                endpoint="GET /v1/timecards/42",
            ),
            FindingSpec(
                library_id="VDB-043",
                header_name="X-Frame-Options",
                channels=("web",),
                endpoint="GET /manager/dashboard",
            ),
            FindingSpec(
                library_id="VDB-044",
                header_name="X-Permitted-Cross-Domain-Policies",
                channels=("api",),
                endpoint="GET /documents/handbook.pdf",
            ),
        ),
    ),
)


def _font(size: int, *, bold: bool = False):
    filename = "seguisb.ttf" if bold else "segoeui.ttf"
    try:
        return ImageFont.truetype(str(Path("C:/Windows/Fonts") / filename), size)
    except OSError:
        return ImageFont.load_default()


def _draw_wrapped(
    draw: ImageDraw.ImageDraw,
    position: tuple[int, int],
    text: str,
    *,
    width: int,
    font,
    fill: str,
    spacing: int = 8,
) -> int:
    lines = []
    for source_line in text.splitlines() or [""]:
        lines.extend(textwrap.wrap(source_line, width=width, replace_whitespace=False) or [""])
    draw.multiline_text(position, "\n".join(lines), font=font, fill=fill, spacing=spacing)
    line_height = draw.textbbox((0, 0), "Ag", font=font)[3] + spacing
    return position[1] + line_height * len(lines)


def _evidence_png(
    spec: ReportSpec,
    title: str,
    endpoint: str,
    request: str,
    observation: str,
    environment: str,
    number: int,
) -> bytes:
    image = Image.new("RGB", (1400, 820), "#f4f7f8")
    draw = ImageDraw.Draw(image)
    accent = "#16705a" if environment == "production" else "#b35a18"
    environment_name = "PRODUCTION" if environment == "production" else spec.non_production_label

    draw.rectangle((0, 0, 1400, 94), fill="#17252b")
    draw.text((48, 25), "SYNTHETIC SECURITY EVIDENCE", font=_font(32, bold=True), fill="white")
    draw.rounded_rectangle((1110, 22, 1352, 72), radius=8, fill=accent)
    draw.text((1134, 33), environment_name, font=_font(21, bold=True), fill="white")

    draw.text((48, 126), f"Finding {number:03d}", font=_font(20, bold=True), fill=accent)
    title_bottom = _draw_wrapped(draw, (48, 158), title, width=68, font=_font(30, bold=True), fill="#17252b")
    app_name_y = title_bottom + 3
    draw.text((48, app_name_y), spec.app_name, font=_font(22), fill="#52646c")

    endpoint_top = max(276, app_name_y + 46)
    draw.rounded_rectangle((48, endpoint_top, 1352, endpoint_top + 68), radius=8, fill="white", outline="#cad5d9", width=2)
    draw.ellipse((72, endpoint_top + 21, 88, endpoint_top + 37), fill=accent)
    draw.text((110, endpoint_top + 17), endpoint, font=_font(22), fill="#20343d")

    panel_top = endpoint_top + 106
    panels = (
        (48, panel_top, 674, 722, "SYNTHETIC REQUEST", request),
        (726, panel_top, 1352, 722, "OBSERVED RESULT", observation),
    )
    for left, top, right, bottom, heading, body in panels:
        draw.rounded_rectangle((left, top, right, bottom), radius=8, fill="white", outline="#cad5d9", width=2)
        draw.rectangle((left, top, right, top + 58), fill="#e4ebed")
        draw.text((left + 24, top + 16), heading, font=_font(19, bold=True), fill="#20343d")
        _draw_wrapped(
            draw,
            (left + 24, top + 86),
            body,
            width=43,
            font=_font(23),
            fill="#20343d",
            spacing=12,
        )

    draw.text(
        (48, 770),
        "Demonstration image generated locally. No live systems or customer data are shown.",
        font=_font(18),
        fill="#65777f",
    )
    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def _item(text: str, *, bold_prefix: str | None = None) -> ListItem:
    if bold_prefix and text.startswith(bold_prefix):
        return ListItem(runs=[Run(text=bold_prefix, bold=True), Run(text=text[len(bold_prefix):])])
    return ListItem(runs=[Run(text=text)])


def _library_entry(finding: FindingSpec) -> dict:
    entry = LIBRARY.get(finding.library_id)
    if entry is None:
        raise RuntimeError(f"Missing vulnerability library entry: {finding.library_id}")
    if finding.library_id not in MISSING_HEADER_LIBRARY_IDS or not entry["title"].startswith("Missing/Misconfigured Security Header:"):
        raise RuntimeError(f"Not an approved missing-header entry: {finding.library_id}")
    return entry


def _finding_contents(
    report_spec: ReportSpec,
    finding: FindingSpec,
    entry: dict,
    report_number: int,
    finding_number: int,
) -> tuple[list[Content], dict[str, EvidenceItem], dict[str, bytes]]:
    base = f"r{report_number}_v{finding_number}"
    now = datetime.now().astimezone()
    channel = finding.channels[0]
    production_target = report_spec.targets[channel][0]
    host = urlsplit(production_target).netloc or production_target
    request = f"{finding.endpoint} HTTP/1.1\nHost: {host}\nAccept: text/html, application/json\nX-Test-Case: missing-security-header"
    observation = f"HTTP 200 response received; the {finding.header_name} header is absent."
    evidence: dict[str, EvidenceItem] = {}
    evidence_files: dict[str, bytes] = {}
    images = []
    for environment, short_name in (("production", "prod"), ("non_production", "nonprod")):
        evidence_id = f"ev_{base}_{short_name}"
        contents = _evidence_png(
            report_spec,
            entry["title"],
            finding.endpoint,
            request,
            observation,
            environment,
            finding_number,
        )
        evidence[evidence_id] = EvidenceItem(
            file=f"evidence/{evidence_id}.png",
            original_name=f"{base}-{short_name}.png",
            width_px=1400,
            height_px=820,
            sha256=hashlib.sha256(contents).hexdigest(),
            uploaded_at=now,
        )
        evidence_files[evidence_id] = contents
        label = "Production" if environment == "production" else report_spec.non_production_label
        images.append(ImageFragment(
            frag_id=f"f_{base}_{short_name}_image",
            type="image",
            environment=environment,
            evidence_id=evidence_id,
            caption=f"{label} response without the {finding.header_name} header",
            width_mm=145,
        ))

    library_contents = [Content.model_validate(content) for content in entry["contents"]]
    remediation = next(content for content in library_contents if content.type == "recommended_remediation")
    remediation.fragments.append(ListFragment(
        frag_id=f"f_{base}_remediation_checks",
        type="bulleted_list",
        items=[
            _item("Apply the header consistently to Production and Non-Production responses."),
            _item("Add an automated regression check for the expected header and value."),
        ],
    ))

    variant = entry["proof_of_concept"].get(channel)
    if not variant:
        raise RuntimeError(f"{finding.library_id} has no {channel} proof-of-concept variant")
    library_steps = Content(type="proof_of_concept", fragments=variant).fragments
    for fragment in library_steps:
        if isinstance(fragment, ListFragment):
            for item in fragment.items:
                for run in item.runs:
                    run.text = run.text.replace("Missing/Misconfigured/Present", "missing")
    proof = Content(
        type="proof_of_concept",
        fragments=[
            *library_steps,
            CodeFragment(
                frag_id=f"f_{base}_code",
                type="code_block",
                caption="Request used to inspect the response headers",
                text=request,
            ),
            TableFragment(
                frag_id=f"f_{base}_table",
                type="table", caption="Observed header comparison",
                header=[_item("Environment"), _item("Observed header state")],
                rows=[
                    [_item("Production"), _item(f"{finding.header_name}: missing")],
                    [_item(report_spec.non_production_label), _item(f"{finding.header_name}: missing")],
                ],
            ),
            ListFragment(
                frag_id=f"f_{base}_continued_steps",
                type="numbered_list",
                continue_numbering=True,
                items=[
                    _item(f"Confirm that {finding.header_name} is absent from the complete Production response."),
                    _item(f"Repeat the same request against {report_spec.non_production_label} and confirm the header is also absent."),
                    _item("Compare the observed behavior with the recommended header values documented above."),
                    _item("Save the complete responses and capture the environment-specific evidence shown below."),
                ],
            ),
            *images,
        ],
    )
    return [*library_contents, proof], evidence, evidence_files


def _build_report(spec: ReportSpec, report_number: int) -> tuple[Report, dict[str, bytes]]:
    now = datetime.now().astimezone()
    report_day = date.today()
    channels = list(spec.targets)
    scope_targets = []
    target_ids: dict[tuple[str, str], str] = {}
    for channel_index, (channel, values) in enumerate(spec.targets.items()):
        for environment, value in zip(("production", "non_production"), values):
            short_environment = "prod" if environment == "production" else "nonprod"
            target_id = f"tgt_r{report_number}_{channel}_{short_environment}"
            target_ids[(channel, environment)] = target_id
            scope_targets.append(ScopeTarget(
                target_id=target_id,
                environment=environment,
                channel=channel,
                value=value,
                order=channel_index,
            ))

    report = Report(
        report_id=spec.report_id,
        app_id=spec.ci_number,
        saved_at=now,
        engagement=Engagement(
            app_name=spec.app_name,
            ci_number=spec.ci_number,
            bsn_number=f"BSN-{report_number:03d}",
            app_owner="Security Engineering",
            segment=spec.segment,
            report_type=spec.report_type,
            start_date=report_day - timedelta(days=14),
            end_date=report_day - timedelta(days=7),
            tested_environments=["production", "non_production"],
            tested_channels=channels,
            non_production_label=spec.non_production_label,
            test_windows={
                "production": TestWindow(
                    start_date=report_day - timedelta(days=10),
                    end_date=report_day - timedelta(days=7),
                    test_time="22:00 EST",
                ),
                "non_production": TestWindow(
                    start_date=report_day - timedelta(days=14),
                    end_date=report_day - timedelta(days=11),
                    test_time="Anytime",
                ),
            },
            test_accounts=[
                TestAccount(user_role="Standard User", username="showcase.user"),
                TestAccount(user_role="Administrator", username="showcase.admin"),
            ],
            limitations="No destructive actions. Testing used synthetic accounts and records.",
            tester="Alex Morgan",
            report_date=report_day,
            classification="Confidential",
        ),
        scope_targets=scope_targets,
    )

    evidence_files: dict[str, bytes] = {}
    for finding_number, finding_spec in enumerate(spec.findings, start=1):
        if len(finding_spec.channels) != 1:
            raise RuntimeError("Each missing-header showcase finding must select one library PoC variant")
        entry = _library_entry(finding_spec)
        contents, evidence, files = _finding_contents(spec, finding_spec, entry, report_number, finding_number)
        selected_targets = [
            target_ids[(channel, environment)]
            for channel in finding_spec.channels
            for environment in ("production", "non_production")
        ]
        vulnerability = Vulnerability(
            uid=f"v_r{report_number}_{finding_number}",
            display_id=f"{finding_number:03d}",
            title=entry["title"],
            likelihood=entry["default_likelihood"],
            impact=entry["default_impact"],
            severity=entry["default_severity"],
            status="open_new",
            scope=Scope(mode="custom", target_ids=selected_targets),
            library_ref=LibraryRef(
                library_id=entry["library_id"],
                source_id=entry["source_id"],
                inserted_at=datetime.now().astimezone(),
            ),
            poc_variants=list(finding_spec.channels),
            contents=contents,
        )
        assign_fresh_fragment_ids(vulnerability)
        report.vulnerabilities.append(vulnerability)
        report.evidence.update(evidence)
        evidence_files.update(files)

    return Report.model_validate(report.model_dump(mode="json", by_alias=True)), evidence_files


def _verify_report(report: Report, spec: ReportSpec) -> None:
    issues = generation_issues(report)
    if issues:
        raise RuntimeError("Report is incomplete: " + "; ".join(issues))
    if len(report.vulnerabilities) != 3:
        raise RuntimeError("Each showcase report must contain exactly three findings")
    if set(report.engagement.tested_environments) != {"production", "non_production"}:
        raise RuntimeError("Each showcase report must cover Production and Non-Production")
    fragment_types = {
        fragment.type
        for finding in report.vulnerabilities
        for content in finding.contents
        for fragment in content.fragments
    }
    missing = EXPECTED_FRAGMENT_TYPES - fragment_types
    if missing:
        raise RuntimeError(f"Report does not exercise these fragment types: {', '.join(sorted(missing))}")
    if "instance_title" in fragment_types:
        raise RuntimeError("Showcase reports must not contain instance-title fragments")
    expected_library_ids = [finding.library_id for finding in spec.findings]
    actual_library_ids = [finding.library_ref.library_id if finding.library_ref else None for finding in report.vulnerabilities]
    if actual_library_ids != expected_library_ids:
        raise RuntimeError("Showcase findings do not match their vulnerability library entries")
    for finding in report.vulnerabilities:
        if set(affected_environments(finding, report)) != {"production", "non_production"}:
            raise RuntimeError(f"{finding.title} does not affect both environments")
        proof = next(content for content in finding.contents if content.type == "proof_of_concept")
        numbered_lists = [fragment for fragment in proof.fragments if isinstance(fragment, ListFragment) and fragment.type == "numbered_list"]
        if len(numbered_lists) != 2 or numbered_lists[0].continue_numbering or not numbered_lists[1].continue_numbering:
            raise RuntimeError(f"{finding.title} does not contain a continued second numbered-list fragment")
        if sum(len(fragment.items) for fragment in numbered_lists) < 7:
            raise RuntimeError(f"{finding.title} proof of concept is not long enough")


def _numbering_id(paragraph) -> str | None:
    properties = paragraph._p.find(qn("w:pPr"))
    numbering = properties.find(qn("w:numPr")) if properties is not None else None
    number = numbering.find(qn("w:numId")) if numbering is not None else None
    return number.get(qn("w:val")) if number is not None else None


def _verify_document(contents: bytes, report: Report) -> None:
    document = Document(BytesIO(contents))
    if len(document.inline_shapes) != 6:
        raise RuntimeError(f"Expected six evidence images in {report.engagement.app_name}")
    summary = next((table for table in document.tables if table.cell(0, 0).text == "Findings"), None)
    if summary is None or len(summary.rows) != 4:
        raise RuntimeError(f"Expected three summary findings in {report.engagement.app_name}")
    text = "\n".join([
        *(paragraph.text for paragraph in document.paragraphs),
        *(cell.text for table in document.tables for row in table.rows for cell in row.cells),
    ])
    if "{{" in text or "-fragments-here" in text.casefold():
        raise RuntimeError(f"Unresolved template token in {report.engagement.app_name}")
    for finding in report.vulnerabilities:
        if finding.title not in text:
            raise RuntimeError(f"Missing finding title in generated document: {finding.title}")
        proof = next(content for content in finding.contents if content.type == "proof_of_concept")
        numbered_lists = [fragment for fragment in proof.fragments if isinstance(fragment, ListFragment) and fragment.type == "numbered_list"]
        first_list_text = "".join(run.text for run in numbered_lists[0].items[-1].runs)
        continued_list_text = "".join(run.text for run in numbered_lists[1].items[0].runs)
        first_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text == first_list_text]
        continued_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text == continued_list_text]
        if len(first_paragraphs) != 1 or len(continued_paragraphs) != 1:
            raise RuntimeError(f"Could not locate both numbered-list fragments for {finding.title}")
        first_numbering = _numbering_id(first_paragraphs[0])
        continued_numbering = _numbering_id(continued_paragraphs[0])
        if first_numbering is None or first_numbering != continued_numbering:
            raise RuntimeError(f"Numbering does not continue across list fragments for {finding.title}")


def main() -> int:
    parser = argparse.ArgumentParser(description="Create three complete showcase reports and DOCX files")
    parser.add_argument("--output-dir", type=Path, default=ROOT / "generated")
    parser.add_argument("--skip-word", action="store_true", help="Skip Microsoft Word field and TOC refresh")
    arguments = parser.parse_args()

    workspace = Workspace(ROOT / "data", "Alex Morgan")
    arguments.output_dir.mkdir(parents=True, exist_ok=True)
    generated = []
    used_library_ids = set()

    for report_number, spec in enumerate(REPORT_SPECS, start=1):
        report, evidence_files = _build_report(spec, report_number)
        _verify_report(report, spec)
        used_library_ids.update(finding.library_ref.library_id for finding in report.vulnerabilities if finding.library_ref)

        draft_path = workspace.save(report)
        for evidence_id, contents in evidence_files.items():
            atomic_write_bytes(draft_path.parent / "evidence" / f"{evidence_id}.png", contents)

        output_path = arguments.output_dir / report_export_filename(report, ".docx")
        output_path.unlink(missing_ok=True)
        contents = render_report_docx(report, TEMPLATE, draft_path.parent)
        if not arguments.skip_word:
            contents = update_docx_bytes_with_word(contents)
        _verify_document(contents, report)
        atomic_write_bytes(output_path, contents)
        generated.append((report.report_id, output_path.resolve()))

    if used_library_ids != MISSING_HEADER_LIBRARY_IDS:
        raise RuntimeError("Showcase set does not use all nine missing-header library entries")

    for report_id, output_path in generated:
        print(f"{report_id}: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())