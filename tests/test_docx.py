from __future__ import annotations

import re
import tempfile
import threading
import unittest
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from typing import get_args
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image, ImageFilter, ImageOps

from app.docx_report import (
    LOCATION_WRAP_CHARACTERS,
    SCOPE_WRAP_CHARACTERS,
    TESTING_RESULT_PARAGRAPH,
    ReportGenerationError,
    _bordered_image_stream,
    _find_table,
    _metadata,
    _populate_scope_tables,
    _unresolved_placeholders,
    _wrap_long_value,
    generation_issues,
    main_template_path,
    render_report_docx,
)
from app.docx_captions import flatten_section_number_fields
from app.docx_components import replace_component_token_runs
from app.docx_import import NO_FINDINGS_TITLE
from app.report_service import provision, sync_evidence_image_slots
from tests.support import numbering_details, off_border_pixels, paragraph_texts, png_bytes
from app.models import CodeFragment, Content, Engagement, EvidenceItem, ImageFragment, InstanceTitleFragment, ListFragment, ListItem, NoteFragment, ParagraphFragment, Report, Run, Scope, ScopeTarget, Segment, TableFragment, TestAccount, TestWindow, Vulnerability


def _keeps_next(element) -> bool:
    properties = element.find(qn("w:pPr"))
    return properties is not None and properties.find(qn("w:keepNext")) is not None


def _is_blank_paragraph(element) -> bool:
    return element.tag == qn("w:p") and not "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()


def _footer_texts(document) -> dict[str, list[str]]:
    """Each footer part of a document by part name, with the text of its paragraphs."""
    return {
        str(part.partname): paragraph_texts(part._element)
        for part in document.part.package.parts
        if "/footer" in str(part.partname)
    }


MASTERS = ("MAIN.docx", "MAIN_ASIA.docx", "MAIN_THICK_MOBILE.docx", "MAIN_THICK_MOBILE_ASIA.docx")


RESOURCES = Path(__file__).resolve().parent.parent / "resources"


class DocxReportTests(unittest.TestCase):
    def test_renders_template_with_repeated_findings_and_no_placeholders(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            evidence_folder = report_folder / "evidence"
            evidence_folder.mkdir()
            evidence = BytesIO()
            Image.new("RGB", (40, 20), "white").save(evidence, format="PNG")
            for evidence_id in ("ev_prod", "ev_prod_extra", "ev_uat", "ev_prod_second", "ev_prev"):
                (evidence_folder / f"{evidence_id}.png").write_bytes(evidence.getvalue())

            now = datetime.now().astimezone()
            report = Report(
                report_id="r_docx_test",
                app_id="CI-DOCX",
                saved_at=now,
                engagement=Engagement(
                    app_name="Northstar Banking",
                    ci_number="CI-DOCX",
                    app_owner="Avery Morgan",
                    segment="JH",
                    report_type="annual_pentest",
                    report_date=date(2026, 9, 9),
                    tester="QA Tester",
                    network="Internal",
                    tested_environments=["production", "non_production"],
                    tested_channels=["web", "api"],
                    # Pinned rather than defaulted, because the assertion below counts "UAT:" headings.
                    non_production_label="UAT",
                    test_windows={
                        "production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2), test_time="22:00 EST"),
                        "non_production": TestWindow(start_date=date(2026, 7, 28), end_date=date(2026, 7, 30), test_time="Anytime"),
                    },
                    test_accounts=[TestAccount(user_role="Customer", username="qa.customer"), TestAccount(user_role="Admin", username="qa.admin")],
                    limitations="No destructive testing.",
                ),
                scope_targets=[
                    ScopeTarget(target_id="t_prod_web", environment="production", channel="web", value="https://prod.example.test"),
                    ScopeTarget(target_id="t_prod_web_2", environment="production", channel="web", value="https://admin.example.test", order=1),
                    ScopeTarget(target_id="t_prod_api", environment="production", channel="api", value="https://prod-api.example.test", order=1),
                    ScopeTarget(target_id="t_prod_api_2", environment="production", channel="api", value="https://prod-api.example.test/v2", order=2),
                    ScopeTarget(target_id="t_uat_web", environment="non_production", channel="web", value="https://uat.example.test"),
                    ScopeTarget(target_id="t_uat_web_2", environment="non_production", channel="web", value="https://staging.example.test", order=1),
                    ScopeTarget(target_id="t_uat_api", environment="non_production", channel="api", value="https://uat-api.example.test"),
                    ScopeTarget(target_id="t_uat_api_2", environment="non_production", channel="api", value="https://uat-api.example.test/v2", order=1),
                ],
                evidence={
                    evidence_id: EvidenceItem(file=f"evidence/{evidence_id}.png", original_name=f"{evidence_id}.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)
                    for evidence_id in ("ev_prod", "ev_prod_extra", "ev_uat", "ev_prod_second", "ev_prev")
                },
            )
            report.vulnerabilities = [
                self._finding("v_high_1", "Authorization bypass", "high", "001", ["t_prod_web", "t_prod_api", "t_uat_api"], [
                    ImageFragment(frag_id="f_img_prod", type="image", environment="production", evidence_id="ev_prod", caption="Production response"),
                    ImageFragment(frag_id="f_img_prod_extra", type="image", environment="production", evidence_id="ev_prod_extra", caption="Additional production response"),
                    ImageFragment(frag_id="f_img_uat", type="image", environment="non_production", evidence_id="ev_uat", caption="Non-Production response"),
                ], likelihood="critical", impact="low"),
                self._finding("v_high_2", "Session fixation", "high", "002", ["t_prod_web", "t_prod_api"], [
                    ImageFragment(frag_id="f_img_prod_2", type="image", environment="production", evidence_id="ev_prod_second", caption="Additional production response"),
                ], likelihood="medium", impact="informational"),
            ]
            report.vulnerabilities[1].status = "resolved"
            provision(report.vulnerabilities[1])
            # The app's own conclusion sentence now blocks generation, so this finding writes its own.
            conclusion = next(content for content in report.vulnerabilities[1].contents if content.type == "in_conclusion")
            conclusion.fragments[0].runs = [Run(text="Session fixation was remediated and could not be reproduced on retest.")]
            previous = next(content for content in report.vulnerabilities[1].contents if content.type == "previous_proof_of_concept")
            next(fragment for fragment in previous.fragments if isinstance(fragment, ListFragment)).items = [
                ListItem(runs=[Run(text="Reproduce the previously reported behavior.")]),
                ListItem(runs=[Run(text="Record the original vulnerable response.")]),
            ]
            # Non-production is outside this finding's scope, so this also pins that history is not
            # filtered by the current engagement.
            previous_image = next(fragment for fragment in previous.fragments if isinstance(fragment, ImageFragment))
            previous_image.environment = "non_production"
            previous_image.evidence_id = "ev_prev"
            previous_image.caption = "Original non-production response"

            template = RESOURCES / "MAIN.docx"
            generated = render_report_docx(report, template, report_folder)
            rendered = Document(BytesIO(generated))
            text = "\n".join([*(paragraph.text for paragraph in rendered.paragraphs), *(cell.text for table in rendered.tables for row in table.rows for cell in row.cells)])
            self.assertNotIn("{{", text)
            self.assertNotIn("-fragments-here", text.casefold())
            self.assertEqual(text.count("Authorization bypass"), 2)
            self.assertEqual(text.count("Session fixation"), 3)
            summary = next(table for table in rendered.tables if table.cell(0, 0).text == "Findings")
            self.assertEqual(len(summary.rows), 3)
            expected_rating_colors = {
                "Authorization bypass": (RGBColor(192, 0, 0), RGBColor(112, 173, 71), RGBColor(237, 125, 49)),
                "Session fixation": (RGBColor(255, 192, 0), RGBColor(81, 81, 81), RGBColor(237, 125, 49)),
            }
            for row in summary.rows[1:]:
                self.assertEqual(
                    tuple(row.cells[index].paragraphs[0].runs[0].font.color.rgb for index in (1, 2, 3)),
                    expected_rating_colors[row.cells[0].text],
                )
                for index in (1, 2, 3):
                    rating_run = next(run for run in row.cells[index].paragraphs[0].runs if run.text)
                    self.assertEqual(rating_run.font.size, Pt(12))
            self.assertEqual(len(rendered.inline_shapes), 5)
            paragraph_texts = [paragraph.text for paragraph in rendered.paragraphs]
            self.assertEqual(paragraph_texts.count("PROD:"), 2)
            # The evidence label follows the engagement's chosen non-production name. The second
            # occurrence is the carried previous proof of concept, which the retest no longer covers.
            self.assertEqual(paragraph_texts.count("UAT:"), 2)
            self.assertEqual(paragraph_texts.count("NON-PROD:"), 0)

            template_document = Document(RESOURCES / "MAIN.docx")
            template_revision = next(
                table for table in template_document.tables if table.cell(0, 0).text == "Version"
            )
            rendered_revision = next(
                table for table in rendered.tables if table.cell(0, 0).text == "Version"
            )
            self.assertEqual(
                [child.tag for child in rendered_revision._tbl.iterchildren()],
                [child.tag for child in template_revision._tbl.iterchildren()],
            )
            self.assertEqual(
                sum(1 for _ in rendered_revision._tbl.iter(qn("w:tr"))),
                sum(1 for _ in template_revision._tbl.iter(qn("w:tr"))),
            )
            revision_text = " ".join(
                node.text or "" for node in rendered_revision._tbl.iter(qn("w:t"))
            )
            self.assertIn("1.0", revision_text)
            self.assertIn("New Report", revision_text)
            self.assertIn("September 9, 2026", revision_text)
            self.assertIn("QA Tester", revision_text)
            self.assertNotIn("{{", revision_text)
            template_header_footer_parts = {
                str(part.partname): part
                for part in template_document.part.package.parts
                if "/header" in str(part.partname) or "/footer" in str(part.partname)
            }
            rendered_header_footer_parts = {
                str(part.partname): part
                for part in rendered.part.package.parts
                if "/header" in str(part.partname) or "/footer" in str(part.partname)
            }
            self.assertEqual(rendered_header_footer_parts.keys(), template_header_footer_parts.keys())
            tokenised_parts = 0
            for part_name, template_part in template_header_footer_parts.items():
                template_text = "".join(node.text or "" for node in template_part._element.iter(qn("w:t")))
                if "{{" not in template_text:
                    continue
                tokenised_parts += 1
                rendered_text = "".join(
                    node.text or ""
                    for node in rendered_header_footer_parts[part_name]._element.iter(qn("w:t"))
                )
                self.assertNotIn("{{", rendered_text)
                self.assertIn("Northstar Banking", rendered_text)
            self.assertGreater(tokenised_parts, 0, "no header or footer carried a token, so nothing was checked")
            page_number_starts = [
                page_numbering.get(qn("w:start")) if page_numbering is not None else None
                for section_properties in rendered.element.body.iter(qn("w:sectPr"))
                if not any(
                    ancestor.tag == qn("w:sectPrChange")
                    for ancestor in section_properties.iterancestors()
                )
                for page_numbering in [section_properties.find(qn("w:pgNumType"))]
            ]
            self.assertEqual(
                page_number_starts,
                [None, "1", "1", None, None, None, None, None, None, None],
            )
            self.assertEqual(sum(paragraph.text == "High Findings" for paragraph in rendered.paragraphs), 1)
            self.assertEqual(sum(paragraph.text == "Previous Proof of Concept: " for paragraph in rendered.paragraphs), 1)
            self.assertEqual(sum(paragraph.text == "In Conclusion: " for paragraph in rendered.paragraphs), 1)
            self._assert_image_fragment_format(rendered)
            web_scope = next(table for table in rendered.tables if table.cell(0, 0).text == "URL(s) in Scope")
            api_scope = next(table for table in rendered.tables if table.cell(0, 0).text == "API Routes")
            expected_scope_lines = [
                (web_scope.cell(2, 0), ["https://prod.example.test", "https://admin.example.test"]),
                (web_scope.cell(4, 0), ["https://uat.example.test", "https://staging.example.test"]),
                (api_scope.cell(2, 0), ["https://prod-api.example.test", "https://prod-api.example.test/v2"]),
                (api_scope.cell(4, 0), ["https://uat-api.example.test", "https://uat-api.example.test/v2"]),
            ]
            for cell, expected_lines in expected_scope_lines:
                self.assertEqual([paragraph.text for paragraph in cell.paragraphs], expected_lines)
                self.assertTrue(all(paragraph.style.name == "List Bullet 2" for paragraph in cell.paragraphs))
            for vulnerability_id in ("001", "002"):
                details = next(
                    table
                    for table in rendered.tables
                    if len(table.rows) >= 5
                    and table.cell(0, 0).text == "Severity"
                    and table.cell(1, 1).text == vulnerability_id
                )
                production_location = details.cell(3, 1)
                self.assertEqual(
                    production_location.text,
                    "Production Environment:\nhttps://prod.example.test\nhttps://prod-api.example.test",
                )
                # The glyph comes from the list style, so short locations need no manual break.
                self.assertEqual(len(production_location._tc.findall(".//" + qn("w:br"))), 0)
                bullets = [item for item in production_location.paragraphs if item.style.name == "List Paragraph"]
                self.assertEqual(len(bullets), 2)
                self.assertTrue(all(item._p.find(qn("w:pPr")).find(qn("w:numPr")) is not None for item in bullets))
            section_titles = {
                "Recommended Remediation:",
                "Proof of Concept:",
                "In Conclusion:",
                "Severity Review Ticket (if applicable):",
            }
            titles_seen = 0
            for paragraph in rendered.paragraphs:
                if paragraph.text.strip() in section_titles:
                    titles_seen += 1
                    previous = paragraph._p.getprevious()
                    blank_count = 0
                    while (
                        previous is not None
                        and previous.tag == qn("w:p")
                        and not "".join(node.text or "" for node in previous.iter(qn("w:t"))).strip()
                    ):
                        blank_count += 1
                        previous = previous.getprevious()
                    self.assertLessEqual(blank_count, 1)
            self.assertGreaterEqual(titles_seen, 3, "the section titles were not found, so the spacing went unchecked")
            description = next(paragraph for paragraph in rendered.paragraphs if paragraph.text == "A complete description.")
            following = description._p.getnext()
            self.assertEqual(following.tag, qn("w:p"))
            self.assertEqual("".join(node.text or "" for node in following.iter(qn("w:t"))), "")
            self.assertTrue(any(paragraph.text == "Request" and paragraph.style.name == "Figures and Tables" for paragraph in rendered.paragraphs))
            image_caption = next(paragraph for paragraph in rendered.paragraphs if paragraph.text.endswith("Production response"))
            self.assertEqual(image_caption.style.name, "Normal")
            self.assertEqual(image_caption._p.find(qn("w:pPr")).find(qn("w:jc")).get(qn("w:val")), "center")
            self.assertTrue(image_caption._p.getprevious().xpath(".//w:drawing"))
            self.assertEqual(image_caption.text, "Figure 2. Production response")
            self.assertEqual(
                [node.text for node in image_caption._p.iter(qn("w:instrText"))],
                [r" SEQ Figure \* ARABIC "],
            )
            component_code = next(paragraph for paragraph in rendered.paragraphs if paragraph.text == "GET /accounts/123")
            self.assertEqual(component_code.runs[0].font.name, "Consolas")
            note = next(paragraph for paragraph in rendered.paragraphs if paragraph.text == "Note: Validate the result independently.")
            self.assertTrue(all(run.italic for run in note.runs if run.text))
            generated_table = next(table for table in rendered.tables if table.cell(0, 0).text == "Generated Header")
            self.assertEqual(sum(int(column.get(qn("w:w"))) for column in generated_table._tbl.tblGrid.iterchildren(qn("w:gridCol"))), 9994)
            with ZipFile(BytesIO(generated)) as archive:
                self.assertIn("word/document.xml", archive.namelist())

    def test_template_without_the_findings_anchor_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            report = Report(
                report_id="r_legacy_template",
                app_id="CI-DOCX",
                saved_at=datetime.now().astimezone(),
                engagement=Engagement(app_name="Northstar Banking"),
            )
            template = RESOURCES / "fixtures" / "report-name.docx"
            with self.assertRaises(ReportGenerationError) as raised:
                render_report_docx(report, template, report_folder, allow_incomplete=True)
            self.assertIn("{{findings}}", str(raised.exception))

    def _supporting_image_report(self, report_folder: Path, tested_environments: list[str], extra_image: ImageFragment) -> Report:
        """A production-only finding carrying one more image, labelled for the other environment."""
        evidence_folder = report_folder / "evidence"
        evidence_folder.mkdir()
        buffer = png_bytes(40, 20)
        for evidence_id in ("ev_prod", "ev_other"):
            (evidence_folder / f"{evidence_id}.png").write_bytes(buffer)
        now = datetime.now().astimezone()
        windows = {
            "production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2), test_time="22:00 EST"),
            "non_production": TestWindow(start_date=date(2026, 7, 28), end_date=date(2026, 7, 30), test_time="Anytime"),
        }
        targets = [ScopeTarget(target_id="t_prod", environment="production", channel="web", value="https://prod.example.test")]
        if "non_production" in tested_environments:
            targets.append(ScopeTarget(target_id="t_uat", environment="non_production", channel="web", value="https://uat.example.test"))
        report = Report(
            report_id="r_supporting_image",
            app_id="CI-DOCX",
            saved_at=now,
            engagement=Engagement(
                app_name="Northstar Banking",
                ci_number="CI-DOCX",
                app_owner="Avery Morgan",
                segment="JH",
                report_type="annual_pentest",
                report_date=date(2026, 9, 9),
                tester="QA Tester",
                network="Internal",
                tested_environments=tested_environments,
                tested_channels=["web"],
                non_production_label="UAT",
                test_windows={environment: windows[environment] for environment in tested_environments},
            ),
            scope_targets=targets,
            evidence={
                evidence_id: EvidenceItem(file=f"evidence/{evidence_id}.png", original_name=f"{evidence_id}.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)
                for evidence_id in ("ev_prod", "ev_other")
            },
        )
        report.vulnerabilities = [
            self._finding("v_prod_only", "Authorization bypass", "high", "001", ["t_prod"], [
                ImageFragment(frag_id="f_img_prod", type="image", environment="production", evidence_id="ev_prod", caption="Production response"),
                extra_image,
            ]),
        ]
        return report

    def test_a_supporting_image_prints_under_its_own_environment_heading(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            supporting = ImageFragment(frag_id="f_img_uat", type="image", environment="non_production", evidence_id="ev_other", caption="Same response in UAT")
            report = self._supporting_image_report(report_folder, ["production", "non_production"], supporting)
            self.assertEqual(generation_issues(report), [])

            rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))
            paragraph_texts = [paragraph.text for paragraph in rendered.paragraphs]
            self.assertEqual(paragraph_texts.count("PROD:"), 1)
            self.assertEqual(paragraph_texts.count("UAT:"), 1)
            self.assertEqual(len(rendered.inline_shapes), 2)
            self.assertIn("Same response in UAT", "\n".join(paragraph_texts))

    def test_an_untested_environment_image_is_left_out_of_the_report(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            untested = ImageFragment(frag_id="f_img_uat", type="image", environment="non_production", evidence_id="ev_other", caption="Response from an untested UAT")
            report = self._supporting_image_report(report_folder, ["production"], untested)
            self.assertEqual(generation_issues(report), [])

            rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))
            paragraph_texts = [paragraph.text for paragraph in rendered.paragraphs]
            self.assertEqual(paragraph_texts.count("PROD:"), 1)
            self.assertEqual(paragraph_texts.count("UAT:"), 0)
            self.assertEqual(len(rendered.inline_shapes), 1)
            self.assertNotIn("Response from an untested UAT", "\n".join(paragraph_texts))

    def test_a_started_supporting_image_must_be_complete_and_an_untouched_one_is_left_out(self) -> None:
        cases = [
            ("untouched", None, "", []),
            ("caption without an image", None, "Same response in UAT", ["Authorization bypass: image required for image fragment"]),
            ("image without a caption", "ev_other", " ", ["Authorization bypass: caption required for image fragment"]),
        ]
        for label, evidence_id, caption, expected in cases:
            with self.subTest(label), tempfile.TemporaryDirectory() as temporary_directory:
                report_folder = Path(temporary_directory)
                supporting = ImageFragment(frag_id="f_img_uat", type="image", environment="non_production", evidence_id=evidence_id, caption=caption)
                report = self._supporting_image_report(report_folder, ["production", "non_production"], supporting)
                self.assertEqual(generation_issues(report), expected)
                if not expected:
                    rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))
                    self.assertEqual([paragraph.text for paragraph in rendered.paragraphs].count("UAT:"), 0)
                    self.assertEqual(len(rendered.inline_shapes), 1)

    @staticmethod
    def _image(frag_id: str, environment: str, caption: str) -> ImageFragment:
        evidence_id = "ev_prod" if environment == "production" else "ev_other"
        return ImageFragment(frag_id=frag_id, type="image", environment=environment, evidence_id=evidence_id, caption=caption)

    @staticmethod
    def _label_sequence(rendered, captions: list[str]) -> list[str]:
        """Environment labels, instance titles and the named captions, in document order."""
        sequence = []
        for text in (paragraph.text.strip() for paragraph in rendered.paragraphs):
            if text in ("PROD:", "UAT:") or text.startswith("Instance "):
                sequence.append(text)
            else:
                sequence.extend(caption for caption in captions if caption in text)
        return sequence

    def test_an_environment_label_prints_whenever_the_environment_changes(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            report = self._supporting_image_report(report_folder, ["production", "non_production"], self._image("f_unused", "non_production", "Unused"))
            images = [
                self._image("f_p1", "production", "Screenshot P1"),
                self._image("f_u1", "non_production", "Screenshot U1"),
                self._image("f_p2", "production", "Screenshot P2"),
                self._image("f_p3", "production", "Screenshot P3"),
                self._image("f_u2", "non_production", "Screenshot U2"),
            ]
            report.vulnerabilities = [self._finding("v_labels", "Authorization bypass", "high", "001", ["t_prod"], images)]
            self.assertEqual(generation_issues(report), [])

            rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))
            self.assertEqual(self._label_sequence(rendered, [image.caption for image in images]), [
                "PROD:", "Screenshot P1", "UAT:", "Screenshot U1", "PROD:", "Screenshot P2", "Screenshot P3", "UAT:", "Screenshot U2",
            ])

    def test_an_instance_title_brings_back_the_environment_label_and_a_note_does_not(self) -> None:
        def sequence() -> list:
            return [
                self._image("f_p1", "production", "Screenshot P1"),
                NoteFragment(frag_id="f_between", type="note", runs=[Run(text="Then refresh the page.")]),
                self._image("f_p2", "production", "Screenshot P2"),
                InstanceTitleFragment(frag_id="f_second", type="instance_title", text="Second account"),
                self._image("f_p3", "production", "Screenshot P3"),
            ]

        captions = ["Screenshot P1", "Screenshot P2", "Screenshot P3", "Current response"]
        expected = ["PROD:", "Screenshot P1", "Screenshot P2", "Instance 1: Second account", "PROD:", "Screenshot P3"]
        for section in ("proof_of_concept", "previous_proof_of_concept"):
            with self.subTest(section), tempfile.TemporaryDirectory() as temporary_directory:
                report_folder = Path(temporary_directory)
                if section == "proof_of_concept":
                    report = self._supporting_image_report(report_folder, ["production", "non_production"], self._image("f_unused", "non_production", "Unused"))
                    report.vulnerabilities = [self._finding("v_labels", "Authorization bypass", "high", "001", ["t_prod"], sequence())]
                else:
                    report = self._retest_report(report_folder, ("ev_prod",))
                    finding = self._finding("v_labels", "Authorization bypass", "high", "001", ["t_prod"], [self._image("f_current", "production", "Current response")])
                    report.vulnerabilities = [finding]
                    self._carry_history(finding, "production", "ev_prod")
                    next(content for content in finding.contents if content.type == "previous_proof_of_concept").fragments = sequence()
                self.assertEqual(generation_issues(report), [])

                rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))
                labels = self._label_sequence(rendered, captions)
                if section == "previous_proof_of_concept":
                    # This finding's own proof of concept prints too; only the carried section is under test.
                    current = next(index for index in range(len(labels)) if labels[index:index + 2] == ["PROD:", "Current response"])
                    del labels[current:current + 2]
                self.assertEqual(labels, expected)

    def _retest_report(self, report_folder: Path, evidence_ids: tuple[str, ...]) -> Report:
        """A production-only finding in an engagement that also tested non-production."""
        evidence_folder = report_folder / "evidence"
        evidence_folder.mkdir()
        buffer = png_bytes(40, 20)
        for evidence_id in evidence_ids:
            (evidence_folder / f"{evidence_id}.png").write_bytes(buffer)
        now = datetime.now().astimezone()
        return Report(
            report_id="r_retest",
            app_id="CI-DOCX",
            saved_at=now,
            engagement=Engagement(
                app_name="Northstar Banking",
                ci_number="CI-DOCX",
                app_owner="Avery Morgan",
                segment="JH",
                report_type="retest",
                report_date=date(2026, 9, 9),
                tester="QA Tester",
                network="Internal",
                tested_environments=["production"],
                tested_channels=["web"],
                # Pinned rather than defaulted, because the caller counts "UAT:" headings.
                non_production_label="UAT",
                test_windows={"production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2), test_time="22:00 EST")},
            ),
            scope_targets=[ScopeTarget(target_id="t_prod", environment="production", channel="web", value="https://prod.example.test")],
            evidence={
                evidence_id: EvidenceItem(file=f"evidence/{evidence_id}.png", original_name=f"{evidence_id}.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)
                for evidence_id in evidence_ids
            },
        )

    @staticmethod
    def _carry_history(finding: Vulnerability, environment: str, evidence_id: str) -> ImageFragment:
        """Promote a finding to previously discovered and fill the seeded previous proof of concept."""
        finding.status = "open_previously_discovered"
        provision(finding)
        # Promoting the status seeds a conclusion, and the app's own sentence blocks generation.
        conclusion = next(content for content in finding.contents if content.type == "in_conclusion")
        conclusion.fragments[0].runs = [Run(text="The original finding remains exploitable on retest.")]
        previous = next(content for content in finding.contents if content.type == "previous_proof_of_concept")
        next(fragment for fragment in previous.fragments if isinstance(fragment, ListFragment)).items = [ListItem(runs=[Run(text="Reproduce the original finding.")])]
        history = next(fragment for fragment in previous.fragments if isinstance(fragment, ImageFragment))
        history.environment = environment
        history.evidence_id = evidence_id
        history.caption = "Original non-production response"
        return history

    def test_previous_proof_of_concept_image_survives_a_narrower_retest(self) -> None:
        """A retest is usually narrower than the test before it, so gating history on the current
        scope would hide, relabel, or drop evidence the previous engagement actually gathered."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            report = self._retest_report(report_folder, ("ev_prod", "ev_history"))
            report.vulnerabilities = [
                self._finding("v_retest", "Authorization bypass", "high", "001", ["t_prod"], [
                    ImageFragment(frag_id="f_img_prod", type="image", environment="production", evidence_id="ev_prod", caption="Production response"),
                ]),
            ]
            finding = report.vulnerabilities[0]
            history = self._carry_history(finding, "non_production", "ev_history")

            sync_evidence_image_slots(finding, report)
            self.assertEqual(history.environment, "non_production", "a carried image keeps the environment it was found in")
            self.assertEqual(generation_issues(report), [])

            template = RESOURCES / "MAIN.docx"
            rendered = Document(BytesIO(render_report_docx(report, template, report_folder)))
            paragraph_texts = [paragraph.text for paragraph in rendered.paragraphs]
            self.assertEqual(paragraph_texts.count("UAT:"), 1, "the carried non-production image must still be labelled and rendered")
            self.assertEqual(len(rendered.inline_shapes), 2)

    def test_a_carried_image_that_already_has_a_border_prints_without_a_second_one(self) -> None:
        """A picture carried over from an earlier report went through this renderer once already."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            report = self._retest_report(report_folder, ("ev_prod", "ev_history"))
            earlier = ImageOps.expand(Image.new("RGB", (1200, 800), "white"), border=2, fill="black")
            earlier.save(report_folder / "evidence" / "ev_history.png")
            report.vulnerabilities = [
                self._finding("v_retest", "Authorization bypass", "high", "001", ["t_prod"], [
                    ImageFragment(frag_id="f_img_prod", type="image", environment="production", evidence_id="ev_prod", caption="Production response"),
                ]),
            ]
            self._carry_history(report.vulnerabilities[0], "production", "ev_history")

            rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))

            printed = []
            for shape in rendered.inline_shapes:
                relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.get(qn("r:embed"))
                printed.append(Image.open(BytesIO(rendered.part.related_parts[relationship_id].blob)).convert("RGB"))
            carried = [image for image in printed if image.size == earlier.size]
            self.assertEqual(len(carried), 1, [image.size for image in printed])
            self.assertTrue(carried[0].tobytes() == earlier.tobytes(), "the carried picture was changed")
            current = [image for image in printed if image.size != earlier.size]
            self.assertEqual(len(current), 1)
            self.assertEqual(off_border_pixels(current[0]), [], "the new screenshot lost its border")

    def test_previous_evidence_does_not_satisfy_the_retest_requirement(self) -> None:
        """Last year's screenshot is not this year's proof, so it must not silently stand in for it."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            report = self._retest_report(report_folder, ("ev_history",))
            report.vulnerabilities = [self._finding("v_retest", "Authorization bypass", "high", "001", ["t_prod"], [])]
            self._carry_history(report.vulnerabilities[0], "production", "ev_history")

            self.assertIn("Authorization bypass: Production evidence image required", generation_issues(report))

    def test_affected_locations_are_real_bullets_and_long_paths_wrap_on_a_slash(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            (report_folder / "evidence").mkdir()
            buffer = png_bytes(40, 20)
            (report_folder / "evidence" / "ev_wrap.png").write_bytes(buffer)
            now = datetime.now().astimezone()
            web_target = "https://prod.example.test/accounts/123/details/extra/settings/preferences/alerts/email"
            api_target = "https://api.example.test/v2/customers/profile/settings/advanced/notifications/preferences"
            report = Report(
                report_id="r_wrap", app_id="CI-DOCX", saved_at=now,
                engagement=Engagement(
                    app_name="Northstar Banking", ci_number="CI-DOCX", app_owner="Avery Morgan", segment="JH", report_type="annual_pentest",
                    report_date=date(2026, 9, 9), tester="QA Tester", network="Internal", tested_environments=["production"], tested_channels=["web", "api"],
                    test_windows={"production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2))},
                ),
                scope_targets=[
                    ScopeTarget(target_id="t_web", environment="production", channel="web", value=web_target),
                    ScopeTarget(target_id="t_api", environment="production", channel="api", value=api_target),
                ],
                evidence={"ev_wrap": EvidenceItem(file="evidence/ev_wrap.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)},
            )
            report.vulnerabilities = [self._finding("v_wrap", "Authorization bypass", "high", "001", ["t_web", "t_api"], [
                ImageFragment(frag_id="f_img", type="image", environment="production", evidence_id="ev_wrap", caption="Production response"),
            ])]
            rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))

            locations = next(
                cell
                for table in rendered.tables if table.cell(0, 0).text not in {"URL(s) in Scope", "API Routes"}
                for row in table.rows for cell in row.cells
                if "prod.example.test/accounts" in cell.text
            )
            bullets = [paragraph for paragraph in locations.paragraphs if paragraph.style.name == "List Paragraph"]
            self.assertEqual(len(bullets), 2)
            for paragraph in bullets:
                properties = paragraph._p.find(qn("w:pPr"))
                self.assertIsNotNone(properties.find(qn("w:numPr")), "affected locations must be a real bullet list")
                self.assertEqual(properties.find(qn("w:jc")).get(qn("w:val")), "left")
                self.assertEqual(len(paragraph._p.findall(".//" + qn("w:br"))), 1, "a long location should break once")
            self.assertNotIn("\u2022", locations.text, "the bullet glyph must come from the list style, not the text")
            self.assertTrue(bullets[0].text.startswith("https://prod.example.test/accounts/"))
            self.assertTrue(all(len(line) <= LOCATION_WRAP_CHARACTERS for line in bullets[0].text.split("\n")))

            for name, target in (("URL(s) in Scope", web_target), ("API Routes", api_target)):
                table = next(candidate for candidate in rendered.tables if candidate.cell(0, 0).text == name)
                paragraph = next(item for item in table.cell(2, 0).paragraphs if item.text.strip())
                self.assertEqual(len(paragraph._p.findall(".//" + qn("w:br"))), 1)
                self.assertEqual(paragraph.text.replace("\n", ""), target)
                self.assertTrue(all(len(line) <= SCOPE_WRAP_CHARACTERS for line in paragraph.text.split("\n")))

    def test_an_environment_with_no_affected_location_still_reads_as_a_list(self) -> None:
        """N/A sat in a plain paragraph while the environment beside it was bulleted, so one cell of
        the pair read as a list and the other as a sentence."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            (report_folder / "evidence").mkdir()
            buffer = png_bytes(40, 20)
            (report_folder / "evidence" / "ev_na.png").write_bytes(buffer)
            now = datetime.now().astimezone()
            report = Report(
                report_id="r_na", app_id="CI-DOCX", saved_at=now,
                engagement=Engagement(
                    app_name="Northstar Banking", ci_number="CI-DOCX", app_owner="Avery Morgan", segment="JH", report_type="annual_pentest",
                    report_date=date(2026, 9, 9), tester="QA Tester", network="Internal",
                    tested_environments=["production", "non_production"], tested_channels=["web"],
                    test_windows={
                        "production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2)),
                        "non_production": TestWindow(start_date=date(2026, 8, 3), end_date=date(2026, 8, 4)),
                    },
                ),
                scope_targets=[
                    ScopeTarget(target_id="t_prod", environment="production", channel="web", value="https://prod.example.test"),
                    ScopeTarget(target_id="t_uat", environment="non_production", channel="web", value="https://uat.example.test"),
                ],
                evidence={"ev_na": EvidenceItem(file="evidence/ev_na.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)},
            )
            # Production only, so the Non-Production cell is the empty one under test.
            report.vulnerabilities = [self._finding("v_na", "Authorization bypass", "high", "001", ["t_prod"], [
                ImageFragment(frag_id="f_img", type="image", environment="production", evidence_id="ev_na", caption="Production response"),
            ])]
            rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))

            cell = next(
                cell
                for table in rendered.tables
                for row in table.rows for cell in row.cells
                # Matched on the heading rather than the label, which follows the configurable
                # non-production name and reads "Lower Region Environment:" by default.
                if "Environment:" in cell.text and "N/A" in cell.text
            )
            bullets = [paragraph for paragraph in cell.paragraphs if paragraph.style.name == "List Paragraph"]
            self.assertEqual([paragraph.text for paragraph in bullets], ["N/A"])
            self.assertIsNotNone(
                bullets[0]._p.find(qn("w:pPr")).find(qn("w:numPr")),
                "the N/A placeholder must be a real bullet, like the locations it stands in for",
            )
            self.assertNotIn("\u2022", cell.text, "the glyph comes from the list style, not the text")
            self.assertNotIn("{{", cell.text, "the token survived, which aborts generation")

    def test_a_value_with_no_separator_inside_the_limit_still_breaks(self) -> None:
        """A break moves back to the nearest special character. A value with none is cut at the
        limit, because letting the line run on would widen the table."""
        # A special character inside the limit is preferred over cutting mid-token.
        self.assertEqual(
            _wrap_long_value("https://prod.example.test/accounts/123/details/extra", 36),
            ["https://prod.example.test/accounts/", "123/details/extra"],
        )
        self.assertEqual(_wrap_long_value("x" * 80, 36), ["x" * 36, "x" * 36, "x" * 8])
        for limit in (LOCATION_WRAP_CHARACTERS, SCOPE_WRAP_CHARACTERS):
            for value in (
                "https://internal-banking-portal-uat.northstar.example.test/accounts/settings/notifications/preferences/email",
                "com.northstar.mobile.banking.application.android.enterprise.edition.release.candidate",
                "https://prod.example.test/reports/export?format=invalid&scope=all&page=2&size=100&sort=desc",
                "x" * 200,
            ):
                with self.subTest(limit=limit, value=value):
                    lines = _wrap_long_value(value, limit)
                    self.assertGreater(len(lines), 1, "fixture must be long enough to wrap")
                    self.assertEqual("".join(lines), value, "wrapping must not lose characters")
                    self.assertTrue(all(len(line) <= limit for line in lines), f"{lines} exceeds {limit}")
                    for line in lines[:-1]:
                        self.assertTrue(
                            not line[-1].isalnum() or len(line) == limit,
                            f"{line!r} should end on a special character or fill the line",
                        )

    def test_a_query_string_breaks_on_its_own_separators(self) -> None:
        # The limit is fixed here, so this pins where the algorithm breaks and not the shipped column
        # widths. Counting back from the 36th character lands on "=" rather than the earlier "?".
        self.assertEqual(
            _wrap_long_value("https://api.example.test/search?q=session+fixation&environment=production", 36),
            ["https://api.example.test/search?q=", "session+fixation&environment=", "production"],
        )

    def test_a_finding_without_a_number_leaves_the_id_blank(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            (report_folder / "evidence").mkdir()
            buffer = png_bytes(40, 20)
            (report_folder / "evidence" / "ev_blank.png").write_bytes(buffer)
            now = datetime.now().astimezone()
            report = Report(
                report_id="r_blank", app_id="CI-DOCX", saved_at=now,
                engagement=Engagement(
                    app_name="Northstar Banking", ci_number="CI-DOCX", app_owner="Avery Morgan", segment="JH", report_type="annual_pentest",
                    report_date=date(2026, 9, 9), tester="QA Tester", network="Internal", tested_environments=["production"], tested_channels=["web"],
                    test_windows={"production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2))},
                ),
                scope_targets=[ScopeTarget(target_id="t_web", environment="production", channel="web", value="https://prod.example.test")],
                evidence={"ev_blank": EvidenceItem(file="evidence/ev_blank.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)},
            )
            image = lambda index: [ImageFragment(frag_id=f"f_img_{index}", type="image", environment="production", evidence_id="ev_blank", caption="Production response")]
            report.vulnerabilities = [
                self._finding("v_numbered", "Numbered finding", "high", "042", ["t_web"], image(1)),
                self._finding("v_unnumbered", "Unnumbered finding", "low", None, ["t_web"], image(2)),
            ]
            rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))

            summary = next(table for table in rendered.tables if table.cell(0, 0).text == "Findings")
            numbers = {row.cells[0].text: row.cells[4].text for row in summary.rows[1:]}
            self.assertEqual(numbers, {"Numbered finding": "042", "Unnumbered finding": ""})

            details = {
                table.cell(0, 1).text: table.cell(1, 1).text
                for table in rendered.tables
                if table.cell(0, 0).text == "Severity" and len(table.rows) >= 5
            }
            self.assertEqual(details.get("High"), "042")
            self.assertEqual(details.get("Low"), "", "an unnumbered finding must not fall back to its internal uid")

            everything = "\n".join([*(p.text for p in rendered.paragraphs), *(c.text for t in rendered.tables for r in t.rows for c in r.cells)])
            self.assertNotIn("v_unnumbered", everything)

    def test_each_severity_section_and_each_later_finding_starts_on_a_new_page(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            (report_folder / "evidence").mkdir()
            (report_folder / "evidence" / "ev_pages.png").write_bytes(png_bytes(40, 20))
            now = datetime.now().astimezone()
            report = Report(
                report_id="r_pages", app_id="CI-DOCX", saved_at=now,
                engagement=Engagement(
                    app_name="Northstar Banking", ci_number="CI-DOCX", app_owner="Avery Morgan", segment="JH", report_type="annual_pentest",
                    report_date=date(2026, 9, 9), tester="QA Tester", network="Internal", tested_environments=["production"], tested_channels=["web"],
                    test_windows={"production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2))},
                ),
                scope_targets=[ScopeTarget(target_id="t_web", environment="production", channel="web", value="https://prod.example.test")],
                evidence={"ev_pages": EvidenceItem(file="evidence/ev_pages.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)},
            )
            image = lambda uid: [ImageFragment(frag_id=f"{uid}_img", type="image", environment="production", evidence_id="ev_pages", caption="Production response")]
            report.vulnerabilities = [
                self._finding("v_low", "Low finding", "low", "003", ["t_web"], image("v_low")),
                self._finding("v_high_b", "Beta high finding", "high", "002", ["t_web"], image("v_high_b")),
                self._finding("v_critical", "Critical finding", "critical", "001", ["t_web"], image("v_critical")),
                self._finding("v_high_a", "Alpha high finding", "high", "004", ["t_web"], image("v_high_a")),
            ]
            rendered = Document(BytesIO(render_report_docx(report, RESOURCES / "MAIN.docx", report_folder)))

            paragraphs = {paragraph.text: paragraph for paragraph in rendered.paragraphs}
            texts = [paragraph.text for paragraph in rendered.paragraphs]
            headings = ["Critical Findings", "High Findings", "Low Findings"]
            self.assertEqual([texts.index(heading) for heading in headings], sorted(texts.index(heading) for heading in headings))
            for heading in headings:
                with self.subTest(heading=heading):
                    self.assertIsNotNone(paragraphs[heading]._p.pPr.pageBreakBefore)
            self.assertLess(texts.index("Alpha high finding"), texts.index("Beta high finding"))
            self.assertIsNotNone(paragraphs["Beta high finding"]._p.pPr.pageBreakBefore)
            self.assertIsNone(paragraphs["Alpha high finding"]._p.pPr.pageBreakBefore)

    def _numbered_sections(self, report_folder: Path, *, continue_second: bool):
        """A proof of concept holding two numbered lists either side of an image, plus one in the
        description, so a continued chain can be told apart from a section boundary."""
        contents = [
            Content(type="description", fragments=[
                ListFragment(frag_id="f_desc_list", type="numbered_list", items=[ListItem(runs=[Run(text="Description step")])]),
            ]),
            Content(type="recommended_remediation", fragments=[
                ParagraphFragment(frag_id="f_rem", type="paragraph", runs=[Run(text="Remediate it")]),
            ]),
            Content(type="proof_of_concept", fragments=[
                ListFragment(frag_id="f_poc_one", type="numbered_list", items=[
                    ListItem(runs=[Run(text="First step")]), ListItem(runs=[Run(text="Second step")]),
                ]),
                ImageFragment(frag_id="f_poc_img", type="image", environment="production", evidence_id="ev_layout", caption="Proof"),
                ListFragment(frag_id="f_poc_two", type="numbered_list", continue_numbering=continue_second, items=[
                    ListItem(runs=[Run(text="Third step")]), ListItem(runs=[Run(text="Fourth step")]),
                ]),
            ]),
        ]
        document = self._layout_document(report_folder, contents)
        listed = [paragraph for paragraph in document.paragraphs if paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None]
        by_text = {paragraph.text: numbering_details(document, paragraph) for paragraph in listed}
        return by_text

    def test_a_continued_numbered_list_shares_one_numbering_with_the_list_above_it(self) -> None:
        """Word reads one numId as one list, so a chain shares it and writes a single startOverride.
        The chain shares its nsid too, which is why the control case below checks both."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            details = self._numbered_sections(Path(temporary_directory), continue_second=True)

            steps = [details[text] for text in ("First step", "Second step", "Third step", "Fourth step")]
            self.assertEqual(len({numbering_id for numbering_id, _, _ in steps}), 1, "a continued list did not join the list above it")
            self.assertEqual(len({nsid for _, nsid, _ in steps}), 1)
            # One numId means one w:num, so there is one startOverride behind all four paragraphs.
            # The count is restarted once, at the top of the chain, and never again inside it.
            self.assertEqual({override for _, _, override in steps}, {1}, "a continued list restarted the count")

            self.assertNotEqual(
                details["Description step"][0], steps[0][0],
                "description shares numbering with the proof of concept, so a chain could cross sections",
            )

    def test_two_numbered_lists_in_one_section_restart_independently_by_default(self) -> None:
        """This is new coverage, not a re-assertion. Nothing else pins this, so a future edit could
        silently renumber every report on disk."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            details = self._numbered_sections(Path(temporary_directory), continue_second=False)

            first, second = details["First step"], details["Third step"]
            self.assertNotEqual(first[0], second[0], "two independent lists shared a numId")
            self.assertNotEqual(first[1], second[1], "two independent lists shared an nsid")
            self.assertEqual((first[2], second[2]), (1, 1), "each independent list needs its own restart")

    def test_a_list_is_followed_by_one_blank_paragraph(self) -> None:
        """Steps and bullets ran straight into whatever followed them. The paragraph and table components
        carry their own trailing blank; the list component does not, because it is cloned once per item."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            document = self._layout_document(report_folder, [
                Content(type="description", fragments=[
                    ListFragment(frag_id="f_steps", type="numbered_list", items=[
                        ListItem(runs=[Run(text="Step one")]), ListItem(runs=[Run(text="Step two")]),
                    ]),
                    ParagraphFragment(frag_id="f_after", type="paragraph", runs=[Run(text="Text after the steps.")]),
                    ListFragment(frag_id="f_bullets", type="bulleted_list", items=[
                        ListItem(runs=[Run(text="A bullet")]), ListItem(runs=[Run(text="Another bullet")]),
                    ]),
                    ParagraphFragment(frag_id="f_last", type="paragraph", runs=[Run(text="Text after the bullets.")]),
                ]),
            ])
            texts = [
                "".join(node.text or "" for node in item.iter(qn("w:t")))
                for item in document.element.body.iterchildren() if item.tag == qn("w:p")
            ]

            def between(first: str, second: str) -> list[str]:
                return texts[texts.index(first) + 1:texts.index(second)]

            for kind, first, last, after in (
                ("numbered", "Step one", "Step two", "Text after the steps."),
                ("bulleted", "A bullet", "Another bullet", "Text after the bullets."),
            ):
                with self.subTest(list=kind):
                    self.assertEqual(between(last, after), [""], f"a {kind} list gained no blank line, or gained two")
                    self.assertEqual(between(first, last), [], f"a {kind} list got its blank between items, not after the last one")

    def test_a_list_that_ends_a_section_adds_no_blank_of_its_own(self) -> None:
        """The blank below a list is spacing between fragments. At the end of a section the template's own
        blank already separates it from the next heading. A second one would be a gap no other
        section has."""
        for kind in ("numbered_list", "bulleted_list"):
            with self.subTest(list=kind), tempfile.TemporaryDirectory() as temporary_directory:
                document = self._layout_document(Path(temporary_directory), [
                    Content(type="description", fragments=[
                        ListFragment(frag_id="f_list", type=kind, items=[ListItem(runs=[Run(text="Last item")])]),
                    ]),
                ])
                last = next(
                    item for item in document.element.body.iterchildren()
                    if item.tag == qn("w:p") and "".join(node.text or "" for node in item.iter(qn("w:t"))) == "Last item"
                )
                blanks = 0
                following = last.getnext()
                while following is not None and _is_blank_paragraph(following):
                    blanks += 1
                    following = following.getnext()
                self.assertEqual(blanks, 1, "a list ending a section kept its spacer on top of the template's blank")

    def test_only_the_generated_instance_label_is_bold(self) -> None:
        """The component is bold throughout and run formatting is additive, so the tester's words
        stay bold unless they are explicitly cleared."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            document = self._layout_document(report_folder, [
                Content(type="proof_of_concept", fragments=[
                    InstanceTitleFragment(frag_id="f_instance", type="instance_title", text="Production login page"),
                ]),
            ])
            paragraph = next(
                item for item in document.element.body.iterchildren()
                if item.tag == qn("w:p") and "Production login page" in "".join(node.text or "" for node in item.iter(qn("w:t")))
            )
            weights = []
            for run in paragraph.findall(qn("w:r")):
                if run.find(qn("w:t")) is None:
                    continue
                setting = run.find(qn("w:rPr"))
                bold = setting is not None and setting.find(qn("w:b")) is not None
                if bold and setting.find(qn("w:b")).get(qn("w:val")) in {"0", "false"}:
                    bold = False
                weights.append(("".join(node.text or "" for node in run.iter(qn("w:t"))), bold))
            self.assertEqual(weights, [("Instance 1:", True), (" Production login page", False)])

    def test_an_instance_label_is_numbered_once_even_when_the_tester_typed_it(self) -> None:
        """Three of the instance titles already saved carry the label in their text, from before the
        generator wrote it. Rendering both would read 'Instance 1: Instance 1: Production'."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            document = self._layout_document(report_folder, [
                Content(type="proof_of_concept", fragments=[
                    InstanceTitleFragment(frag_id="f_typed", type="instance_title", text="Instance 7: Production"),
                    InstanceTitleFragment(frag_id="f_plain", type="instance_title", text="Non-Production"),
                ]),
            ])
            text_of = lambda element: "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()
            titles = [
                text_of(item) for item in document.element.body.iterchildren()
                if item.tag == qn("w:p") and "Production" in text_of(item)
            ]
            # The tester's own number goes too. The position in the section decides the number,
            # not the typing.
            self.assertEqual(titles, ["Instance 1: Production", "Instance 2: Non-Production"])

    def test_an_instance_label_is_normalized_regardless_of_case(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            document = self._layout_document(report_folder, [
                Content(type="proof_of_concept", fragments=[
                    InstanceTitleFragment(frag_id="f_lower", type="instance_title", text="instance 7: Production"),
                    InstanceTitleFragment(frag_id="f_upper", type="instance_title", text="INSTANCE 3 - Mobile"),
                ]),
            ])
            titles = [paragraph.text for paragraph in document.paragraphs if paragraph.text.startswith("Instance ")]

            self.assertEqual(titles, ["Instance 1: Production", "Instance 2: Mobile"])

    def test_instance_label_normalization_requires_a_boundary_after_the_number(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            document = self._layout_document(report_folder, [
                Content(type="proof_of_concept", fragments=[
                    InstanceTitleFragment(frag_id="f_7zip", type="instance_title", text="Instance 7zip package"),
                ]),
            ])
            titles = [paragraph.text for paragraph in document.paragraphs if paragraph.text.startswith("Instance ")]

            self.assertEqual(titles, ["Instance 1: Instance 7zip package"])

    def test_instance_label_punctuation_requires_a_boundary_before_normalization(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            document = self._layout_document(report_folder, [
                Content(type="proof_of_concept", fragments=[
                    InstanceTitleFragment(frag_id="f_hyphen", type="instance_title", text="Instance 7-Zip package"),
                    InstanceTitleFragment(frag_id="f_decimal", type="instance_title", text="Instance 7.5 cluster"),
                ]),
            ])
            titles = [paragraph.text for paragraph in document.paragraphs if paragraph.text.startswith("Instance ")]

            self.assertEqual(titles, ["Instance 1: Instance 7-Zip package", "Instance 2: Instance 7.5 cluster"])

    def test_repeated_instance_labels_are_normalized_once(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            document = self._layout_document(report_folder, [
                Content(type="proof_of_concept", fragments=[
                    InstanceTitleFragment(frag_id="f_repeat", type="instance_title", text="Instance 9: Instance 3: Production"),
                ]),
            ])
            titles = [paragraph.text for paragraph in document.paragraphs if paragraph.text.startswith("Instance ")]

            self.assertEqual(titles, ["Instance 1: Production"])

    def _layout_document(self, report_folder: Path, contents: list, *, status: str = "open_new", tickets: str = "", template: str = "MAIN.docx", segment: str = "JH"):
        """Render one finding through the shipped template, for page-layout assertions."""
        evidence_folder = report_folder / "evidence"
        evidence_folder.mkdir(exist_ok=True)
        buffer = png_bytes(40, 20)
        (evidence_folder / "ev_layout.png").write_bytes(buffer)
        now = datetime.now().astimezone()
        report = Report(
            report_id="r_layout", app_id="CI-DOCX", saved_at=now,
            engagement=Engagement(
                app_name="Northstar Banking", ci_number="CI-DOCX", app_owner="Avery Morgan", segment=segment, report_type="annual_pentest",
                report_date=date(2026, 9, 9), tester="QA Tester", network="Internal", tested_environments=["production"], tested_channels=["web"],
                test_windows={"production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2))},
            ),
            scope_targets=[ScopeTarget(target_id="t_web", environment="production", channel="web", value="https://prod.example.test")],
            evidence={"ev_layout": EvidenceItem(file="evidence/ev_layout.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)},
        )
        report.vulnerabilities = [Vulnerability(
            uid="v_layout", display_id="001", title="Layout finding",
            likelihood="high", impact="high", severity="high", status=status,
            scope=Scope(mode="custom", target_ids=["t_web"]), contents=contents,
            severity_review_tickets=tickets,
        )]
        return Document(BytesIO(render_report_docx(report, RESOURCES / template, report_folder, allow_incomplete=True)))

    @staticmethod
    def _table(suffix: str) -> TableFragment:
        return TableFragment(
            frag_id=f"f_table_{suffix}", type="table",
            header=[ListItem(runs=[Run(text="Header")])],
            rows=[[ListItem(runs=[Run(text=f"Row {suffix}")])]],
        )

    def test_a_table_repeats_its_header_row_and_keeps_each_row_on_one_page(self) -> None:
        """table_fragment.docx sets neither, so Word would split a row across two pages and leave the columns of a
        long table unlabelled from its second page on. The header repeats only when the table has a header row."""
        cell = lambda text: ListItem(runs=[Run(text=text)])
        cases = [
            ("a header and rows", [cell("Alpha"), cell("Beta")], [[cell("gamma"), cell("delta")], [cell("epsilon"), cell("zeta")]], True),
            ("a header only", [cell("Alpha"), cell("Beta")], [], True),
            ("no header", [], [[cell("gamma"), cell("delta")]], False),
        ]
        for name, header, rows, repeats in cases:
            with self.subTest(name), tempfile.TemporaryDirectory() as temporary_directory:
                fragment = TableFragment(frag_id="f_rows", type="table", header=header, rows=rows)
                document = self._layout_document(Path(temporary_directory), [Content(type="description", fragments=[fragment])])
                table = next(table for table in document.tables if any(item.text in {"Alpha", "gamma"} for row in table.rows for item in row.cells))
                properties = [row._tr.find(qn("w:trPr")) for row in table.rows]
                self.assertGreaterEqual(len(properties), 1)
                for position, row_properties in enumerate(properties):
                    self.assertIsNotNone(row_properties, f"row {position} has no properties")
                    self.assertIsNotNone(row_properties.find(qn("w:cantSplit")), f"row {position} can split across pages")
                    self.assertEqual(row_properties.find(qn("w:tblHeader")) is not None, repeats and position == 0, f"row {position}")
                # w:trPr comes before the cells, as the schema says.
                self.assertTrue(all(row._tr[0].tag == qn("w:trPr") for row in table.rows))

    def test_a_table_is_followed_by_one_spacer_so_two_tables_never_merge(self) -> None:
        """A table fragment contributes only its w:tbl, and Word silently merges adjacent tables."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            document = self._layout_document(report_folder, [
                Content(type="description", fragments=[self._table("a"), self._table("b")]),
                Content(type="recommended_remediation", fragments=[ParagraphFragment(frag_id="f_fix", type="paragraph", runs=[Run(text="Apply the fix.")])]),
                Content(type="proof_of_concept", fragments=[
                    ListFragment(frag_id="f_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])]),
                    ImageFragment(frag_id="f_img", type="image", environment="production", evidence_id="ev_layout", caption="Production response"),
                ]),
            ])
            body = list(document.element.body.iterchildren())
            for left, right in zip(body, body[1:]):
                self.assertFalse(left.tag == qn("w:tbl") and right.tag == qn("w:tbl"), "adjacent tables merge into one in Word")

            first = next(index for index, element in enumerate(body) if element.tag == qn("w:tbl") and "Row a" in "".join(node.text or "" for node in element.iter(qn("w:t"))))
            spacer = body[first + 1]
            self.assertTrue(_is_blank_paragraph(spacer), "a mid-section table needs a spacer below it")
            self.assertEqual(body[first + 2].tag, qn("w:tbl"), "exactly one spacer, not two")
            properties = spacer.find(qn("w:pPr"))
            # An empty paragraph carrying either of these is read back as a fragment by docx_import.
            self.assertIsNone(properties.find(qn("w:numPr")), "a numbered spacer imports as a phantom list")
            self.assertIsNone(properties.find(qn("w:pStyle")), "a caption-styled spacer erases the previous image caption")

            last_table = body[first + 2]
            blanks = 0
            following = last_table.getnext()
            while following is not None and _is_blank_paragraph(following):
                blanks += 1
                following = following.getnext()
            self.assertEqual(blanks, 1, "a table ending a section keeps only the template blank; its spacer is trimmed")

    def test_section_headings_are_kept_on_the_page_with_the_content_below_them(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            paragraph = lambda name: [ParagraphFragment(frag_id=f"f_{name}", type="paragraph", runs=[Run(text=f"The {name} text.")])]
            document = self._layout_document(report_folder, [
                Content(type="description", fragments=paragraph("description")),
                Content(type="recommended_remediation", fragments=paragraph("remediation")),
                Content(type="previous_proof_of_concept", fragments=[ListFragment(frag_id="f_prev", type="numbered_list", items=[ListItem(runs=[Run(text="Reproduce it.")])])]),
                Content(type="proof_of_concept", fragments=[
                    ListFragment(frag_id="f_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])]),
                    ImageFragment(frag_id="f_img", type="image", environment="production", evidence_id="ev_layout", caption="Production response"),
                ]),
                Content(type="in_conclusion", fragments=paragraph("conclusion")),
            ], status="open_previously_discovered")

            blanks_checked = 0
            for title in ("Description:", "Recommended Remediation:", "Previous Proof of Concept:", "Proof of Concept:", "In Conclusion:", "The following demonstrates the vulnerability:"):
                heading = next(item for item in document.paragraphs if item.text.strip() == title)
                self.assertTrue(_keeps_next(heading._p), f"{title!r} must be kept with its content")
                previous = heading._p.getprevious()
                if previous is not None and _is_blank_paragraph(previous):
                    blanks_checked += 1
                    self.assertFalse(_keeps_next(previous), f"the walk above {title!r} must stop at the blank")
            self.assertGreater(blanks_checked, 0, "no heading had a blank above it, so the walk's stop went unchecked")

            title = next(item for item in document.paragraphs if item.text.strip() == "Layout finding")
            self.assertFalse(_keeps_next(title._p), "a finding title already carries its own page break")

    def test_the_severity_review_ticket_label_is_kept_on_the_page_with_its_value(self) -> None:
        """Word can break the page between the label and its value, which leaves the label alone at the
        foot of one page with N/A or the tickets at the top of the next."""
        paragraph = lambda name: [ParagraphFragment(frag_id=f"f_{name}", type="paragraph", runs=[Run(text=f"The {name} text.")])]
        contents = lambda: [
            Content(type="description", fragments=paragraph("description")),
            Content(type="recommended_remediation", fragments=paragraph("remediation")),
            Content(type="previous_proof_of_concept", fragments=paragraph("previous")),
            Content(type="proof_of_concept", fragments=paragraph("proof")),
            Content(type="in_conclusion", fragments=paragraph("conclusion")),
        ]
        cases = [
            ("open_previously_discovered", "", "N/A"),
            ("open_previously_discovered", "1234\n5678", "GRIMPEN-1234GRIMPEN-5678"),
            ("open_resolved_on_non_prod", "", "N/A"),
            ("resolved", "", "N/A"),
            ("closed", "", "N/A"),
        ]
        for status, tickets, printed in cases:
            with self.subTest(status=status, tickets=tickets), tempfile.TemporaryDirectory() as temporary_directory:
                document = self._layout_document(Path(temporary_directory), contents(), status=status, tickets=tickets)
                label = next(item for item in document.paragraphs if item.text.strip() == "Severity Review Ticket (if applicable):")
                self.assertTrue(_keeps_next(label._p), "the label must stay on the page where its value starts")

                # keepNext only helps when the very next paragraph is the value, not a template blank.
                value = label._p.getnext()
                self.assertEqual("".join(node.text or "" for node in value.iter(qn("w:t"))), printed)
                self.assertFalse(_keeps_next(value), "the chain ends at the value")

                above = label._p.getprevious()
                self.assertTrue(_is_blank_paragraph(above), "the walk relies on a blank paragraph above the label to stop at")
                self.assertFalse(_keeps_next(above), "the walk above the label must stop at the blank")

    def test_severity_review_tickets_print_one_prefixed_value_per_line(self) -> None:
        """The prefix belongs to the document. The draft stores bare digits, so a tester never types
        it and an imported value never carries it back in."""
        paragraph = lambda name: [ParagraphFragment(frag_id=f"f_{name}", type="paragraph", runs=[Run(text=f"The {name} text.")])]
        retest_contents = lambda: [
            Content(type="description", fragments=paragraph("description")),
            Content(type="recommended_remediation", fragments=paragraph("remediation")),
            Content(type="previous_proof_of_concept", fragments=paragraph("previous")),
            Content(type="proof_of_concept", fragments=paragraph("proof")),
            Content(type="in_conclusion", fragments=paragraph("conclusion")),
        ]

        def ticket_value(document):
            """The first paragraph carrying text below the label, so a template blank cannot fool it."""
            label = next(item for item in document.paragraphs if item.text.strip() == "Severity Review Ticket (if applicable):")
            node = label._p.getnext()
            while node is not None and not "".join(part.text or "" for part in node.iter(qn("w:t"))).strip():
                node = node.getnext()
            self.assertIsNotNone(node, "the label must be followed by its value")
            return node

        with tempfile.TemporaryDirectory() as temporary_directory:
            document = self._layout_document(
                Path(temporary_directory), retest_contents(), status="open_previously_discovered", tickets="1234\n5678\n90123",
            )
            value = ticket_value(document)
            self.assertEqual("".join(node.text or "" for node in value.iter(qn("w:t"))), "GRIMPEN-1234GRIMPEN-5678GRIMPEN-90123")
            # Three values are two breaks. Joined onto one line the text above would still match.
            self.assertEqual(len(value.findall(".//" + qn("w:br"))), 2)

        with tempfile.TemporaryDirectory() as temporary_directory:
            document = self._layout_document(Path(temporary_directory), retest_contents(), status="open_previously_discovered")
            value = ticket_value(document)
            self.assertEqual("".join(node.text or "" for node in value.iter(qn("w:t"))), "N/A")
            self.assertEqual(len(value.findall(".//" + qn("w:br"))), 0)

        with tempfile.TemporaryDirectory() as temporary_directory:
            document = self._layout_document(Path(temporary_directory), [
                Content(type="description", fragments=paragraph("description")),
                Content(type="recommended_remediation", fragments=paragraph("remediation")),
                Content(type="proof_of_concept", fragments=paragraph("proof")),
            ], tickets="1234")
            body = "\n".join(item.text for item in document.paragraphs)
            self.assertNotIn("Severity Review Ticket", body, "new_finding.docx carries no such label")
            self.assertNotIn("GRIMPEN-", body, "a value on an Open (New) finding must reach nothing")

    def _code_block_above(self, fragments, *, section: str = "proof_of_concept", status: str = "open_new"):
        """Render `fragments` as one section of a finding and return the element directly above the
        code block that says GET /accounts/123, plus the document body for further looking."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            contents = [
                Content(type="description", fragments=[ParagraphFragment(frag_id="f_d", type="paragraph", runs=[Run(text="Description text")])]),
                Content(type="recommended_remediation", fragments=[ParagraphFragment(frag_id="f_r", type="paragraph", runs=[Run(text="Remediation text")])]),
                Content(type="proof_of_concept", fragments=[ListFragment(frag_id="f_s", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])])]),
            ]
            contents = [content for content in contents if content.type != section] + [Content(type=section, fragments=fragments)]
            document = self._layout_document(Path(temporary_directory), contents, status=status)
            body = list(document.element.body.iterchildren())
        text_of = lambda element: "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()
        code = next(item for item in body if item.tag == qn("w:p") and text_of(item) == "GET /accounts/123")
        return code.getprevious(), body, text_of

    def test_a_code_block_in_a_proof_of_concept_has_no_blank_line_above_it(self) -> None:
        """Whatever introduces a code block leaves a spacer paragraph behind it, which put an empty
        line between the steps and the command they lead into."""
        steps = ListFragment(frag_id="f_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])])
        for name, introduction in (
            ("a numbered list", steps),
            ("a paragraph", ParagraphFragment(frag_id="f_p", type="paragraph", runs=[Run(text="The request is:")])),
            ("a note", NoteFragment(frag_id="f_n", type="note", runs=[Run(text="Use a valid session.")])),
        ):
            for section, status in (("proof_of_concept", "open_new"), ("previous_proof_of_concept", "open_previously_discovered")):
                with self.subTest(after=name, section=section):
                    code = CodeFragment(frag_id="f_code", type="code_block", caption=None, text="GET /accounts/123")
                    above, body, text_of = self._code_block_above([introduction, code], section=section, status=status)
                    self.assertFalse(_is_blank_paragraph(above), f"an empty line sits between {name} and its code block")
                    self.assertTrue(text_of(above), "the element above the code block is not the text that introduces it")

    def test_a_captioned_code_block_keeps_the_space_above_its_caption_and_loses_the_one_below_it(self) -> None:
        steps = ListFragment(frag_id="f_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])])
        code = CodeFragment(frag_id="f_code", type="code_block", caption="Request", text="GET /accounts/123")
        above, body, text_of = self._code_block_above([steps, code])
        self.assertEqual(text_of(above), "Request", "the caption must sit directly on its code block")
        self.assertTrue(_keeps_next(above), "the caption must stay with the code block under it")
        self.assertTrue(_is_blank_paragraph(above.getprevious()), "the caption lost the space that separates it from the steps")

    def test_two_code_blocks_in_a_proof_of_concept_are_laid_out_as_they_were(self) -> None:
        """Adjacent code blocks never had a spacer between them. Removing the spacer above a code block
        must not remove the first block's own paragraph."""
        first = CodeFragment(frag_id="f_first", type="code_block", caption=None, text="POST /login")
        second = CodeFragment(frag_id="f_code", type="code_block", caption=None, text="GET /accounts/123")
        above, body, text_of = self._code_block_above([first, second])
        self.assertEqual(text_of(above), "POST /login")

    def test_a_code_block_outside_the_proof_of_concept_keeps_its_blank_line(self) -> None:
        """Only the proof of concept loses the blank line above a code block. A Description keeps
        the layout it has."""
        intro = ParagraphFragment(frag_id="f_p", type="paragraph", runs=[Run(text="The request is:")])
        code = CodeFragment(frag_id="f_code", type="code_block", caption=None, text="GET /accounts/123")
        above, body, text_of = self._code_block_above([intro, code], section="description")
        self.assertTrue(_is_blank_paragraph(above))

    def test_labels_titles_and_lead_ins_are_kept_with_what_they_introduce(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            document = self._layout_document(report_folder, [
                Content(type="description", fragments=[
                    ParagraphFragment(frag_id="f_lead", type="paragraph", runs=[Run(text="Use these values:")]),
                    self._table("a"),
                    CodeFragment(frag_id="f_code", type="code_block", caption="Request", text="GET /accounts/123"),
                    CodeFragment(frag_id="f_bare", type="code_block", caption=None, text="GET /health"),
                ]),
                Content(type="recommended_remediation", fragments=[ParagraphFragment(frag_id="f_fix", type="paragraph", runs=[Run(text="Apply the fix.")])]),
                Content(type="proof_of_concept", fragments=[
                    ListFragment(frag_id="f_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])]),
                    InstanceTitleFragment(frag_id="f_instance", type="instance_title", text="Instance 1: Production"),
                    ImageFragment(frag_id="f_img", type="image", environment="production", evidence_id="ev_layout", caption="Production response"),
                ]),
            ])
            body = list(document.element.body.iterchildren())
            text_of = lambda element: "".join(node.text or "" for node in element.iter(qn("w:t"))).strip()

            for label in ("PROD:", "Instance 1: Production", "Request"):
                element = next(item for item in body if item.tag == qn("w:p") and text_of(item) == label)
                self.assertTrue(_keeps_next(element), f"{label!r} must be kept with what follows it")

            lead_in = next(item for item in body if item.tag == qn("w:p") and text_of(item) == "Use these values:")
            self.assertTrue(_keeps_next(lead_in), "a table lead-in must not be stranded above its table")
            self.assertTrue(_keeps_next(lead_in.getnext()), "the blank between a lead-in and its table must not break the chain")

            # Located from the caption, because the template's own cover art is the first drawing.
            caption = next(item for item in body if item.tag == qn("w:p") and text_of(item).endswith("Production response"))
            image = caption.getprevious()
            self.assertTrue(image.xpath(".//w:drawing"), "the caption must sit directly below its image")
            self.assertTrue(_keeps_next(image), "an image must stay with its caption")
            self.assertFalse(_keeps_next(caption), "a caption ends a chain and must not start another")

            bare_code = next(item for item in body if item.tag == qn("w:p") and text_of(item) == "GET /health")
            self.assertFalse(_keeps_next(bare_code.getprevious()), "an uncaptioned code block has no heading to keep")

    def test_multiline_text_fragments_render_as_paragraphs_not_manual_breaks(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            document = self._layout_document(Path(temporary_directory), [
                Content(type="description", fragments=[
                    ParagraphFragment(
                        frag_id="f_multiline_paragraph",
                        type="paragraph",
                        runs=[Run(text="First paragraph.\r\n"), Run(text="Second paragraph.", bold=True)],
                    ),
                    NoteFragment(
                        frag_id="f_multiline_note",
                        type="note",
                        runs=[Run(text="First note paragraph.\nSecond note paragraph.")],
                    ),
                ]),
                Content(type="recommended_remediation", fragments=[
                    ParagraphFragment(frag_id="f_fix", type="paragraph", runs=[Run(text="Apply the fix.")]),
                ]),
                Content(type="proof_of_concept", fragments=[
                    ListFragment(frag_id="f_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])]),
                    ImageFragment(frag_id="f_img", type="image", environment="production", evidence_id="ev_layout", caption="Production response"),
                ]),
            ])

            expected = [
                "First paragraph.",
                "Second paragraph.",
                "Note: First note paragraph.",
                "Second note paragraph.",
            ]
            paragraph_texts = [paragraph.text for paragraph in document.paragraphs]
            indexes = [paragraph_texts.index(text) for text in expected]
            paragraphs = [document.paragraphs[index] for index in indexes]
            self.assertEqual(indexes[1], indexes[0] + 1)
            self.assertEqual(indexes[3], indexes[2] + 1)
            self.assertTrue(paragraphs[1].runs[0].bold, "formatting after the newline was lost")
            self.assertEqual(sum(len(paragraph._p.findall(".//" + qn("w:br"))) for paragraph in paragraphs), 0)
            self.assertEqual(sum(paragraph.text.startswith("Note:") for paragraph in paragraphs), 1)

    def _assert_image_fragment_format(self, document) -> None:
        for index in range(len(document.inline_shapes)):
            shape = document.inline_shapes[index]
            line = shape._inline.graphic.graphicData.pic.spPr.find(qn("a:ln"))
            self.assertIsNone(line)
            relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.get(qn("r:embed"))
            image = Image.open(BytesIO(document.part.related_parts[relationship_id].blob)).convert("RGB")
            self.assertEqual(off_border_pixels(image), [], "the screenshot lost its black border")
            paragraph = next(shape._inline.iterancestors(qn("w:p")))
            alignment = paragraph.find(qn("w:pPr") + "/" + qn("w:jc"))
            self.assertIsNotNone(alignment)
            self.assertEqual(alignment.get(qn("w:val")), "center")

    def _component_report(self, report_folder: Path, channel: str, segment: str, targets: list[ScopeTarget]) -> Report:
        """A complete single-finding report covering one component app type."""
        evidence_folder = report_folder / "evidence"
        evidence_folder.mkdir(exist_ok=True)
        buffer = png_bytes(40, 20)
        (evidence_folder / "ev_prod.png").write_bytes(buffer)
        now = datetime.now().astimezone()
        report = Report(
            report_id="r_component",
            app_id="CI-DOCX",
            saved_at=now,
            engagement=Engagement(
                app_name="Northstar Banking",
                ci_number="CI-DOCX",
                app_owner="Avery Morgan",
                segment=segment,
                report_type="annual_pentest",
                report_date=date(2026, 9, 9),
                tester="QA Tester",
                network="Internal",
                tested_environments=["production"],
                tested_channels=[channel],
                test_windows={"production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2), test_time="22:00 EST")},
            ),
            scope_targets=targets,
            evidence={"ev_prod": EvidenceItem(file="evidence/ev_prod.png", original_name="ev_prod.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)},
        )
        report.vulnerabilities = [
            self._finding("v_component", "Authorization bypass", "high", "001", [next(target.target_id for target in targets if target.environment == "production")], [
                ImageFragment(frag_id="f_img_prod", type="image", environment="production", evidence_id="ev_prod", caption="Production response"),
            ]),
        ]
        if segment == "Asia":
            # Asia requires both, so the fixture is not generation-clean without them.
            report.vulnerabilities[0].cvss_score = "8.1"
            report.vulnerabilities[0].cvss_vector = "CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:U/C:H/I:H/A:N"
        if segment == "GWAM":
            # A GWAM report types its own name, so the fixture is not generation-clean without one.
            report.engagement.report_name = "GWAM - Northstar Banking - Annual Pentest 2026"
        return report

    def test_main_template_path_selects_on_both_axes(self) -> None:
        resources = RESOURCES
        covered = set()
        with tempfile.TemporaryDirectory() as temporary_directory:
            for channel, segment, expected in (
                ("web", "JH", "MAIN.docx"),
                ("web", "Asia", "MAIN_ASIA.docx"),
                ("thick_client", "JH", "MAIN_THICK_MOBILE.docx"),
                ("mobile", "Asia", "MAIN_THICK_MOBILE_ASIA.docx"),
                # Asia is a boolean, not one value among several, so GDT and GFT land in the JH
                # column by design. The decision lives here rather than in the code it does not change.
                ("web", "GDT", "MAIN.docx"),
                ("thick_client", "GDT", "MAIN_THICK_MOBILE.docx"),
                ("web", "GFT", "MAIN.docx"),
                ("thick_client", "GFT", "MAIN_THICK_MOBILE.docx"),
                ("web", "GWAM", "MAIN.docx"),
                ("mobile", "GWAM", "MAIN_THICK_MOBILE.docx"),
            ):
                covered.add(segment)
                report = self._component_report(Path(temporary_directory), channel, segment, [
                    ScopeTarget(target_id="t_one", environment="production", channel=channel, value="Acme.exe", description="Main client"),
                ])
                chosen = main_template_path(report, resources)
                self.assertEqual(chosen.name, expected)
                # The parent doubles as the component fragment root, so every branch must stay in resources.
                self.assertEqual(chosen.parent, resources)
                self.assertTrue(chosen.is_file())
        # Written out by hand on purpose, so a new segment must be added here rather than guessed.
        self.assertEqual(covered, set(get_args(Segment)), "a segment has no row in the template table")

    def test_both_cover_blocks_print_the_chosen_network_access_on_every_template(self) -> None:
        """Each cover is drawn twice (the shape and its fallback); if only one carried the token the
        two would disagree and import would refuse the document. The Remediation Timelines header
        also reads Internal and must survive."""
        for channel, segment in (("web", "JH"), ("web", "Asia"), ("thick_client", "JH"), ("mobile", "Asia")):
            for network in ("External", "Internal"):
                with self.subTest(channel=channel, segment=segment, network=network), tempfile.TemporaryDirectory() as temporary_directory:
                    report_folder = Path(temporary_directory)
                    report = self._component_report(report_folder, channel, segment, [
                        ScopeTarget(target_id="t_one", environment="production", channel=channel, value="Acme.exe", description="Main client"),
                    ])
                    report.engagement.network = network
                    rendered = Document(BytesIO(render_report_docx(report, main_template_path(report, RESOURCES), report_folder)))
                    texts = [
                        "".join(node.text or "" for node in paragraph.iter(qn("w:t"))).strip()
                        for paragraph in rendered.element.body.iter(qn("w:p"))
                    ]

                    self.assertEqual([texts[index + 1] for index, text in enumerate(texts) if text == "Application Type"], [network, network])
                    header = next(table for table in rendered.tables if any(cell.text.strip() == "Emergency" for row in table.rows for cell in row.cells))
                    self.assertIn("Internal", [cell.text.strip() for row in header.rows for cell in row.cells])

    def test_each_master_prints_its_report_name_from_one_token_on_the_cover_and_in_both_footers(self) -> None:
        for master in MASTERS:
            with self.subTest(master):
                document = Document(RESOURCES / master)
                body = paragraph_texts(document.element.body)
                footers = _footer_texts(document)

                self.assertEqual(body.count("{{report-name}}"), 1, "the cover title")
                self.assertEqual(sorted(name for name, texts in footers.items() if "{{report-name}}" in texts), ["/word/footer5.xml", "/word/footer8.xml"])
                every_text = [*body, *(text for texts in footers.values() for text in texts)]
                for left_over in ("{{segment}}", "{{test-type}}"):
                    self.assertEqual([text for text in every_text if left_over in text], [], f"{left_over} is still spelled out in a title line")

    def test_each_master_keeps_its_application_cell_and_its_footer_template_id(self) -> None:
        for master in MASTERS:
            with self.subTest(master):
                document = Document(RESOURCES / master)

                self.assertIn("{{app-name}}", paragraph_texts(document.element.body), "the cover table's Application cell")
                self.assertTrue(any("GCSAppSec_Report2026v1" in text for texts in _footer_texts(document).values() for text in texts), "the template ID in the footers")

    def test_the_report_name_prints_on_the_cover_and_in_both_footers_while_the_application_name_stays_put(self) -> None:
        for channel in ("web", "thick_client"):
            for segment in get_args(Segment):
                with self.subTest(channel=channel, segment=segment), tempfile.TemporaryDirectory() as temporary_directory:
                    report_folder = Path(temporary_directory)
                    report = self._component_report(report_folder, channel, segment, [
                        ScopeTarget(target_id="t_one", environment="production", channel=channel, value="Acme.exe", description="Main client"),
                    ])
                    report.engagement.report_date = date(2027, 3, 4)
                    # Typed under every segment, so a segment that has no report name of its own must ignore it.
                    typed = "GWAM - Meridian Portal \u2013 Annual Pentest 2027" if segment == "GWAM" else "GWAM - Typed Name - Retest 2026"
                    report.engagement.report_name = typed
                    expected = typed if segment == "GWAM" else f"{segment} \u2013 Northstar Banking \u2013 Annual Pentest 2027"
                    self.assertEqual(generation_issues(report), [])

                    rendered = Document(BytesIO(render_report_docx(report, main_template_path(report, RESOURCES), report_folder)))

                    body = paragraph_texts(rendered.element.body)
                    footers = _footer_texts(rendered)
                    self.assertEqual(body.count(expected), 1, "the cover title")
                    self.assertEqual(sorted(name for name, texts in footers.items() if expected in texts), ["/word/footer5.xml", "/word/footer8.xml"])
                    self.assertTrue(any("GCSAppSec_Report2026v1" in text for texts in footers.values() for text in texts), "the template ID in the footers")
                    self.assertEqual([body[index + 1] for index, text in enumerate(body) if text == "Application"], ["Northstar Banking", "Northstar Banking"])
                    # The places that name the application, not the report.
                    body_text = "\n".join(body)
                    for place in (
                        "requested by the Northstar Banking application team",
                        "Northstar Banking PENTEST",
                        "team members of the Northstar Banking application",
                        "used by the Northstar Banking web application",
                        "Summary of findings for Northstar Banking",
                    ):
                        self.assertIn(place, body_text)
                    if segment != "GWAM":
                        self.assertNotIn("Typed Name", "\n".join([body_text, *(text for texts in footers.values() for text in texts)]))

    def test_a_report_name_holding_a_token_word_prints_as_typed(self) -> None:
        """The value is free text and goes in last, so no later pass may rewrite it, and a pass that
        found its own output again would never finish."""
        typed = "GWAM - app-name and report-name for app-owner - tester-name 2026"
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            report = self._component_report(report_folder, "web", "GWAM", [
                ScopeTarget(target_id="t_one", environment="production", channel="web", value="https://prod.example.test"),
            ])
            report.engagement.report_name = typed
            self.assertEqual(generation_issues(report), [])
            outcome: dict = {}

            def render() -> None:
                try:
                    outcome["document"] = Document(BytesIO(render_report_docx(report, main_template_path(report, RESOURCES), report_folder)))
                except Exception as error:  # reported by the assertions below
                    outcome["error"] = error

            worker = threading.Thread(target=render, daemon=True)
            worker.start()
            worker.join(60)
            self.assertFalse(worker.is_alive(), "rendering a name that holds a token word never finished")
            self.assertNotIn("error", outcome)

            body = paragraph_texts(outcome["document"].element.body)
            self.assertEqual(body.count(typed), 1, "the cover title")
            self.assertEqual(sorted(name for name, texts in _footer_texts(outcome["document"]).items() if typed in texts), ["/word/footer5.xml", "/word/footer8.xml"])

    def test_no_shipped_document_spells_report_name_without_braces(self) -> None:
        """`report-name` is no longer a word the renderer replaces wherever it finds it, because the
        value is free text. Only the braced token is replaced, so a bare spelling would print as is."""
        documents = sorted([*RESOURCES.glob("*.docx"), *(path for folder in ("fragments", "finding_types", "severity_titles") for path in (RESOURCES / folder).glob("*.docx"))])
        self.assertGreater(len(documents), 15, "too few documents were checked")
        for path in documents:
            with self.subTest(str(path.relative_to(RESOURCES))):
                document = Document(path)
                roots = [document.element, *(part._element for part in document.part.package.parts if "/header" in str(part.partname) or "/footer" in str(part.partname))]
                spelled = [text for root in roots for text in paragraph_texts(root) if "report-name" in re.sub(r"\{\{\s*report-name\s*\}\}", "", text)]
                self.assertEqual(spelled, [])

    def test_a_blank_account_row_is_left_out_of_the_user_roles_table(self) -> None:
        cases = [
            ("among filled rows", [TestAccount(user_role="Admin", username="qa.admin"), TestAccount(user_role=" ", username="")], [["Admin", "qa.admin"]]),
            ("the only row", [TestAccount(user_role="", username="")], [["N/A", "N/A"]]),
        ]
        for label, accounts, expected in cases:
            with self.subTest(label):
                document = Document(RESOURCES / "MAIN.docx")
                report = Report(report_id="r_accounts", app_id="CI-ACCOUNTS", saved_at=datetime.now().astimezone())
                report.engagement.test_accounts = accounts

                _populate_scope_tables(document, report)

                table = _find_table(document, "User Roles")
                self.assertEqual([[cell.text for cell in row.cells] for row in table.rows[1:]], expected)
                self.assertEqual(report.engagement.test_accounts, accounts, "the draft's rows must not change")

    def test_every_shipped_template_renders_without_unresolved_placeholders(self) -> None:
        """None of the four has been through this renderer before. Each carries its own anchors,
        table headers and tokens, and every one of them is a hard precondition."""
        resources = RESOURCES
        for channel, segment in (("web", "JH"), ("web", "Asia"), ("thick_client", "JH"), ("mobile", "Asia"), ("web", "GDT"), ("thick_client", "GDT"), ("web", "GFT"), ("thick_client", "GFT"), ("web", "GWAM"), ("mobile", "GWAM")):
            for has_findings in (True, False):
                with self.subTest(channel=channel, segment=segment, has_findings=has_findings), tempfile.TemporaryDirectory() as temporary_directory:
                    report_folder = Path(temporary_directory)
                    report = self._component_report(report_folder, channel, segment, [
                        ScopeTarget(target_id="t_one", environment="production", channel=channel, value="Acme.exe", description="Main client"),
                    ])
                    if not has_findings:
                        report.vulnerabilities = []
                    self.assertEqual(generation_issues(report), [])
                    render_report_docx(report, main_template_path(report, resources), report_folder)

    def test_component_scope_fills_the_binaries_table_production_first(self) -> None:
        resources = RESOURCES
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            report = self._component_report(report_folder, "thick_client", "JH", [
                ScopeTarget(target_id="t_nonprod", environment="non_production", channel="thick_client", value="Acme.Staging.exe", description="UAT build", order=0),
                ScopeTarget(target_id="t_prod", environment="production", channel="thick_client", value="Acme.exe", description="Main client", order=0),
                ScopeTarget(target_id="t_prod_2", environment="production", channel="thick_client", value="Acme.Updater.exe", description="Background updater", order=1),
            ])
            report.engagement.tested_environments = ["production", "non_production"]
            report.engagement.test_windows["non_production"] = TestWindow(start_date=date(2026, 7, 28), end_date=date(2026, 7, 30))
            rendered = Document(BytesIO(render_report_docx(report, main_template_path(report, resources), report_folder)))

            table = next(table for table in rendered.tables if table.cell(0, 0).text.strip() == "Component")
            # One cloned row per component, because stacked paragraphs would misalign every row
            # below a wrap.
            self.assertEqual([(row.cells[0].text, row.cells[1].text) for row in table.rows[1:]], [
                ("Acme.exe", "Main client"),
                ("Acme.Updater.exe", "Background updater"),
                ("Acme.Staging.exe", "UAT build"),
            ])
            self.assertIn("Thick Client", "\n".join(paragraph.text for paragraph in rendered.paragraphs))

    def test_an_empty_component_list_leaves_one_placeholder_row_not_the_prototype(self) -> None:
        """An untouched {{binaries}} row would fail the unresolved-placeholder check at the very end
        of generation, which is the least useful place to discover an empty scope."""
        resources = RESOURCES
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            report = self._component_report(report_folder, "mobile", "JH", [
                ScopeTarget(target_id="t_web", environment="production", channel="web", value="https://prod.example.test"),
            ])
            report.engagement.tested_channels = ["web", "mobile"]
            rendered = Document(BytesIO(render_report_docx(report, main_template_path(report, resources), report_folder)))

            table = next(table for table in rendered.tables if table.cell(0, 0).text.strip() == "Component")
            self.assertEqual([(row.cells[0].text, row.cells[1].text) for row in table.rows[1:]], [("N/A", "N/A")])
            self.assertIn("Mobile", "\n".join(paragraph.text for paragraph in rendered.paragraphs))

    @staticmethod
    def _fragments_holding(texts: list[str]) -> dict[str, list]:
        """One section's worth of every fragment type that carries free text, each holding every text given."""
        cell = lambda text: ListItem(runs=[Run(text=text)])
        return {
            "paragraphs": [ParagraphFragment(frag_id=f"f_p{index}", type="paragraph", runs=[Run(text=text)]) for index, text in enumerate(texts)],
            "notes": [NoteFragment(frag_id=f"f_n{index}", type="note", runs=[Run(text=text)]) for index, text in enumerate(texts)],
            "a numbered list": [ListFragment(frag_id="f_nl", type="numbered_list", items=[cell(text) for text in texts])],
            "a bulleted list": [ListFragment(frag_id="f_bl", type="bulleted_list", items=[cell(text) for text in texts])],
            "code blocks": [CodeFragment(frag_id=f"f_c{index}", type="code_block", text=text) for index, text in enumerate(texts)],
            "table cells": [TableFragment(frag_id="f_tc", type="table", header=[cell("Header")], rows=[[cell(text)] for text in texts])],
            "table headers": [TableFragment(frag_id="f_th", type="table", header=[cell(text) for text in texts], rows=[[cell("x") for _ in texts]])],
        }

    def test_text_that_looks_like_a_template_placeholder_prints_as_typed(self) -> None:
        """A tester documents payloads such as {{7*7}} and writes words such as non-vulnerable. Neither is the
        template's own text, so neither may stop generation, and each has to print as it was typed."""
        typed = [
            "The field reflects {{7*7}} back.",
            "The host is non-vulnerable.",
            "{{table-body-cell}}",
            "{{findings}}",
            "{{{x}}}",
            "open {{ here",
            "close }} there",
            "Words like brief-explanation-here and step-1-here are plain text.",
        ]
        for kind, fragments in self._fragments_holding(typed).items():
            with self.subTest(kind=kind):
                with tempfile.TemporaryDirectory() as temporary_directory:
                    document = self._layout_document(Path(temporary_directory), [Content(type="description", fragments=fragments)])
                printed = paragraph_texts(document.element.body)
                for text in typed:
                    self.assertTrue(any(text in line for line in printed), f"{text!r} did not print as typed")
                self.assertNotRegex("".join(printed), "[\ue000-\uf8ff]", "a private marker survived into the document")

    def test_control_characters_in_text_are_dropped_and_do_not_stop_generation(self) -> None:
        """lxml refuses them, and the error that reached the tester named nothing."""
        for kind, fragments in self._fragments_holding(["a\x01b\x1bc", "d\x00e\x0bf"]).items():
            with self.subTest(kind=kind):
                with tempfile.TemporaryDirectory() as temporary_directory:
                    document = self._layout_document(Path(temporary_directory), [Content(type="description", fragments=fragments)])
                printed = paragraph_texts(document.element.body)
                for text in ("abc", "def"):
                    self.assertTrue(any(text in line for line in printed), f"{text!r} did not print")

    def test_a_leftover_template_placeholder_still_stops_generation(self) -> None:
        """The check reads the template's own text. Text a tester typed is hidden from it while the document is built."""
        leftover = Document()
        leftover.add_paragraph("For {{app-name}} on the cover")
        leftover.add_paragraph("Findings for the xxx-vuln-name here")
        self.assertEqual(_unresolved_placeholders(leftover), ["-vuln", "{{app-name}}"])

        typed = Document()
        paragraph = typed.add_paragraph("{{x}}")
        replace_component_token_runs([paragraph._p], "x", [Run(text="{{7*7}} and non-vulnerable")])
        self.assertEqual(_unresolved_placeholders(typed), [])

    def test_a_template_token_nothing_replaced_stops_a_full_render(self) -> None:
        real_metadata = _metadata
        without_app_name = lambda report: {token: value for token, value in real_metadata(report).items() if token != "app-name"}
        with patch("app.docx_report._metadata", side_effect=without_app_name), tempfile.TemporaryDirectory() as temporary_directory:
            with self.assertRaisesRegex(ReportGenerationError, r"Unresolved template placeholders: .*\{\{app-name\}\}"):
                self._layout_document(Path(temporary_directory), [])

    def test_a_table_headed_like_a_template_table_is_never_taken_for_it(self) -> None:
        """The renderer finds the Section table by its header once the findings are in, and a tester's table can be
        headed anything. Each master lacks a different template table, so all four are tried."""
        cell = lambda text: ListItem(runs=[Run(text=text)])
        headers = ("Section", "Component", "Findings", "Limitations", "User Roles", "URL(s) in Scope", "API Routes")
        for channel, segment in (("web", "JH"), ("web", "Asia"), ("thick_client", "JH"), ("mobile", "Asia")):
            for header in headers:
                with self.subTest(channel=channel, segment=segment, header=header), tempfile.TemporaryDirectory() as temporary_directory:
                    report_folder = Path(temporary_directory)
                    report = self._component_report(report_folder, channel, segment, [
                        ScopeTarget(target_id="t_one", environment="production", channel=channel, value="Acme.exe", description="Main client"),
                    ])
                    description = next(content for content in report.vulnerabilities[0].contents if content.type == "description")
                    description.fragments.append(TableFragment(frag_id="f_collide", type="table", header=[cell(header), cell("Notes")], rows=[[cell("a"), cell("b")]]))
                    rendered = Document(BytesIO(render_report_docx(report, main_template_path(report, RESOURCES), report_folder)))
                    ours = [table for table in rendered.tables if [item.text for item in table.rows[0].cells] == [header, "Notes"]]
                    self.assertEqual(len(ours), 1)
                    self.assertEqual([[item.text for item in row.cells] for row in ours[0].rows], [[header, "Notes"], ["a", "b"]])
                    if segment == "Asia":
                        carrying = [table for table in rendered.tables if any(item.text.strip() == "8.1" for row in table.rows for item in row.cells)]
                        self.assertEqual(len(carrying), 1, "the Section table lost its CVSS row")

    def _no_findings_document(self, report_folder: Path, segment: str = "JH"):
        report = self._component_report(report_folder, "web", segment, [
            ScopeTarget(target_id="t_prod", environment="production", channel="web", value="https://prod.example.test"),
        ])
        report.vulnerabilities = []
        return Document(BytesIO(render_report_docx(report, main_template_path(report, RESOURCES), report_folder)))

    def test_a_report_with_no_findings_prints_one_plain_summary_row(self) -> None:
        """The rating columns would keep the template's 10 pt token run, and a row that rates nothing
        has no rating colour to show."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            rendered = self._no_findings_document(Path(temporary_directory))
            summary = next(table for table in rendered.tables if table.cell(0, 0).text == "Findings")
            self.assertEqual(
                [[cell.text for cell in row.cells] for row in summary.rows[1:]],
                [[NO_FINDINGS_TITLE, "N/A", "N/A", "N/A", "N/A", "N/A"]],
            )
            for index, cell in enumerate(summary.rows[1].cells):
                run = next(run for run in cell.paragraphs[0].runs if run.text)
                self.assertIsNone(run.font.color.rgb, f"column {index} is coloured")
                if index in (1, 2, 3):
                    self.assertEqual(run.font.size, Pt(12), f"column {index} kept the token's size")

    def test_a_report_with_no_findings_prints_the_testing_result_in_place_of_the_findings(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            rendered = self._no_findings_document(Path(temporary_directory))
            paragraphs = [paragraph for paragraph in rendered.paragraphs if paragraph.text.strip()]
            result = next(index for index, paragraph in enumerate(paragraphs) if paragraph.text == TESTING_RESULT_PARAGRAPH)
            heading = paragraphs[result - 1]
            self.assertEqual(heading.style.style_id, "ReportHeading1")
            self.assertIsNotNone(heading._p.find(qn("w:pPr")).find(qn("w:pageBreakBefore")), "the section does not start a page")
            severity_headings = [
                paragraph.text for paragraph in paragraphs
                if paragraph.style.style_id == "ReportHeading1" and paragraph.text.endswith("Findings")
            ]
            self.assertEqual(severity_headings, [])

    def test_an_asia_report_with_no_findings_prints_one_plain_section_row(self) -> None:
        """The empty row this used to print would not import back, and its Severity cell kept the
        template's colour."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            rendered = self._no_findings_document(Path(temporary_directory), "Asia")
            table = next(table for table in rendered.tables if table.cell(0, 0).text.strip() == "Section")
            self.assertEqual([[cell.text for cell in row.cells] for row in table.rows[1:]], [["N/A"] * 5])
            self.assertEqual(list(table.rows[1]._tr.iter(qn("w:color"))), [], "an N/A cell kept a colour")
            self.assertEqual(list(table.rows[1]._tr.iter(qn("w:instrText"))), [], "an N/A cell holds a field")

    def _asia_multi_finding_document(self, report_folder: Path):
        report = self._component_report(report_folder, "thick_client", "Asia", [
            ScopeTarget(target_id="t_prod", environment="production", channel="thick_client", value="Acme.exe", description="Main client"),
        ])
        report.vulnerabilities = []
        for uid, title, severity, display_id, score, vector in (
            ("v_low", "Verbose error messages", "low", "003", "3.1", "CVSS:3.1/AV:N/AC:H/PR:H/UI:R/S:U/C:L/I:N/A:N"),
            ("v_crit", "Remote code execution", "critical", "001", "9.8", "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H"),
            ("v_crit_b", "Authentication bypass", "critical", "002", "9.1", "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:N"),
        ):
            finding = self._finding(uid, title, severity, display_id, ["t_prod"], [
                ImageFragment(frag_id=f"f_img_{uid}", type="image", environment="production", evidence_id="ev_prod", caption="Production response"),
            ])
            # Distinct per finding, so a row paired with the wrong finding fails rather than passing.
            finding.cvss_score = score
            finding.cvss_vector = vector
            report.vulnerabilities.append(finding)
        resources = RESOURCES
        return Document(BytesIO(render_report_docx(report, main_template_path(report, resources), report_folder)))

    def test_asia_section_column_references_each_findings_own_heading(self) -> None:
        """Word owns the numbering, so the reference cannot disagree with the heading it points at.
        Counting headings in Python could, silently, in a delivered report."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            rendered = self._asia_multi_finding_document(Path(temporary_directory))
            table = next(table for table in rendered.tables if table.cell(0, 0).text.strip() == "Section")

            bookmarks = [node.get(qn("w:name")) for node in rendered.element.body.iter(qn("w:bookmarkStart"))]
            referenced = []
            for row in table.rows[1:]:
                instructions = [node.text for node in row.cells[0]._tc.iter(qn("w:instrText"))]
                self.assertEqual(len(instructions), 1, "one field per Section cell, or Word shows the wrong one")
                name = instructions[0].split()[1]
                self.assertEqual(instructions[0], f" REF {name} \\w \\h ")
                referenced.append(name)

            # A duplicate name resolves to whichever bookmark Word finds first, without complaining.
            self.assertEqual(len(set(referenced)), len(referenced))
            for name in referenced:
                self.assertEqual(bookmarks.count(name), 1)
            bookmarked_styles = {
                paragraph.style.style_id
                for paragraph in rendered.paragraphs
                if any(node.get(qn("w:name")) in referenced for node in paragraph._p.iter(qn("w:bookmarkStart")))
            }
            self.assertEqual(bookmarked_styles, {"ReportHeading2"})
            self.assertNotIn("{{section-number}}", "\n".join(cell.text for row in table.rows for cell in row.cells))

    def test_a_computed_section_number_loses_its_trailing_period(self) -> None:
        """Word reports the heading's list label, "7.1. ", period and all. Trimming the field result
        would not survive the refresh settings.xml asks for, so a computed field becomes plain text.
        An uncomputed one is left alone, or a render without Word would deliver an empty cell."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            rendered = self._asia_multi_finding_document(Path(temporary_directory))
            table = next(table for table in rendered.tables if table.cell(0, 0).text.strip() == "Section")
            cell = table.rows[1].cells[0]
            # Stands in for Word, which is the only thing that can compute the number.
            end = next(
                node.getparent() for node in cell._tc.iter(qn("w:fldChar"))
                if node.get(qn("w:fldCharType")) == "end"
            )
            computed = OxmlElement("w:r")
            written = OxmlElement("w:t")
            written.text = "7.1. "
            computed.append(written)
            end.addprevious(computed)

            self.assertEqual(flatten_section_number_fields(rendered), 1)
            self.assertEqual(cell.text.strip(), "7.1")
            self.assertEqual(len(list(cell._tc.iter(qn("w:fldChar")))), 0, "the field outlived the flattening")
            untouched = [len(list(row.cells[0]._tc.iter(qn("w:instrText")))) for row in table.rows[2:]]
            self.assertEqual(untouched, [1] * len(untouched), "a field Word had not computed was destroyed")

    def test_asia_section_rows_follow_the_rendered_finding_order(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            rendered = self._asia_multi_finding_document(Path(temporary_directory))
            table = next(table for table in rendered.tables if table.cell(0, 0).text.strip() == "Section")

            # ReportHeading2 also carries static sections like "Findings Summary", so the bookmarks
            # are what identify a finding title. They are also what a Section cell can point at.
            headings = [
                paragraph.text
                for paragraph in rendered.paragraphs
                if any(node.get(qn("w:name")).startswith("vuln_") for node in paragraph._p.iter(qn("w:bookmarkStart")))
            ]
            self.assertEqual([row.cells[1].text for row in table.rows[1:]], headings)
            # Severity order first, then title. The body and this table must not sort independently.
            self.assertEqual(headings, ["Authentication bypass", "Remote code execution", "Verbose error messages"])
            self.assertEqual([row.cells[2].text for row in table.rows[1:]], ["Critical", "Critical", "Low"])
            self.assertEqual([row.cells[3].text for row in table.rows[1:]], ["9.1", "9.8", "3.1"])
            self.assertEqual(
                [row.cells[4].text for row in table.rows[1:]],
                [
                    "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:N",
                    "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H",
                    "CVSS:3.1/AV:N/AC:H/PR:H/UI:R/S:U/C:L/I:N/A:N",
                ],
            )

    def test_asia_section_references_are_marked_for_word_to_compute(self) -> None:
        """A field left clean renders blank until someone presses F9, which nobody does."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            report_folder = Path(temporary_directory)
            report = self._component_report(report_folder, "thick_client", "Asia", [
                ScopeTarget(target_id="t_prod", environment="production", channel="thick_client", value="Acme.exe", description="Main client"),
            ])
            resources = RESOURCES
            contents = render_report_docx(report, main_template_path(report, resources), report_folder)
            rendered = Document(BytesIO(contents))

            table = next(table for table in rendered.tables if table.cell(0, 0).text.strip() == "Section")
            starts = [node for row in table.rows[1:] for node in row.cells[0]._tc.iter(qn("w:fldChar")) if node.get(qn("w:fldCharType")) == "begin"]
            self.assertTrue(starts)
            self.assertTrue(all(node.get(qn("w:dirty")) == "true" for node in starts))
            with ZipFile(BytesIO(contents)) as archive:
                self.assertIn("w:updateFields", archive.read("word/settings.xml").decode("utf-8"))

    @staticmethod
    def _finding(
        uid: str,
        title: str,
        severity: str,
        display_id: str,
        target_ids: list[str],
        images: list[ImageFragment],
        *,
        likelihood: str = "high",
        impact: str = "high",
    ) -> Vulnerability:
        return Vulnerability(
            uid=uid,
            display_id=display_id,
            title=title,
            likelihood=likelihood,
            impact=impact,
            severity=severity,
            status="open_new",
            scope=Scope(mode="custom", target_ids=target_ids),
            contents=[
                Content(type="description", fragments=[ParagraphFragment(frag_id=f"{uid}_description", type="paragraph", runs=[Run(text="A complete "), Run(text="description.", bold=True)]), NoteFragment(frag_id=f"{uid}_note", type="note", runs=[Run(text="Validate the result independently.")])]),
                Content(type="recommended_remediation", fragments=[ListFragment(frag_id=f"{uid}_remediation", type="bulleted_list", items=[ListItem(runs=[Run(text="Apply authorization controls.")])])]),
                Content(type="proof_of_concept", fragments=[ListFragment(frag_id=f"{uid}_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")]), ListItem(runs=[Run(text="Observe the response.")])]), CodeFragment(frag_id=f"{uid}_code", type="code_block", caption="Request", text="GET /accounts/123"), TableFragment(frag_id=f"{uid}_table", type="table", caption="Generated table", header=[ListItem(runs=[Run(text="Generated Header")]), ListItem(runs=[Run(text="Result")])], rows=[[ListItem(runs=[Run(text="Request")]), ListItem(runs=[Run(text="Allowed")])]]), *images]),
            ],
        )


class ImageBorderTests(unittest.TestCase):
    """The border is part of the picture's pixels, so a picture that came out of an earlier report
    already has it and must not get a second."""

    TEMPLATE = RESOURCES / "fragments" / "image_fragment.docx"

    def _embedded(self, folder: Path, image: Image.Image) -> Image.Image:
        """The picture the generator embeds for this image."""
        path = folder / "source.png"
        image.save(path)
        return Image.open(_bordered_image_stream(path, 155, self.TEMPLATE)).convert("RGB")

    def test_an_image_that_already_has_a_border_gets_no_second_one(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            screenshot = Image.new("RGB", (1920, 1080), "white")
            once = self._embedded(folder, screenshot)
            twice = self._embedded(folder, once)

            self.assertGreater(once.width, screenshot.width, "a plain screenshot must still get a border")
            self.assertEqual(twice.size, once.size, "a second border was drawn around the first")
            self.assertTrue(twice.tobytes() == once.tobytes(), "the bordered picture was changed")
            self.assertEqual(off_border_pixels(twice), [])

    def test_a_border_shrunk_and_blurred_by_a_copy_is_still_found(self) -> None:
        """Copying a picture out of Word may hand back a smaller, softer one."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            bordered = self._embedded(folder, Image.new("RGB", (1920, 1080), "white"))
            for scale in (1.0, 0.75, 0.5, 0.4):
                with self.subTest(scale=scale):
                    smaller = bordered.resize((round(bordered.width * scale), round(bordered.height * scale)), Image.Resampling.LANCZOS)
                    smaller = smaller.filter(ImageFilter.GaussianBlur(0.6))

                    self.assertEqual(self._embedded(folder, smaller).size, smaller.size)

    def test_only_a_dark_edge_all_the_way_round_counts_as_a_border(self) -> None:
        white = Image.new("RGB", (1920, 1080), "white")
        dark_without_a_left_edge = Image.new("RGB", (1920, 1080), (30, 30, 30))
        dark_without_a_left_edge.paste((255, 255, 255), (0, 0, 5, 1080))
        cases = (
            ("dark-mode screenshot", Image.new("RGB", (1920, 1080), (30, 30, 30)), False),
            ("black frame", ImageOps.expand(white, border=3, fill="black"), False),
            ("blue frame", ImageOps.expand(white, border=5, fill=(0, 0, 255)), True),
            ("grey frame", ImageOps.expand(white, border=6, fill=(128, 128, 128)), True),
            ("dark on three sides only", dark_without_a_left_edge, True),
        )
        with tempfile.TemporaryDirectory() as temporary_directory:
            for label, image, gets_a_border in cases:
                with self.subTest(label):
                    self.assertEqual(self._embedded(Path(temporary_directory), image).size != image.size, gets_a_border)


if __name__ == "__main__":
    unittest.main()