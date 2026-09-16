from __future__ import annotations

import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm
from PIL import Image

from app.docx_captions import postprocess_image_captions


class DocxCaptionTests(unittest.TestCase):
    def test_postprocesses_three_images_into_native_word_captions(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            source = folder / "source.docx"
            output = folder / "captioned.docx"
            second_output = folder / "captioned-again.docx"
            document = Document()
            captions = [
                "Authentication response.",
                "Figure 99: Authorization response.",
                "Remediation verification.",
            ]
            for index, caption_text in enumerate(captions, 1):
                image = BytesIO()
                Image.new("RGB", (320, 180), (40 * index, 80, 120)).save(image, format="PNG")
                image.seek(0)
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.add_run().add_picture(image, width=Mm(80))
                caption = document.add_paragraph(caption_text, style="Caption")
                caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
            document.save(source)

            destination, converted = postprocess_image_captions(source, output)
            self.assertEqual(destination, output)
            self.assertEqual(converted, 3)
            with ZipFile(output) as package:
                self.assertIsNone(package.testzip())
                self.assertEqual(len([name for name in package.namelist() if name.startswith("word/media/")]), 3)

            rendered = Document(output)
            self.assertEqual(len(rendered.inline_shapes), 3)
            caption_paragraphs = [
                paragraph
                for paragraph in rendered.paragraphs
                if any("SEQ Figure" in (node.text or "") for node in paragraph._p.iter(qn("w:instrText")))
            ]
            self.assertEqual(
                [paragraph.text for paragraph in caption_paragraphs],
                [
                    "Figure 1. Authentication response.",
                    "Figure 2. Authorization response.",
                    "Figure 3. Remediation verification.",
                ],
            )
            for paragraph in caption_paragraphs:
                self.assertEqual(paragraph.style.name, "Caption")
                self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.CENTER)
                self.assertTrue(paragraph._p.getprevious().xpath(".//w:drawing"))
                instructions = [node.text for node in paragraph._p.iter(qn("w:instrText"))]
                self.assertEqual(instructions, [r" SEQ Figure \* ARABIC "])
                field_types = [
                    node.get(qn("w:fldCharType"))
                    for node in paragraph._p.iter(qn("w:fldChar"))
                ]
                self.assertEqual(field_types, ["begin", "separate", "end"])
            update = rendered.settings._element.find(qn("w:updateFields"))
            self.assertIsNotNone(update)
            self.assertEqual(update.get(qn("w:val")), "true")

            _, converted_again = postprocess_image_captions(output, second_output)
            self.assertEqual(converted_again, 0)
            reopened = Document(second_output)
            self.assertEqual(len(list(reopened.element.body.iter(qn("w:instrText")))), 3)
            self.assertTrue(
                all(
                    paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
                    for paragraph in reopened.paragraphs
                    if any("SEQ Figure" in (node.text or "") for node in paragraph._p.iter(qn("w:instrText")))
                )
            )


if __name__ == "__main__":
    unittest.main()