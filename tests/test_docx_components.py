from __future__ import annotations

import re
import threading
import unittest
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from lxml import etree

from app.docx_components import clone_component_elements, replace_component_token, replace_component_token_runs, replace_pattern_across_text_nodes, restore_user_text, user_text
from app.models import Run
from tests.support import numbering_details

ROOT = Path(__file__).resolve().parent.parent
FRAGMENTS = ROOT / "resources" / "fragments"


class DocxComponentTests(unittest.TestCase):
    RATING_FONT_COLORS = {
        "Critical": RGBColor(192, 0, 0),
        "High": RGBColor(237, 125, 49),
        "Medium": RGBColor(255, 192, 0),
        "Low": RGBColor(112, 173, 71),
        "Informational": RGBColor(81, 81, 81),
    }

    def test_preserves_tag_styles_and_changes_only_rating_colors(self) -> None:
        document = Document()
        tags = (
            ("ordinary-tag", "Ordinary value", RGBColor(12, 34, 56)),
            ("likelihood-rating", "Critical", self.RATING_FONT_COLORS["Critical"]),
            ("impact-rating", "Low", self.RATING_FONT_COLORS["Low"]),
            ("severity-rating", "High", self.RATING_FONT_COLORS["High"]),
        )
        for token, _, _ in tags:
            paragraph = document.add_paragraph(style="Quote")
            run = paragraph.add_run(f"{{{{{token}}}}}")
            run.bold = True
            run.italic = True
            run.underline = True
            run.font.name = "Courier New"
            run.font.size = Pt(13)
            run.font.small_caps = True
            run.font.color.rgb = RGBColor(12, 34, 56)
        paragraph_properties = [paragraph._p.pPr.xml for paragraph in document.paragraphs]
        run_properties = [self._run_properties_without_color(paragraph.runs[0]) for paragraph in document.paragraphs]

        for token, value, _ in tags:
            replace_component_token([document.element.body], token, value)
        for index, (_, value, expected_color) in enumerate(tags):
            with self.subTest(value=value):
                paragraph = document.paragraphs[index]
                run = paragraph.runs[0]
                self.assertEqual(paragraph.text, value)
                self.assertEqual(paragraph._p.pPr.xml, paragraph_properties[index])
                self.assertEqual(self._run_properties_without_color(run), run_properties[index])
                self.assertEqual(run.font.color.rgb, expected_color)

    def test_applies_all_rating_font_colors(self) -> None:
        for rating, expected_color in self.RATING_FONT_COLORS.items():
            with self.subTest(rating=rating):
                document = Document()
                document.add_paragraph("{{rating}}")
                replace_component_token([document.element.body], "rating", rating)
                self.assertEqual(document.paragraphs[0].runs[0].font.color.rgb, expected_color)

    def test_a_cloned_fragment_keeps_its_formatting_and_list_definition(self) -> None:
        for filename in ("paragraph_fragment.docx", "numbered_fragment.docx", "bulleted_fragment.docx"):
            with self.subTest(filename=filename):
                source = Document(FRAGMENTS / filename)
                target = Document(ROOT / "resources" / "MAIN.docx")
                cloned = next(
                    element for element in clone_component_elements(target, FRAGMENTS / filename)
                    if element.tag == qn("w:p")
                )
                self.assertEqual(self._paragraph_properties(cloned), self._paragraph_properties(source.paragraphs[0]._p))
                self.assertEqual(self._run_formatting(cloned), self._run_formatting(source.paragraphs[0]._p))

    def test_cloned_lists_share_a_numbering_only_when_given_the_same_map(self) -> None:
        target = Document(ROOT / "resources" / "MAIN.docx")

        def numbered(numbering_ids=None):
            elements = clone_component_elements(target, FRAGMENTS / "numbered_fragment.docx", numbering_ids=numbering_ids)
            paragraph = next(element for element in elements if element.tag == qn("w:p"))
            target.element.body.sectPr.addprevious(paragraph)
            return numbering_details(target, target.paragraphs[-1])

        shared: dict = {}
        first, second = numbered(shared), numbered(shared)
        self.assertEqual(first[:2], second[:2])
        self.assertEqual((first[2], second[2]), (1, 1))
        separate = numbered()
        self.assertNotEqual(separate[0], first[0])
        self.assertNotEqual(separate[1], first[1])

    @staticmethod
    def _finishes(call, seconds: float = 5.0) -> bool:
        """Whether call() returns in time, so a loop that never ends fails the test instead of hanging the suite."""
        worker = threading.Thread(target=call, daemon=True)
        worker.start()
        worker.join(seconds)
        return not worker.is_alive()

    def test_a_value_that_spells_its_own_token_is_written_once_and_the_replacement_ends(self) -> None:
        """The loop looked for the next token from the start of the paragraph, so text it had just written that
        spelled the token was found and replaced again, for ever."""
        runs = lambda value: lambda element: replace_component_token_runs([element], "x", [Run(text=value)])
        plain = lambda value: lambda element: replace_component_token([element], "x", value)
        bare_word = lambda value: lambda element: replace_pattern_across_text_nodes(element, re.compile(r"(?<![\w-])name(?![\w-])"), value)
        cases = [
            ("rich runs", "Before {{x}} after {{ x }}.", runs("a {{x}} b"), "Before a {{x}} b after a {{x}} b."),
            ("exactly the token", "{{x}}", runs("{{x}}"), "{{x}}"),
            ("one string", "Before {{x}} after {{ x }}.", plain("a {{x}} b"), "Before a {{x}} b after a {{x}} b."),
            ("a bare word", "Hello name and name.", bare_word("my name"), "Hello my name and my name."),
        ]
        for name, template, replace, expected in cases:
            with self.subTest(name):
                paragraph = Document().add_paragraph(template)
                self.assertTrue(self._finishes(lambda: replace(paragraph._p)), "the replacement re-matched the text it had just written")
                self.assertEqual(restore_user_text(paragraph.text), expected)

    def test_marks_keep_braces_and_markers_out_of_the_check_and_come_out_exactly(self) -> None:
        """What a tester typed goes in with a mark inside each pair of braces and each unfinished-template marker.
        Taking the marks out has to give back exactly what was typed, even when the tester typed a mark."""
        for typed in ("{{7*7}}", "}}}", "{{{x}}}", "non-vulnerable", "Brief-Explanation-Here", "a\ue000b", "\ue000m", "\ue000s{{", "plain"):
            with self.subTest(typed=typed):
                hidden = user_text(typed)
                self.assertNotIn("{{", hidden)
                self.assertNotIn("}}", hidden)
                self.assertNotIn("-vuln", hidden.casefold())
                self.assertEqual(restore_user_text(hidden), typed)
        self.assertEqual(user_text("plain text, one { and one }"), "plain text, one { and one }", "text with nothing to hide was changed")

    def test_a_token_split_across_runs_is_still_replaced(self) -> None:
        """The loop now carries on from where it wrote, so a token the template split over runs, and one that
        follows a replacement in the same paragraph, must both still be found."""
        paragraph = Document().add_paragraph()
        for piece in ("A {{", "x}} then {", "{x}} end"):
            paragraph.add_run(piece)
        replace_component_token_runs([paragraph._p], "x", [Run(text="one"), Run(text="two", bold=True)])
        self.assertEqual(paragraph.text, "A onetwo then onetwo end")

    def test_characters_xml_cannot_hold_are_dropped_where_text_is_written(self) -> None:
        """lxml refuses these, which made Generate fail for text a tester had pasted, and nothing named it."""
        typed = "a\x01b\x1bc\x00d\ud800e\ufffff\x0bg"
        runs = lambda element: replace_component_token_runs([element], "x", [Run(text=typed)])
        plain = lambda element: replace_component_token([element], "x", typed)
        for name, replace in (("rich runs", runs), ("one string", plain)):
            with self.subTest(name):
                paragraph = Document().add_paragraph("{{x}}")
                replace(paragraph._p)
                self.assertEqual(paragraph.text, "abcdefg")

    def test_a_line_break_a_tab_and_a_carriage_return_survive_the_character_filter(self) -> None:
        paragraph = Document().add_paragraph("{{x}}")
        replace_component_token_runs([paragraph._p], "x", [Run(text="one\r\ntwo\rthree\tfour\nfive")])
        self.assertEqual(paragraph.text, "one\ntwo\nthree\tfour\nfive")

    @staticmethod
    def _run_formatting(paragraph) -> list[str | None]:
        return [
            etree.tostring(run.find(qn("w:rPr")), method="c14n", exclusive=True).decode() if run.find(qn("w:rPr")) is not None else None
            for run in paragraph.iter(qn("w:r"))
        ]

    @staticmethod
    def _paragraph_properties(paragraph) -> str:
        properties = deepcopy(paragraph.find(qn("w:pPr")))
        for numbering_id in properties.iter(qn("w:numId")):
            numbering_id.set(qn("w:val"), "0")
        return etree.tostring(properties, method="c14n", exclusive=True).decode()

    @staticmethod
    def _run_properties_without_color(run) -> str | None:
        properties = deepcopy(run._r.rPr)
        if properties is None:
            return None
        color = properties.find(qn("w:color"))
        if color is not None:
            properties.remove(color)
        return etree.tostring(properties, method="c14n", exclusive=True).decode()


if __name__ == "__main__":
    unittest.main()
