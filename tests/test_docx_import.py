"""Reading a generated report back in.

Every test here renders with ``render_report_docx`` directly. The generate *routes* call Word
through ``update_docx_bytes_with_word`` and return 422 on any machine without it, so going through
them would test the environment rather than the document.
"""
from __future__ import annotations

import copy
import functools
import hashlib
import json
import tempfile
import unittest
import uuid
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from threading import Event
from typing import get_args
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from fastapi.testclient import TestClient

from app import main
from app.docx_import import NETWORK_IMPORT_WARNING, SEVERITY_TICKET_LABEL, ReportImportError, ReportImportLimitError, _parsed_ticket_lines, classify_paragraph, first_change, numbering_formats, parse_report_docx
from app.docx_report import generation_issues, main_template_path, render_report_docx
from app.report_service import RESOLVED_REMEDIATION, character_issue, finding_input_issues, setup_issues
from tests.support import png_bytes, use_temp_workspace
from app.models import (
    CodeFragment,
    Content,
    Engagement,
    EvidenceItem,
    ImageFragment,
    InstanceTitleFragment,
    ListFragment,
    ListItem,
    NoteFragment,
    ParagraphFragment,
    Report,
    Run,
    Scope,
    ScopeTarget,
    Segment,
    TableFragment,
    TestAccount,
    TestWindow,
    Vulnerability,
)

TEMPLATE = Path(__file__).resolve().parent.parent / "resources" / "MAIN.docx"


@functools.cache
def base_render() -> bytes:
    """The shared fixture, rendered once. Tests reopen it from these bytes, so none can alter it."""
    with tempfile.TemporaryDirectory() as temporary_directory:
        folder = Path(temporary_directory)
        return render_report_docx(FragmentRecognitionTests._report(folder), TEMPLATE, folder)


def retest_shaped(finding: Vulnerability, conclusion: list[Run], status: str = "open_previously_discovered") -> None:
    """Give a finding the five sections a previously discovered one prints, ending in ``conclusion``."""
    finding.status = status
    finding.contents.extend([
        Content(type="previous_proof_of_concept", fragments=[
            ListFragment(frag_id="f_prev", type="numbered_list", items=[ListItem(runs=[Run(text="Last year's step.")])]),
        ]),
        Content(type="in_conclusion", fragments=[ParagraphFragment(frag_id="f_conc", type="paragraph", runs=conclusion)]),
    ])


def conclusion_render(runs: list[Run], status: str = "open_previously_discovered", title: str | None = None) -> bytes:
    """The shared finding as a retest-shaped finding whose In Conclusion section holds ``runs``."""
    with tempfile.TemporaryDirectory() as temporary_directory:
        folder = Path(temporary_directory)
        report = FragmentRecognitionTests._report(folder)
        if title:
            report.vulnerabilities[0].title = title
        retest_shaped(report.vulnerabilities[0], runs, status)
        return render_report_docx(report, TEMPLATE, folder)


def _twin(finding: Vulnerability, uid: str, display_id: str, title: str) -> Vulnerability:
    """A copy of ``finding`` under another identity, with its fragment IDs made unique."""
    twin = copy.deepcopy(finding)
    twin.uid, twin.display_id, twin.title = uid, display_id, title
    for content in twin.contents:
        for fragment in content.fragments:
            fragment.frag_id += f"_{uid}"
    return twin


@functools.cache
def two_findings_render() -> bytes:
    """Two previously discovered findings in different severity groups, so each heading sits under its
    own group heading and one can be damaged without the other."""
    with tempfile.TemporaryDirectory() as temporary_directory:
        folder = Path(temporary_directory)
        report = FragmentRecognitionTests._report(folder)
        first = report.vulnerabilities[0]
        retest_shaped(first, [Run(text="Still reachable on retest.")])
        second = _twin(first, "v_second", "002", "Second finding")
        second.severity = "medium"
        report.vulnerabilities.append(second)
        return render_report_docx(report, TEMPLATE, folder)


@functools.cache
def lettered_findings_render(letters: str) -> bytes:
    """The shared finding plus one more for each letter, titled "Finding B" and so on, all High and Open (New)."""
    with tempfile.TemporaryDirectory() as temporary_directory:
        folder = Path(temporary_directory)
        report = FragmentRecognitionTests._report(folder)
        first = report.vulnerabilities[0]
        for number, letter in enumerate(letters, start=2):
            report.vulnerabilities.append(_twin(first, f"v_{letter}", f"{number:03d}", f"Finding {letter}"))
        return render_report_docx(report, TEMPLATE, folder)


@functools.cache
def asia_findings_render() -> bytes:
    """Two findings on the Asia template, the one that has a CVSS table."""
    with tempfile.TemporaryDirectory() as temporary_directory:
        folder = Path(temporary_directory)
        report = FragmentRecognitionTests._report(folder)
        report.engagement.segment = "Asia"
        first = report.vulnerabilities[0]
        first.cvss_score, first.cvss_vector = "8.1", "CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:U/C:H/I:H/A:N"
        second = _twin(first, "v_second", "002", "Second finding")
        second.cvss_score, second.cvss_vector = "3.1", "CVSS:3.1/AV:N/AC:H/PR:H/UI:R/S:U/C:L/I:N/A:N"
        report.vulnerabilities.append(second)
        return render_report_docx(report, main_template_path(report, Path(__file__).resolve().parent.parent / "resources"), folder)


def _heading(document, text: str):
    """The first body paragraph whose whole text is ``text``."""
    return next(paragraph for paragraph in document.paragraphs if " ".join(paragraph.text.split()) == text)


def _set_text(paragraph, text: str) -> None:
    for run in paragraph.runs[1:]:
        run.text = ""
    paragraph.runs[0].text = text


def _table_starting(document, header: str):
    return next(table for table in document.tables if table.rows[0].cells[0].text.strip() == header)


def _table_after(document, heading: str):
    """The first table after the paragraph whose whole text is ``heading``."""
    element = _heading(document, heading)._p.getnext()
    while element.tag != qn("w:tbl"):
        element = element.getnext()
    return next(table for table in document.tables if table._tbl is element)


def _edited(data: bytes, change) -> bytes:
    document = Document(BytesIO(data))
    change(document)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class FragmentRecognitionTests(unittest.TestCase):
    """The importer recognises a fragment by the component that produced it, so these pin the
    signals that survive generation. A template edit that erases one fails here rather than
    silently mislabelling every fragment in every future import."""

    @staticmethod
    def _report(folder: Path) -> Report:
        (folder / "evidence").mkdir()
        buffer = png_bytes(40, 20)
        (folder / "evidence" / "ev_shot.png").write_bytes(buffer)
        now = datetime.now().astimezone()
        report = Report(
            report_id="r_import", app_id="CI-IMPORT", saved_at=now,
            engagement=Engagement(
                app_name="Northstar Banking", ci_number="CI-IMPORT", segment="JH",
                report_type="annual_pentest", report_date=date(2026, 9, 9), tester="QA Tester", network="Internal",
                tested_environments=["production"], tested_channels=["web"],
                test_windows={"production": TestWindow(start_date=date(2026, 8, 1), end_date=date(2026, 8, 2))},
            ),
            scope_targets=[ScopeTarget(target_id="t_web", environment="production", channel="web", value="https://prod.example.test")],
            evidence={"ev_shot": EvidenceItem(file="evidence/ev_shot.png", original_name="shot.png", width_px=40, height_px=20, sha256="0" * 64, uploaded_at=now)},
        )
        report.vulnerabilities = [Vulnerability(
            uid="v_all", display_id="001", title="Authorization bypass",
            likelihood="high", impact="high", severity="high", status="open_new",
            scope=Scope(mode="custom", target_ids=["t_web"]),
            contents=[
                Content(type="description", fragments=[
                    ParagraphFragment(frag_id="f_para", type="paragraph", runs=[Run(text="A justified paragraph.")]),
                    NoteFragment(frag_id="f_note", type="note", runs=[Run(text="Validate independently.")]),
                    ListFragment(frag_id="f_bullet", type="bulleted_list", items=[ListItem(runs=[Run(text="A bullet.")])]),
                ]),
                Content(type="recommended_remediation", fragments=[
                    ParagraphFragment(frag_id="f_rem", type="paragraph", runs=[Run(text="Apply controls.")]),
                ]),
                Content(type="proof_of_concept", fragments=[
                    ListFragment(frag_id="f_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])]),
                    InstanceTitleFragment(frag_id="f_title", type="instance_title", text="First instance"),
                    CodeFragment(frag_id="f_code", type="code_block", caption="Request", text="GET /accounts/123"),
                    TableFragment(frag_id="f_table", type="table", caption="Results",
                                  header=[ListItem(runs=[Run(text="Header")])],
                                  rows=[[ListItem(runs=[Run(text="Cell")])]]),
                    ImageFragment(frag_id="f_image", type="image", environment="production",
                                  evidence_id="ev_shot", caption="Production response"),
                ]),
            ],
        )]
        return report

    def _rendered(self):
        return Document(BytesIO(base_render()))

    def _retest_report(self, folder: Path, tickets: str = "") -> Report:
        """Generation-clean and previously discovered, so it renders ``retest_finding.docx``.

        The shared fixture is ``open_new``, which renders ``new_finding.docx`` and carries no ticket
        paragraph at all, so nothing already here can reach this."""
        report = self._report(folder)
        finding = report.vulnerabilities[0]
        finding.status = "open_previously_discovered"
        finding.severity_review_tickets = tickets
        finding.contents.append(Content(type="previous_proof_of_concept", fragments=[
            ListFragment(frag_id="f_prev", type="numbered_list", items=[ListItem(runs=[Run(text="Last year's step.")])]),
        ]))
        finding.contents.append(Content(type="in_conclusion", fragments=[
            ParagraphFragment(frag_id="f_conc", type="paragraph", runs=[Run(text="Still reachable on retest.")]),
        ]))
        return report

    @staticmethod
    def _normalized_retest_projection(payload: dict, evidence: dict[str, bytes], summary: dict) -> dict:
        report = Report.model_validate(payload)
        target_indexes = {target.target_id: index for index, target in enumerate(report.scope_targets)}

        def fragment_value(fragment) -> dict:
            value = fragment.model_dump(mode="json", exclude={"frag_id", "evidence_id"}, exclude_none=True)
            return value

        return {
            "engagement": report.engagement.model_dump(mode="json"),
            "targets": [
                target.model_dump(mode="json", exclude={"target_id"})
                for target in report.scope_targets
            ],
            "findings": [{
                "display_id": finding.display_id,
                "title": finding.title,
                "likelihood": finding.likelihood,
                "impact": finding.impact,
                "severity": finding.severity,
                "status": finding.status,
                "scope_targets": [target_indexes[target_id] for target_id in finding.scope.target_ids],
                "tickets": finding.severity_review_tickets,
                "cvss_score": finding.cvss_score,
                "cvss_vector": finding.cvss_vector,
                "contents": [
                    {"type": content.type, "fragments": [fragment_value(fragment) for fragment in content.fragments]}
                    for content in finding.contents
                ],
            } for finding in report.vulnerabilities],
            "evidence": sorted((hashlib.sha256(data).hexdigest(), len(data)) for data in evidence.values()),
            "summary": summary,
        }

    def test_both_import_modes_read_network_access_from_the_cover(self) -> None:
        """Retest is the default mode, so reading it only on the editable path would leave the
        commoner one unset."""
        for mode in ("retest", "editable"):
            for network in ("Internal", "External"):
                with self.subTest(mode=mode, network=network), tempfile.TemporaryDirectory() as temporary_directory:
                    folder = Path(temporary_directory)
                    report = self._report(folder)
                    report.engagement.network = network

                    payload, _evidence, summary = parse_report_docx(render_report_docx(report, TEMPLATE, folder), mode=mode)

                    self.assertEqual(Report.model_validate(payload).engagement.network, network)
                    self.assertNotIn(NETWORK_IMPORT_WARNING, summary.get("warnings", []))
                    if mode == "editable":
                        self.assertIn("network access", summary["restored_engagement"])

    def test_a_cover_without_a_network_access_imports_unset_and_says_so(self) -> None:
        """Generation refuses an unset value, so N/A here means a hand-edited cover; it is not a
        choice to restore."""
        document = Document(BytesIO(base_render()))
        for node in document.element.body.iter(qn("w:t")):
            if node.text == "Internal" and any(ancestor.tag == qn("w:txbxContent") for ancestor in node.iterancestors()):
                node.text = "N/A"
        buffer = BytesIO()
        document.save(buffer)
        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, _evidence, summary = parse_report_docx(buffer.getvalue(), mode=mode)

                self.assertIsNone(Report.model_validate(payload).engagement.network)
                self.assertIn(NETWORK_IMPORT_WARNING, summary["warnings"])
                self.assertIn("Setup", NETWORK_IMPORT_WARNING)
                self.assertNotIn("network access", summary.get("restored_engagement", []))

    def test_severity_review_tickets_survive_a_round_trip(self) -> None:
        """Rendered with the prefix, read back without it. Asserting either half alone would pass
        while the two disagreed, which is the only failure worth catching here."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._retest_report(folder, tickets="1234\n5678\n90123")
            self.assertEqual(generation_issues(report), [])
            document = render_report_docx(report, TEMPLATE, folder)
            self.assertIn("GRIMPEN-1234", "\n".join(item.text for item in Document(BytesIO(document)).paragraphs))

            payload, _evidence, _summary = parse_report_docx(document)
            self.assertEqual(payload["vulnerabilities"][0]["severity_review_tickets"], "1234\n5678\n90123")

    def test_retest_default_preserves_the_retest_projection_except_named_neutral_fixes(self) -> None:
        """The proof of concept the document carries is the latest one, so it becomes the draft's
        *previous* proof of concept and the retest starts with an empty one."""
        data = base_render()

        payload, evidence, summary = parse_report_docx(data)
        projection = self._normalized_retest_projection(payload, evidence, summary)

        self.assertEqual(projection["engagement"], {
            "app_name": "Northstar Banking", "app_owner": "", "ci_number": "", "bsn_number": "",
            "segment": "JH", "report_type": None, "report_date": None, "tester": "",
            "network": "Internal",
            "tested_environments": ["production"], "tested_channels": ["web"],
            "non_production_label": "NON-PROD", "start_date": None, "end_date": None,
            "test_windows": {}, "test_accounts": [{"user_role": "N/A", "username": "N/A"}],
            "limitations": "N/A", "classification": "Confidential", "template_set": "default-v1",
        })
        self.assertEqual(projection["targets"], [{
            "environment": "production", "channel": "web",
            "value": "https://prod.example.test", "description": "", "order": 0,
        }])
        finding = projection["findings"][0]
        self.assertEqual(
            {key: finding[key] for key in ("display_id", "title", "likelihood", "impact", "severity", "status", "scope_targets")},
            {"display_id": "001", "title": "Authorization bypass", "likelihood": "high", "impact": "high",
             "severity": "high", "status": "open_previously_discovered", "scope_targets": [0]},
        )
        self.assertEqual([content["type"] for content in finding["contents"]], [
            "description", "recommended_remediation", "previous_proof_of_concept", "proof_of_concept", "in_conclusion",
        ])
        self.assertEqual(
            [fragment["type"] for fragment in finding["contents"][2]["fragments"]],
            ["numbered_list", "instance_title", "code_block", "table", "image"],
        )
        self.assertEqual([fragment["type"] for fragment in finding["contents"][3]["fragments"]], ["numbered_list", "image"])
        self.assertEqual([fragment["type"] for fragment in finding["contents"][0]["fragments"]], ["paragraph", "note", "bulleted_list"])
        self.assertEqual(len(projection["evidence"]), 1)
        record = payload["evidence"][next(iter(evidence))]
        self.assertEqual(record["sha256"], hashlib.sha256(next(iter(evidence.values()))).hexdigest())
        self.assertEqual(projection["summary"], {
            "retained": 1, "dropped_resolved": [], "statuses_rewritten": ["Authorization bypass"],
        })
        Report.model_validate(payload)

    def test_retest_mode_and_omitted_mode_are_equivalent(self) -> None:
        data = base_render()

        implicit = parse_report_docx(data)
        explicit = parse_report_docx(data, mode="retest")

        self.assertEqual(
            self._normalized_retest_projection(*implicit),
            self._normalized_retest_projection(*explicit),
        )

    def test_duplicate_title_asia_rows_pair_by_occurrence(self) -> None:
        resources = Path(__file__).resolve().parent.parent / "resources"
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            report.engagement.segment = "Asia"
            first = report.vulnerabilities[0]
            first.cvss_score = "3.1"
            first.cvss_vector = "CVSS:3.1/AV:N/AC:H/PR:H/UI:R/S:U/C:L/I:N/A:N"
            second = copy.deepcopy(first)
            second.uid = "v_duplicate"
            second.display_id = "002"
            second.cvss_score = "8.1"
            second.cvss_vector = "CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:U/C:H/I:H/A:N"
            for content in second.contents:
                for fragment in content.fragments:
                    fragment.frag_id += "_duplicate"
            report.vulnerabilities = [first, second]

            payload, _evidence, _summary = parse_report_docx(
                render_report_docx(report, main_template_path(report, resources), folder),
            )

            self.assertEqual(
                [(finding["display_id"], finding["cvss_score"], finding["cvss_vector"]) for finding in payload["vulnerabilities"]],
                [("001", first.cvss_score, first.cvss_vector), ("002", second.cvss_score, second.cvss_vector)],
            )

    def test_an_empty_ticket_value_prints_n_a_and_imports_as_nothing(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            document = render_report_docx(self._retest_report(folder), TEMPLATE, folder)
            payload, _evidence, _summary = parse_report_docx(document)
            self.assertEqual(payload["vulnerabilities"][0]["severity_review_tickets"], "")
            # The label is the template's, not a fragment the tester wrote.
            conclusion = next(content for content in payload["vulnerabilities"][0]["contents"] if content["type"] == "in_conclusion")
            self.assertEqual(conclusion["fragments"], [])

    def test_ticket_formats_from_older_reports_normalise_to_bare_digits(self) -> None:
        """Liberal in what it accepts, strict in what survives. The last case is the point: mining
        digits out of surrounding text would invent 2024 as a ticket, and a fabricated reference in
        a delivered report reads exactly as plausibly as a real one."""
        for printed, expected in (
            ("GRIMPEN-1234\nGRIMPEN-5678", "1234\n5678"),
            ("3454, 3453, 2323", "3454\n3453\n2323"),
            ("GRIMPEN-1234; 5678", "1234\n5678"),
            ("3523", "3523"),
            ("OLDKEY-3523", "3523"),
            ("  3523  ", "3523"),
            ("GRIMPEN-3523 (closed 2024)", ""),
            ("N/A", ""),
            ("", ""),
        ):
            with self.subTest(printed=printed):
                self.assertEqual(_parsed_ticket_lines(printed)[0], expected)

    def test_cvss_values_are_read_back_from_the_row_that_names_the_finding(self) -> None:
        """The table carries a row per rendered finding while the importer drops every Resolved one,
        so a Resolved finding first in the document is what proves the match is not positional."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._retest_report(folder)
            report.engagement.segment = "Asia"
            retained = report.vulnerabilities[0]
            retained.severity = retained.likelihood = retained.impact = "low"
            retained.cvss_score, retained.cvss_vector = "3.1", "CVSS:3.1/AV:N/AC:H/PR:H/UI:R/S:U/C:L/I:N/A:N"
            resolved = copy.deepcopy(retained)
            resolved.uid, resolved.display_id, resolved.title = "v_done", "002", "Fixed last year"
            resolved.status = "resolved"
            resolved.severity = resolved.likelihood = resolved.impact = "critical"
            resolved.cvss_score, resolved.cvss_vector = "9.8", "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H"
            # Critical sorts above low, so the dropped finding owns the first row of the table.
            report.vulnerabilities = [retained, resolved]
            resources = Path(__file__).resolve().parent.parent / "resources"
            document = render_report_docx(report, main_template_path(report, resources), folder, allow_incomplete=True)

            payload, _evidence, summary = parse_report_docx(document)
            imported = payload["vulnerabilities"]
            self.assertEqual([finding["title"] for finding in imported], ["Authorization bypass"])
            self.assertEqual(imported[0]["cvss_score"], "3.1", "the dropped finding's row was taken by position")
            self.assertEqual(imported[0]["cvss_vector"], "CVSS:3.1/AV:N/AC:H/PR:H/UI:R/S:U/C:L/I:N/A:N")
            self.assertIn("Fixed last year", summary["dropped_resolved"])

    def test_cvss_round_trip_preserves_every_value_save_validation_accepts(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._retest_report(folder)
            report.engagement.segment = "Asia"
            finding = report.vulnerabilities[0]
            finding.cvss_score = "٩.٨"
            finding.cvss_vector = "CVSS:3.1/AV:Ν"
            self.assertEqual(finding_input_issues(report), [])

            resources = Path(__file__).resolve().parent.parent / "resources"
            document = render_report_docx(report, main_template_path(report, resources), folder)
            payload, _evidence, _summary = parse_report_docx(document)
            imported = payload["vulnerabilities"][0]
            self.assertEqual(
                (imported["cvss_score"], imported["cvss_vector"]),
                (finding.cvss_score, finding.cvss_vector),
            )

    def test_a_template_without_the_section_table_imports_with_the_fields_empty(self) -> None:
        """Three of the four templates have no such table, so its absence is normal and never raises."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            document = render_report_docx(self._retest_report(folder), TEMPLATE, folder)
            payload, _evidence, _summary = parse_report_docx(document)
            finding = payload["vulnerabilities"][0]
            self.assertEqual((finding["cvss_score"], finding["cvss_vector"]), ("", ""))

    def test_a_report_with_no_findings_imports_as_a_draft_with_none(self) -> None:
        """Next year's pentest starts by importing this one, and no finding section claims the
        placeholder rows a no-findings report prints."""
        resources = Path(__file__).resolve().parent.parent / "resources"
        for segment in ("JH", "Asia"):
            for mode in ("editable", "retest"):
                with self.subTest(segment=segment, mode=mode), tempfile.TemporaryDirectory() as temporary_directory:
                    folder = Path(temporary_directory)
                    report = self._report(folder)
                    report.engagement.segment = segment
                    report.vulnerabilities = []
                    document = render_report_docx(report, main_template_path(report, resources), folder)

                    payload, evidence, summary = parse_report_docx(document, mode=mode)
                    self.assertEqual(payload["vulnerabilities"], [])
                    self.assertEqual(evidence, {})
                    self.assertEqual(summary["retained"], 0)

    def _non_production_report(self, folder: Path, label: str) -> Report:
        """The production fixture never prints a non-production heading, so it cannot catch a
        mislabelled round trip."""
        report = self._report(folder)
        report.engagement.non_production_label = label
        report.engagement.tested_environments = ["production", "non_production"]
        report.engagement.test_windows["non_production"] = TestWindow(start_date=date(2026, 8, 3), end_date=date(2026, 8, 4))
        report.scope_targets.append(
            ScopeTarget(target_id="t_web_np", environment="non_production", channel="web", value="https://uat.example.test")
        )
        finding = report.vulnerabilities[0]
        finding.scope = Scope(mode="custom", target_ids=["t_web", "t_web_np"])
        proof = next(content for content in finding.contents if content.type == "proof_of_concept")
        proof.fragments.append(ImageFragment(
            frag_id="f_image_np", type="image", environment="non_production",
            evidence_id="ev_shot", caption="Non-production response",
        ))
        return report

    def test_a_non_production_evidence_image_keeps_its_environment_through_a_round_trip(self) -> None:
        """The heading above the screenshots is the label's only appearance in the document: the
        scope rows read "Non-Production Environment" from the template, not the tester's label. Miss
        it and the images below inherit the previous environment, so non-prod evidence files as prod."""
        for label in ("NON-PROD", "STAGE", "TEST/MO"):
            with self.subTest(label=label), tempfile.TemporaryDirectory() as temporary_directory:
                folder = Path(temporary_directory)
                report = self._non_production_report(folder, label)
                document = render_report_docx(report, TEMPLATE, folder)
                payload, _evidence, _summary = parse_report_docx(document)

                self.assertEqual(payload["engagement"]["non_production_label"], label, "the label was not recovered")
                fragments = [
                    fragment for finding in payload["vulnerabilities"]
                    for content in finding["contents"] if content["type"] == "previous_proof_of_concept"
                    for fragment in content["fragments"]
                ]
                environments = [fragment["environment"] for fragment in fragments if fragment["type"] == "image"]
                self.assertIn("non_production", environments, "non-production evidence came back filed elsewhere")
                stray = [fragment for fragment in fragments if fragment["type"] == "instance_title" and fragment["text"].rstrip(":").upper() == label]
                self.assertEqual(stray, [], "the environment heading survived as a visible fragment")

    def test_an_environment_label_repeated_after_an_instance_title_imports_cleanly(self) -> None:
        """Reports written before this app repeat the label after each instance title too."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            proof = next(content for content in report.vulnerabilities[0].contents if content.type == "proof_of_concept")
            proof.fragments = [
                ListFragment(frag_id="f_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])]),
                ImageFragment(frag_id="f_first", type="image", environment="production", evidence_id="ev_shot", caption="First response"),
                InstanceTitleFragment(frag_id="f_title", type="instance_title", text="Second account"),
                ImageFragment(frag_id="f_second", type="image", environment="production", evidence_id="ev_shot", caption="Second response"),
            ]
            document = render_report_docx(report, TEMPLATE, folder)
            self.assertEqual([item.text for item in Document(BytesIO(document)).paragraphs].count("PROD:"), 2, "the label must print twice for this to test anything")

            payload, _evidence, _summary = parse_report_docx(document)
            imported = next(content for content in payload["vulnerabilities"][0]["contents"] if content["type"] == "previous_proof_of_concept")
            self.assertEqual([fragment["type"] for fragment in imported["fragments"]], ["numbered_list", "image", "instance_title", "image"])
            self.assertEqual(imported["fragments"][2]["text"], "Second account")
            self.assertEqual([fragment["environment"] for fragment in imported["fragments"] if fragment["type"] == "image"], ["production", "production"])

    def test_every_fragment_type_is_recognisable_in_a_generated_report(self) -> None:
        document = self._rendered()
        formats = numbering_formats(document)
        found = {}
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            kind = classify_paragraph(paragraph, formats)
            if paragraph._p.findall(".//" + qn("a:blip")):
                found["image"] = kind
            for marker, expected in (
                ("A justified paragraph.", "paragraph"),
                ("Validate independently.", "note"),
                ("A bullet.", "bulleted_list"),
                ("Send the request.", "numbered_list"),
                ("Instance 1: First instance", "instance_title"),
                ("GET /accounts/123", "code_block"),
                ("Request", "caption"),
            ):
                if text == marker or (marker == "Validate independently." and text.endswith(marker)):
                    found[expected] = kind
        for expected, actual in found.items():
            self.assertEqual(actual, expected, f"{expected} was read as {actual}")
        self.assertEqual(
            set(found),
            {"paragraph", "note", "bulleted_list", "numbered_list", "instance_title", "code_block", "caption", "image"},
            "a fragment type never appeared in the rendered document",
        )

    def test_a_code_block_directly_under_its_steps_is_still_read_as_code(self) -> None:
        """A proof of concept now prints a code block right under the step before it, with no blank
        paragraph between. The importer tells a code block by its shading, so it must not need one."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            proof = next(content for content in report.vulnerabilities[0].contents if content.type == "proof_of_concept")
            screenshot = next(fragment for fragment in proof.fragments if fragment.type == "image")
            proof.fragments = [
                ListFragment(frag_id="f_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Send the request.")])]),
                CodeFragment(frag_id="f_code", type="code_block", caption=None, text="GET /accounts/123"),
                screenshot,
            ]
            document = Document(BytesIO(render_report_docx(report, TEMPLATE, folder)))
        formats = numbering_formats(document)
        paragraphs = document.paragraphs
        index = next(position for position, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "GET /accounts/123")
        self.assertEqual(classify_paragraph(paragraphs[index], formats), "code_block")
        self.assertEqual(paragraphs[index - 1].text.strip(), "Send the request.", "a blank paragraph still sits above the code block")
        self.assertEqual(classify_paragraph(paragraphs[index - 1], formats), "numbered_list")

    def test_a_table_is_a_body_element_not_a_paragraph(self) -> None:
        """``table_fragment.docx`` adds no paragraph style of its own, so a table is only ever found
        structurally."""
        payload, _evidence, _summary = parse_report_docx(base_render(), mode="editable")
        proof = next(content for content in payload["vulnerabilities"][0]["contents"] if content["type"] == "proof_of_concept")
        table = next(fragment for fragment in proof["fragments"] if fragment["type"] == "table")
        text = lambda cell: "".join(run["text"] for run in cell["runs"])
        self.assertEqual([text(cell) for cell in table["header"]], ["Header"])
        self.assertEqual([[text(cell) for cell in row] for row in table["rows"]], [["Cell"]])

    def test_list_format_is_resolved_rather_than_read_from_the_numbering_id(self) -> None:
        """Merging the components renumbers the lists, so trusting ``numId`` swaps bulleted and
        numbered lists. Only the resolved format survives."""
        document = self._rendered()
        formats = numbering_formats(document)
        by_text = {}
        for paragraph in document.paragraphs:
            properties = paragraph._p.find(qn("w:pPr"))
            numbering = properties.find(qn("w:numPr")) if properties is not None else None
            if numbering is None:
                continue
            number_id = numbering.find(qn("w:numId")).get(qn("w:val"))
            by_text[paragraph.text.strip()] = (number_id, formats.get(number_id))
        self.assertEqual(by_text["A bullet."][1], "bullet")
        self.assertEqual(by_text["Send the request."][1], "decimal")
        self.assertNotEqual(by_text["A bullet."][0], by_text["Send the request."][0])


    def test_editable_mode_keeps_the_source_status_and_proof_role(self) -> None:
        data = base_render()

        payload, evidence, _summary = parse_report_docx(data, mode="editable")

        finding = payload["vulnerabilities"][0]
        self.assertEqual(finding["status"], "open_new")
        sections = {content["type"]: content["fragments"] for content in finding["contents"]}
        self.assertNotIn("previous_proof_of_concept", sections)
        self.assertEqual(
            [fragment["type"] for fragment in sections["proof_of_concept"]],
            ["numbered_list", "instance_title", "code_block", "table", "image"],
        )
        self.assertEqual(len(evidence), 1)
        Report.model_validate(payload)

    def test_editable_mode_restores_every_printed_engagement_field(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            report.engagement.app_owner = "Application Owner"
            report.engagement.limitations = "No SMS access"
            report.engagement.test_windows["production"].test_time = "22:00 EST"
            report.engagement.test_accounts = [
                TestAccount(user_role="Administrator", username="alice"),
                TestAccount(user_role="Viewer", username="bob"),
            ]

            payload, _evidence, _summary = parse_report_docx(
                render_report_docx(report, TEMPLATE, folder), mode="editable",
            )

            engagement = payload["engagement"]
            self.assertEqual(engagement["app_name"], "Northstar Banking")
            self.assertEqual(engagement["app_owner"], "Application Owner")
            self.assertEqual(engagement["segment"], "JH")
            self.assertEqual(engagement["report_type"], "annual_pentest")
            self.assertEqual(engagement["report_date"], "2026-09-09")
            self.assertEqual(engagement["tester"], "QA Tester")
            self.assertEqual(engagement["test_windows"]["production"], {
                "start_date": "2026-08-01", "end_date": "2026-08-02", "test_time": "22:00 EST",
            })
            self.assertEqual(engagement["test_accounts"], [
                {"user_role": "Administrator", "username": "alice"},
                {"user_role": "Viewer", "username": "bob"},
            ])
            self.assertEqual(engagement["limitations"], "No SMS access")
            self.assertEqual((engagement.get("ci_number", ""), engagement.get("bsn_number", "")), ("", ""))
            Report.model_validate(payload)

    def test_editable_mode_preserves_visible_whitespace_in_user_fields(self) -> None:
        resources = Path(__file__).resolve().parent.parent / "resources"
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            report.engagement.limitations = "No SMS access\nRead only after  17:00"
            report.engagement.tested_channels = ["api", "thick_client"]
            report.scope_targets = [
                ScopeTarget(target_id="t_api", environment="production", channel="api", value="POST  /v1/pay"),
                ScopeTarget(
                    target_id="t_component", environment="production", channel="thick_client",
                    value="Acme.exe", description="Desktop  client",
                ),
            ]
            report.vulnerabilities[0].scope.target_ids = ["t_api", "t_component"]

            payload, _evidence, _summary = parse_report_docx(
                render_report_docx(report, main_template_path(report, resources), folder),
                mode="editable",
            )

            self.assertEqual(payload["engagement"]["limitations"], report.engagement.limitations)
            targets = {target["channel"]: target for target in payload["scope_targets"]}
            self.assertEqual(targets["api"]["value"], "POST  /v1/pay")
            self.assertEqual(targets["thick_client"]["description"], "Desktop  client")

    def test_editable_mode_rejects_invalid_nonempty_finding_scalars(self) -> None:
        resources = Path(__file__).resolve().parent.parent / "resources"
        # None is the unedited document: it must import, or every refusal below proves nothing.
        cases = {
            None: None,
            "display_id": '"Authorization bypass" has an invalid finding ID "ABC". A finding ID is one to five digits.',
            "cvss_score": '"Authorization bypass" has an invalid CVSS score "8.1?". CVSS Score cannot have "?".',
            "cvss_vector": '"Authorization bypass" has an invalid CVSS vector "CVSS:3.1/AV:N (bad)". CVSS Vector cannot have " ", "(" or ")".',
            "severity_review_tickets": '"Authorization bypass" has an invalid severity review ticket: "not-a-ticket".',
        }
        for field, message in cases.items():
            with self.subTest(field=field), tempfile.TemporaryDirectory() as temporary_directory:
                folder = Path(temporary_directory)
                report = self._retest_report(folder, tickets="1234")
                report.engagement.segment = "Asia"
                finding = report.vulnerabilities[0]
                finding.cvss_score = "8.1"
                finding.cvss_vector = "CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:U/C:H/I:H/A:N"
                document = Document(BytesIO(render_report_docx(report, main_template_path(report, resources), folder)))
                if field == "display_id":
                    summary = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Findings")
                    summary.rows[1].cells[4].text = "ABC"
                elif field == "cvss_score":
                    cvss = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Section")
                    cvss.rows[1].cells[3].text = "8.1?"
                elif field == "cvss_vector":
                    cvss = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Section")
                    cvss.rows[1].cells[4].text = "CVSS:3.1/AV:N (bad)"
                elif field == "severity_review_tickets":
                    paragraph = next(paragraph for paragraph in document.paragraphs if "GRIMPEN-1234" in paragraph.text)
                    paragraph.text = "GRIMPEN-1234\nnot-a-ticket"
                output = BytesIO()
                document.save(output)

                if message is None:
                    parse_report_docx(output.getvalue(), mode="editable")
                    continue
                with self.assertRaises(ReportImportError) as raised:
                    parse_report_docx(output.getvalue(), mode="editable")
                self.assertIn(message, str(raised.exception))

    def test_editable_scope_recovers_all_four_template_variants(self) -> None:
        resources = Path(__file__).resolve().parent.parent / "resources"
        cases = (
            ("web", "JH", "https://prod.example.test", ""),
            ("api", "Asia", "/v1/accounts/{id}", ""),
            ("thick_client", "JH", "Acme.exe", "Desktop client"),
            ("mobile", "Asia", "Acme.ipa", "iOS client"),
        )
        for channel, segment, value, description in cases:
            with self.subTest(channel=channel, segment=segment), tempfile.TemporaryDirectory() as temporary_directory:
                folder = Path(temporary_directory)
                report = self._report(folder)
                report.engagement.segment = segment
                report.engagement.tested_channels = [channel]
                report.scope_targets = [ScopeTarget(
                    target_id="t_scope", environment="production", channel=channel,
                    value=value, description=description,
                )]
                report.vulnerabilities[0].scope.target_ids = ["t_scope"]
                if segment == "Asia":
                    report.vulnerabilities[0].cvss_score = "8.1"
                    report.vulnerabilities[0].cvss_vector = "CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:U/C:H/I:H/A:N"

                data = render_report_docx(report, main_template_path(report, resources), folder)
                payload, _evidence, summary = parse_report_docx(data, mode="editable")

                self.assertEqual(payload["engagement"]["tested_channels"], [channel])
                self.assertEqual(payload["scope_targets"][0]["value"], value)
                self.assertEqual(payload["scope_targets"][0]["description"], description)
                self.assertFalse(any("Review the environment" in warning or "review the app type" in warning for warning in summary["warnings"]))
                self.assertTrue(any("rendered Word copies" in warning for warning in summary["warnings"]))
                Report.model_validate(payload)

    def test_non_production_component_environment_uses_finding_evidence_when_dates_are_missing(self) -> None:
        resources = Path(__file__).resolve().parent.parent / "resources"
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            report.engagement.tested_environments = ["non_production"]
            report.engagement.tested_channels = ["thick_client"]
            report.engagement.test_windows = {}
            report.scope_targets = [ScopeTarget(
                target_id="t_component", environment="non_production", channel="thick_client",
                value="Acme.Staging.exe", description="UAT desktop client",
            )]
            report.vulnerabilities[0].scope.target_ids = ["t_component"]

            payload, _evidence, summary = parse_report_docx(
                render_report_docx(report, main_template_path(report, resources), folder, allow_incomplete=True),
                mode="editable",
            )

            self.assertEqual(payload["scope_targets"], [{
                **payload["scope_targets"][0],
                "environment": "non_production",
                "channel": "thick_client",
                "value": "Acme.Staging.exe",
                "description": "UAT desktop client",
                "order": 0,
            }])
            self.assertEqual(payload["vulnerabilities"][0]["scope"]["target_ids"], [payload["scope_targets"][0]["target_id"]])
            self.assertFalse(any("Review the environment" in warning for warning in summary["warnings"]))

    def test_component_environment_defaults_once_when_both_environments_are_observed(self) -> None:
        resources = Path(__file__).resolve().parent.parent / "resources"
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            report.engagement.tested_environments = ["production", "non_production"]
            report.engagement.tested_channels = ["thick_client"]
            report.engagement.test_windows["non_production"] = TestWindow(
                start_date=date(2026, 8, 3), end_date=date(2026, 8, 4),
            )
            report.scope_targets = [ScopeTarget(
                target_id="t_component", environment="non_production", channel="thick_client",
                value="Acme.exe", description="Desktop client",
            )]
            report.vulnerabilities[0].scope.target_ids = ["t_component"]

            payload, _evidence, summary = parse_report_docx(
                render_report_docx(report, main_template_path(report, resources), folder, allow_incomplete=True),
                mode="editable",
            )

            self.assertEqual(len(payload["scope_targets"]), 2)
            targets = {target["environment"]: target for target in payload["scope_targets"]}
            self.assertEqual(targets["production"]["description"], "Desktop client")
            self.assertEqual(targets["non_production"]["description"], "")
            self.assertEqual(payload["vulnerabilities"][0]["scope"]["target_ids"], [targets["non_production"]["target_id"]])
            self.assertTrue(any('component "Acme.exe"' in warning for warning in summary["warnings"]))
            self.assertTrue(any('"Acme.exe"; it was added' in warning for warning in summary["warnings"]))

    def test_editable_mode_keeps_all_statuses_and_printed_sections(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            open_new = report.vulnerabilities[0]

            previously_discovered = copy.deepcopy(open_new)
            previously_discovered.uid = "v_previous"
            previously_discovered.display_id = "002"
            previously_discovered.title = "Previously discovered issue"
            previously_discovered.status = "open_previously_discovered"
            previously_discovered.contents.extend([
                Content(type="previous_proof_of_concept", fragments=[
                    ListFragment(frag_id="f_previous_steps", type="numbered_list", items=[ListItem(runs=[Run(text="Prior evidence.")])]),
                ]),
                Content(type="in_conclusion", fragments=[
                    ParagraphFragment(frag_id="f_previous_conclusion", type="paragraph", runs=[Run(text="Still open.")]),
                ]),
            ])
            for content in previously_discovered.contents:
                for fragment in content.fragments:
                    fragment.frag_id += "_previous"

            resolved = copy.deepcopy(previously_discovered)
            resolved.uid = "v_resolved"
            resolved.display_id = "003"
            resolved.title = "Resolved issue"
            resolved.status = "resolved"
            for content in resolved.contents:
                for fragment in content.fragments:
                    fragment.frag_id += "_resolved"
            remediation = next(content for content in resolved.contents if content.type == "recommended_remediation")
            remediation.fragments = [ParagraphFragment(
                frag_id="f_resolved_remediation",
                type="paragraph",
                runs=[Run(text=RESOLVED_REMEDIATION)],
                generated="resolved_remediation",
            )]
            resolved_on_non_prod = copy.deepcopy(previously_discovered)
            resolved_on_non_prod.uid = "v_non_prod"
            resolved_on_non_prod.display_id = "004"
            resolved_on_non_prod.title = "Fixed in the lower region only"
            resolved_on_non_prod.status = "open_resolved_on_non_prod"
            for content in resolved_on_non_prod.contents:
                for fragment in content.fragments:
                    fragment.frag_id += "_non_prod"
            report.vulnerabilities = [open_new, previously_discovered, resolved, resolved_on_non_prod]

            data = render_report_docx(report, TEMPLATE, folder)
            payload, evidence, summary = parse_report_docx(data, mode="editable")

            findings = {finding["title"]: finding for finding in payload["vulnerabilities"]}
            self.assertEqual({title: finding["status"] for title, finding in findings.items()}, {
                "Authorization bypass": "open_new",
                "Previously discovered issue": "open_previously_discovered",
                "Resolved issue": "resolved",
                "Fixed in the lower region only": "open_resolved_on_non_prod",
            })
            self.assertEqual(
                [content["type"] for content in findings["Fixed in the lower region only"]["contents"]],
                ["description", "recommended_remediation", "previous_proof_of_concept", "proof_of_concept", "in_conclusion"],
                "the fourth status prints what previously discovered prints",
            )
            self.assertEqual(
                [content["type"] for content in findings["Authorization bypass"]["contents"]],
                ["description", "recommended_remediation", "proof_of_concept"],
            )
            self.assertEqual(
                [content["type"] for content in findings["Previously discovered issue"]["contents"]],
                ["description", "recommended_remediation", "previous_proof_of_concept", "proof_of_concept", "in_conclusion"],
            )
            resolved_remediation = next(
                content for content in findings["Resolved issue"]["contents"]
                if content["type"] == "recommended_remediation"
            )
            self.assertEqual(resolved_remediation["fragments"][0]["generated"], "resolved_remediation")
            referenced = {
                fragment["evidence_id"]
                for finding in findings.values()
                for content in finding["contents"]
                for fragment in content["fragments"]
                if fragment["type"] == "image" and fragment["evidence_id"]
            }
            self.assertEqual(referenced, set(payload["evidence"]))
            self.assertEqual(referenced, set(evidence))
            self.assertEqual(summary["retained"], 4)
            Report.model_validate(payload)

    def test_editable_mode_rejects_edited_resolved_remediation(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._retest_report(folder)
            report.vulnerabilities[0].status = "resolved"
            remediation = next(content for content in report.vulnerabilities[0].contents if content.type == "recommended_remediation")
            remediation.fragments = [ParagraphFragment(
                frag_id="f_resolved_remediation",
                type="paragraph",
                runs=[Run(text="Keep this manually edited advice.")],
            )]

            data = render_report_docx(report, TEMPLATE, folder, allow_incomplete=True)
            with self.assertRaisesRegex(
                ReportImportError,
                r'"Authorization bypass" has edited remediation while Resolved\..*but the report has: "Keep this manually edited advice\."',
            ):
                parse_report_docx(data, mode="editable")

    def test_an_unmatched_api_location_becomes_a_visible_review_target(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            report.engagement.tested_channels = ["api"]
            report.scope_targets = [ScopeTarget(
                target_id="t_api", environment="production", channel="api", value="/v1/accounts/{id}",
            )]
            report.vulnerabilities[0].scope.target_ids = ["t_api"]
            document = Document(BytesIO(render_report_docx(report, TEMPLATE, folder)))
            detail = next(table for table in document.tables if any(row.cells[0].text.strip() == "Location" for row in table.rows))
            production = next(row for row in detail.rows if row.cells[1].text.startswith("Production Environment"))
            production.cells[1].paragraphs[1].text = "/v1/accounts/{account_id}"
            output = BytesIO()
            document.save(output)

            payload, _evidence, summary = parse_report_docx(output.getvalue(), mode="editable")

            finding = payload["vulnerabilities"][0]
            selected = next(target for target in payload["scope_targets"] if target["target_id"] in finding["scope"]["target_ids"])
            self.assertEqual((selected["channel"], selected["value"]), ("api", "/v1/accounts/{account_id}"))
            self.assertEqual(finding["scope"]["custom_locations"], {})
            self.assertTrue(any("review the app type" in warning for warning in summary["warnings"]))
            Report.model_validate(payload)

    def test_retest_summary_discloses_shared_scope_corrections(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            report.engagement.tested_channels = ["api"]
            report.scope_targets = [ScopeTarget(
                target_id="t_api", environment="production", channel="api", value="/v1/accounts/{id}",
            )]
            report.vulnerabilities[0].scope.target_ids = ["t_api"]
            document = Document(BytesIO(render_report_docx(report, TEMPLATE, folder)))
            detail = next(table for table in document.tables if any(row.cells[0].text.strip() == "Location" for row in table.rows))
            production = next(row for row in detail.rows if row.cells[1].text.startswith("Production Environment"))
            production.cells[1].paragraphs[1].text = "/v1/accounts/{account_id}"
            output = BytesIO()
            document.save(output)

            _payload, _evidence, summary = parse_report_docx(output.getvalue())

            self.assertEqual(summary["retained"], 1)
            self.assertTrue(any("review the app type" in warning for warning in summary["warnings"]))

    def test_wrapped_scope_and_location_text_stays_one_target(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            value = "https://prod.example.test/" + "very-long-segment/" * 8
            report.scope_targets[0].value = value

            payload, _evidence, _summary = parse_report_docx(
                render_report_docx(report, TEMPLATE, folder), mode="editable",
            )

            self.assertEqual([target["value"] for target in payload["scope_targets"]], [value])
            self.assertEqual(payload["vulnerabilities"][0]["scope"]["target_ids"], [payload["scope_targets"][0]["target_id"]])

    def test_note_runs_and_numbered_list_boundaries_survive(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            description = next(content for content in report.vulnerabilities[0].contents if content.type == "description")
            note = next(fragment for fragment in description.fragments if fragment.type == "note")
            note.runs = [Run(text="Validate ", bold=True), Run(text="independently.", italic=True)]
            proof = next(content for content in report.vulnerabilities[0].contents if content.type == "proof_of_concept")
            image = next(fragment for fragment in proof.fragments if fragment.type == "image")
            proof.fragments = [
                ListFragment(frag_id="f_first", type="numbered_list", items=[
                    ListItem(runs=[Run(text="Step one.")]), ListItem(runs=[Run(text="Step two.")]),
                ]),
                image,
                ListFragment(frag_id="f_continued", type="numbered_list", continue_numbering=True, items=[
                    ListItem(runs=[Run(text="Step three.")]),
                ]),
                ListFragment(frag_id="f_restarted", type="numbered_list", items=[
                    ListItem(runs=[Run(text="New step one.")]),
                ]),
            ]

            payload, _evidence, _summary = parse_report_docx(
                render_report_docx(report, TEMPLATE, folder), mode="editable",
            )

            imported_description = payload["vulnerabilities"][0]["contents"][0]["fragments"]
            imported_note = next(fragment for fragment in imported_description if fragment["type"] == "note")
            self.assertEqual(imported_note["runs"], [
                {"text": "Validate ", "bold": True, "italic": True},
                {"text": "independently.", "italic": True},
            ])
            imported_proof = next(
                content for content in payload["vulnerabilities"][0]["contents"]
                if content["type"] == "proof_of_concept"
            )
            lists = [fragment for fragment in imported_proof["fragments"] if fragment["type"] == "numbered_list"]
            self.assertEqual([len(fragment["items"]) for fragment in lists], [2, 1, 1])
            self.assertEqual([fragment.get("continue_numbering", False) for fragment in lists], [False, True, False])

    def test_the_blank_line_after_a_bulleted_list_is_not_imported(self) -> None:
        """A generated list is followed by an empty paragraph. Read back as a fragment it would be text the
        tester never wrote, and it would stand between a list and the one after it."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            description = next(content for content in report.vulnerabilities[0].contents if content.type == "description")
            description.fragments = [
                ListFragment(frag_id="f_first", type="bulleted_list", items=[
                    ListItem(runs=[Run(text="One.")]), ListItem(runs=[Run(text="Two.")]),
                ]),
                ListFragment(frag_id="f_second", type="bulleted_list", items=[ListItem(runs=[Run(text="Three.")])]),
                ParagraphFragment(frag_id="f_after", type="paragraph", runs=[Run(text="After the lists.")]),
            ]

            payload, _evidence, _summary = parse_report_docx(
                render_report_docx(report, TEMPLATE, folder), mode="editable",
            )

            imported = payload["vulnerabilities"][0]["contents"][0]["fragments"]
            self.assertEqual([fragment["type"] for fragment in imported], ["bulleted_list", "bulleted_list", "paragraph"])
            self.assertEqual([len(fragment["items"]) for fragment in imported[:2]], [2, 1])

    def test_editable_image_preserves_embedded_png_bytes_and_display_width(self) -> None:
        data = base_render()
        document = Document(BytesIO(data))
        paragraph = next(paragraph for paragraph in document.paragraphs if paragraph._p.findall(".//" + qn("a:blip")))
        blip = paragraph._p.find(".//" + qn("a:blip"))
        expected = document.part.related_parts[blip.get(qn("r:embed"))].blob

        payload, evidence, _summary = parse_report_docx(data, mode="editable")

        image = next(
            fragment
            for content in payload["vulnerabilities"][0]["contents"]
            for fragment in content["fragments"]
            if fragment["type"] == "image"
        )
        self.assertEqual(evidence[image["evidence_id"]], expected)
        self.assertEqual(image["width_mm"], 155.0)

    def test_an_imported_image_regenerates_at_the_size_it_was_imported(self) -> None:
        """The embedded picture already carries the border, so drawing one again would grow it."""
        payload, evidence, _summary = parse_report_docx(base_render(), mode="editable")
        report = Report.model_validate(payload)
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            (folder / "evidence").mkdir()
            for evidence_id, data in evidence.items():
                (folder / report.evidence[evidence_id].file).write_bytes(data)
            regenerated = Document(BytesIO(render_report_docx(report, TEMPLATE, folder)))

        embedded = sorted(
            (part.image.px_width, part.image.px_height)
            for part in (
                regenerated.part.related_parts[shape._inline.graphic.graphicData.pic.blipFill.blip.get(qn("r:embed"))]
                for shape in regenerated.inline_shapes
            )
        )
        imported = sorted((item.width_px, item.height_px) for item in report.evidence.values())
        self.assertEqual(embedded, imported)

    def test_the_editable_import_names_rendered_copies_but_promises_no_second_border(self) -> None:
        _payload, _evidence, summary = parse_report_docx(base_render(), mode="editable")

        self.assertTrue(any("rendered Word copies" in warning for warning in summary["warnings"]))
        self.assertFalse(any("another border" in warning for warning in summary["warnings"]))

    def test_docx_resource_limits_reject_before_import(self) -> None:
        data = base_render()
        cases = (
            {"max_package_bytes": len(data) - 1},
            {"max_files": 1},
            {"max_uncompressed_bytes": 1},
            {"max_image_bytes": 1},
            {"max_image_pixels": 1},
            {"max_evidence_bytes": 1},
        )
        for limits in cases:
            with self.subTest(limits=limits), self.assertRaises(ReportImportLimitError):
                parse_report_docx(data, mode="editable", **limits)

    def test_oversized_media_member_rejects_before_python_docx_materializes_it(self) -> None:
        data = base_render()

        with patch("app.docx_import.Document") as document, self.assertRaises(ReportImportLimitError):
            parse_report_docx(data, mode="editable", max_image_bytes=1)

        document.assert_not_called()

    def test_editable_mode_rejects_image_width_above_renderer_maximum(self) -> None:
        document = Document(BytesIO(base_render()))
        paragraph = next(paragraph for paragraph in document.paragraphs if paragraph._p.findall(".//" + qn("a:blip")))
        extent = paragraph._p.find(".//" + qn("wp:extent"))
        original_width, original_height = int(extent.get("cx")), int(extent.get("cy"))
        oversized_width = 156 * 36_000
        extent.set("cx", str(oversized_width))
        extent.set("cy", str(round(original_height * oversized_width / original_width)))
        output = BytesIO()
        document.save(output)

        with self.assertRaisesRegex(ReportImportError, "155 mm"):
            parse_report_docx(output.getvalue(), mode="editable")

    def test_editable_mode_rejects_rotated_evidence_geometry(self) -> None:
        document = Document(BytesIO(base_render()))
        paragraph = next(paragraph for paragraph in document.paragraphs if paragraph._p.findall(".//" + qn("a:blip")))
        paragraph._p.find(".//" + qn("a:xfrm")).set("rot", "60000")
        output = BytesIO()
        document.save(output)

        with self.assertRaisesRegex(ReportImportError, r'^In the Proof of Concept section of "Authorization bypass": Rotated or flipped'):
            parse_report_docx(output.getvalue(), mode="editable")

    def test_the_note_prefix_is_not_kept_as_content(self) -> None:
        """``Note: `` is printed by note_fragment.docx, so keeping it would render ``Note: Note: ``
        the next time the report is generated."""
        payload, _, _ = parse_report_docx(base_render())
        description = payload["vulnerabilities"][0]["contents"][0]["fragments"]
        note = next(fragment for fragment in description if fragment["type"] == "note")
        self.assertEqual("".join(run["text"] for run in note["runs"]), "Validate independently.")

    def test_two_findings_may_share_a_title(self) -> None:
        """Nothing stops a tester naming two findings the same. Keying the summary by title gave
        both the last row's number, and the import then failed on a duplicate finding number."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = self._report(folder)
            twin = copy.deepcopy(report.vulnerabilities[0])
            twin.uid, twin.display_id = "v_twin", "002"
            for content in twin.contents:
                for fragment in content.fragments:
                    fragment.frag_id += "_twin"
            report.vulnerabilities.append(twin)

            payload, _, _ = parse_report_docx(render_report_docx(report, TEMPLATE, folder))
            self.assertEqual([finding["display_id"] for finding in payload["vulnerabilities"]], ["001", "002"])
            Report.model_validate(payload)

    def test_a_title_with_extra_spaces_imports_back_as_typed(self) -> None:
        """The heading and the Findings Summary table print the title as typed, so a report the app wrote
        must pair them whatever spaces the title holds, and the finding keeps the title as typed."""
        titles = (
            ("a double space", "Cross  site scripting"),
            ("a non-breaking space", "Cross\u00a0site scripting"),
            ("a trailing space", "Cross site scripting "),
        )
        for label, title in titles:
            with tempfile.TemporaryDirectory() as temporary_directory:
                folder = Path(temporary_directory)
                report = self._report(folder)
                report.vulnerabilities[0].title = title
                data = render_report_docx(report, TEMPLATE, folder)
            for mode in ("retest", "editable"):
                with self.subTest(title=label, mode=mode):
                    payload, _evidence, _summary = parse_report_docx(data, mode=mode)
                    self.assertEqual(payload["vulnerabilities"][0]["title"], title)

    def test_a_file_that_is_not_a_word_document_is_refused_by_name(self) -> None:
        with self.assertRaisesRegex(ReportImportError, "not a readable Word document"):
            parse_report_docx(b"not a document")


class FindingRowTests(unittest.TestCase):
    """A finding takes its title, severity, ID and status from its finding row. A finding heading finds its
    row by equal title, then by title ignoring capitals and extra spaces, then by position."""

    def test_an_unlabelled_table_before_description_is_not_silently_lost(self) -> None:
        def insert_table(document):
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Account"
            table.cell(0, 1).text = "Sensitive evidence"
            _heading(document, "Description:")._p.addprevious(table._tbl)

        data = _edited(base_render(), insert_table)
        for mode in ("retest", "editable"):
            with self.subTest(mode=mode), self.assertRaisesRegex(
                ReportImportError, r'"Authorization bypass".*table.*before.*Description',
            ):
                parse_report_docx(data, mode=mode)

    @staticmethod
    def _parse(data: bytes, mode: str) -> tuple[dict, list[str]]:
        payload, _evidence, summary = parse_report_docx(data, mode=mode)
        return payload, summary.get("warnings", [])

    def test_a_closed_finding_row_in_other_capitals_keeps_its_status(self) -> None:
        def close(document):
            _set_text(_table_starting(document, "Findings").rows[1].cells[5].paragraphs[0], "CLOSED")

        data = _edited(two_findings_render(), close)
        payload, _evidence, summary = parse_report_docx(data, mode="editable")
        self.assertEqual(payload["vulnerabilities"][0]["status"], "closed")
        self.assertFalse(any("unrecognised status" in warning for warning in summary.get("warnings", [])))
        payload, _evidence, summary = parse_report_docx(data, mode="retest")
        self.assertEqual([finding["title"] for finding in payload["vulnerabilities"]], ["Second finding"])
        self.assertEqual(summary["dropped_resolved"], ["Authorization bypass"])

    def test_hyperlink_text_in_a_description_survives_import(self) -> None:
        def paste_link(document):
            paragraph = _heading(document, "A justified paragraph.")
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "endpoint")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = " Sensitive endpoint"
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)

        data = _edited(base_render(), paste_link)
        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, _evidence, _summary = parse_report_docx(data, mode=mode)
                description = payload["vulnerabilities"][0]["contents"][0]
                self.assertEqual(
                    "".join(run["text"] for run in description["fragments"][0]["runs"]),
                    "A justified paragraph. Sensitive endpoint",
                )

    def test_explicitly_disabled_emphasis_does_not_import_as_enabled(self) -> None:
        for tag, value, field in (("w:b", "0", "bold"), ("w:i", "false", "italic"), ("w:u", "none", "underline")):
            def turn_off(document):
                run = _heading(document, "A justified paragraph.").runs[0]
                mark = OxmlElement(tag)
                mark.set(qn("w:val"), value)
                run._r.get_or_add_rPr().append(mark)

            data = _edited(base_render(), turn_off)
            with self.subTest(field=field):
                payload, _evidence, _summary = parse_report_docx(data, mode="editable")
                paragraph = payload["vulnerabilities"][0]["contents"][0]["fragments"][0]
                self.assertFalse(paragraph["runs"][0].get(field, False))

    def test_a_ticket_label_in_description_does_not_consume_authored_text(self) -> None:
        def explain_tickets(document):
            label = document.add_paragraph(SEVERITY_TICKET_LABEL)
            explanation = document.add_paragraph("The reviewer can request a reassessment.")
            anchor = _heading(document, "A justified paragraph.")
            anchor._p.addnext(explanation._p)
            anchor._p.addnext(label._p)

        data = _edited(base_render(), explain_tickets)
        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, _evidence, _summary = parse_report_docx(data, mode=mode)
                description = payload["vulnerabilities"][0]["contents"][0]
                text = " ".join(
                    run["text"] for fragment in description["fragments"]
                    for run in fragment.get("runs", [])
                )
                self.assertIn(SEVERITY_TICKET_LABEL, text)
                self.assertIn("The reviewer can request a reassessment.", text)
                self.assertEqual(payload["vulnerabilities"][0]["severity_review_tickets"], "")

    def test_unattached_or_repeated_captions_are_not_silently_lost(self) -> None:
        def caption(document, text):
            paragraph = document.add_paragraph(text)
            paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = "FiguresandTables"
            return paragraph

        def orphan(document):
            _heading(document, "Recommended Remediation:")._p.addprevious(caption(document, "An orphan observation")._p)

        def repeated(document):
            original = next(paragraph for paragraph in document.paragraphs if "Production response" in paragraph.text)
            original._p.addnext(caption(document, "A second observation")._p)

        def separated(document):
            table = _table_starting(document, "Header")
            table._tbl.addprevious(document.add_paragraph("This is prose, not a caption.")._p)

        for label, change in (("orphan", orphan), ("repeated image caption", repeated), ("separated table caption", separated)):
            with self.subTest(label=label), self.assertRaisesRegex(ReportImportError, r'section of "Authorization bypass":.*caption'):
                parse_report_docx(_edited(base_render(), change), mode="editable")

    def test_images_in_mixed_text_containers_are_not_silently_lost(self) -> None:
        def mixed_paragraph(document):
            picture = next(paragraph for paragraph in document.paragraphs if paragraph._p.findall(".//" + qn("a:blip")))
            picture.add_run("The response exposes another customer's account.")

        def picture_in_table(document):
            picture = next(paragraph for paragraph in document.paragraphs if paragraph._p.findall(".//" + qn("a:blip")))
            cell = _table_starting(document, "Header").cell(1, 0)
            cell.paragraphs[0]._p.extend(copy.deepcopy(run) for run in picture._p.findall(qn("w:r")))

        for label, change in (("text beside a picture", mixed_paragraph), ("picture inside a table", picture_in_table)):
            data = _edited(base_render(), change)
            for mode in ("retest", "editable"):
                with self.subTest(label=label, mode=mode), self.assertRaisesRegex(
                    ReportImportError, r'section of "Authorization bypass":.*(?:picture|image)',
                ):
                    parse_report_docx(data, mode=mode)

    def test_content_control_text_is_not_silently_lost(self) -> None:
        def wrap_description(document):
            paragraph = _heading(document, "A justified paragraph.")
            control = OxmlElement("w:sdt")
            contents = OxmlElement("w:sdtContent")
            control.append(contents)
            paragraph._p.addprevious(control)
            contents.append(paragraph._p)

        data = _edited(base_render(), wrap_description)
        with self.assertRaisesRegex(ReportImportError, r'Description section of "Authorization bypass":.*content control'):
            parse_report_docx(data, mode="editable")

    def test_a_header_only_table_does_not_gain_an_invented_row(self) -> None:
        def remove_data(document):
            table = _table_starting(document, "Header")
            table._tbl.remove(table.rows[1]._tr)

        payload, _evidence, _summary = parse_report_docx(_edited(base_render(), remove_data), mode="editable")
        proof = next(content for content in payload["vulnerabilities"][0]["contents"] if content["type"] == "proof_of_concept")
        table = next(fragment for fragment in proof["fragments"] if fragment["type"] == "table")
        self.assertEqual(table["header"], [{"runs": [{"text": "Header", "bold": True}]}])
        self.assertEqual(table["rows"], [])

    def test_disabled_paragraph_marks_do_not_change_fragment_type(self) -> None:
        for tag in ("w:b", "w:i"):
            def plain_paragraph(document):
                paragraph = document.add_paragraph("Plain authored observation.")
                properties = OxmlElement("w:rPr")
                mark = OxmlElement(tag)
                mark.set(qn("w:val"), "0")
                properties.append(mark)
                paragraph._p.get_or_add_pPr().append(properties)
                _heading(document, "A justified paragraph.")._p.addnext(paragraph._p)

            with self.subTest(tag=tag):
                payload, _evidence, _summary = parse_report_docx(_edited(base_render(), plain_paragraph), mode="editable")
                fragment = payload["vulnerabilities"][0]["contents"][0]["fragments"][1]
                self.assertEqual(fragment["type"], "paragraph")
                self.assertEqual(fragment["runs"], [{"text": "Plain authored observation."}])

    def test_explicitly_removed_numbering_does_not_become_a_numbered_list(self) -> None:
        def remove_numbering(document):
            paragraph = _heading(document, "A justified paragraph.")
            paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0

        data = _edited(base_render(), remove_numbering)
        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, _evidence, _summary = parse_report_docx(data, mode=mode)
                self.assertEqual(payload["vulnerabilities"][0]["contents"][0]["fragments"][0]["type"], "paragraph")

    def test_an_unexpected_final_group_heading_does_not_truncate_a_finding(self) -> None:
        def insert_heading(document):
            conclusion = [paragraph for paragraph in document.paragraphs if paragraph.text == "Still reachable on retest."][-1]
            heading = document.add_paragraph("Additional retest details")
            heading._p.get_or_add_pPr().get_or_add_pStyle().val = "ReportHeading1"
            evidence = document.add_paragraph("The second endpoint still exposes customer data.")
            conclusion._p.addnext(evidence._p)
            conclusion._p.addnext(heading._p)

        data = _edited(two_findings_render(), insert_heading)
        with self.assertRaisesRegex(ReportImportError, r'"Second finding".*Additional retest details'):
            parse_report_docx(data, mode="editable")

    def test_merged_table_cells_do_not_import_as_duplicated_cells(self) -> None:
        for direction in ("horizontal", "vertical"):
            def merge(document):
                table = _table_starting(document, "Header")
                table.add_column(Mm(20))
                table.cell(0, 1).text = "Second header"
                table.cell(1, 1).text = "Second value"
                other = table.cell(0, 1) if direction == "horizontal" else table.cell(1, 0)
                table.cell(0, 0).merge(other)

            with self.subTest(direction=direction), self.assertRaisesRegex(
                ReportImportError, r'Proof of Concept section of "Authorization bypass":.*merged',
            ):
                parse_report_docx(_edited(base_render(), merge), mode="editable")

    def test_blank_paragraphs_inside_table_cells_keep_their_line_breaks(self) -> None:
        for label, lines in (("leading", ["", "Cell"]), ("middle", ["First", "", "Last"]), ("trailing", ["Cell", ""])):
            def add_lines(document):
                cell = _table_starting(document, "Header").cell(1, 0)
                cell.text = lines[0]
                for line in lines[1:]:
                    cell.add_paragraph(line)

            with self.subTest(label=label):
                payload, _evidence, _summary = parse_report_docx(_edited(base_render(), add_lines), mode="editable")
                proof = next(content for content in payload["vulnerabilities"][0]["contents"] if content["type"] == "proof_of_concept")
                table = next(fragment for fragment in proof["fragments"] if fragment["type"] == "table")
                self.assertEqual("".join(run["text"] for run in table["rows"][0][0]["runs"]), "\n".join(lines))

    def _assert_pairs(self, data: bytes, titles: list[str], position_warnings: list[str]) -> None:
        """Both modes import the findings under ``titles`` and warn about pairs made by position as given."""
        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, warnings = self._parse(data, mode)
                self.assertEqual([finding["title"] for finding in payload["vulnerabilities"]], titles)
                self.assertEqual([warning for warning in warnings if "by position" in warning], position_warnings)

    def test_a_heading_that_differs_from_its_finding_row_only_in_capitals_pairs_without_a_warning(self) -> None:
        data = _edited(two_findings_render(), lambda document: _set_text(_heading(document, "Second finding"), "SECOND Finding"))

        self._assert_pairs(data, ["Authorization bypass", "Second finding"], [])

    def test_a_heading_in_other_words_pairs_with_its_finding_row_by_position_and_warns(self) -> None:
        data = _edited(two_findings_render(), lambda document: _set_text(_heading(document, "Second finding"), "Totally different words"))

        self._assert_pairs(data, ["Authorization bypass", "Second finding"], [
            'The finding heading "Totally different words" was paired with the finding row "Second finding" by position. The finding uses the row\'s title.',
        ])

    def test_a_finding_row_in_other_words_pairs_with_its_heading_by_position_and_gives_its_title(self) -> None:
        def rename_row(document):
            _set_text(_table_starting(document, "Findings").rows[2].cells[0].paragraphs[0], "Renamed finding")

        self._assert_pairs(_edited(two_findings_render(), rename_row), ["Authorization bypass", "Renamed finding"], [
            'The finding heading "Second finding" was paired with the finding row "Renamed finding" by position. The finding uses the row\'s title.',
        ])

    def test_two_retitled_headings_pair_with_their_finding_rows_in_order(self) -> None:
        def retitle(document):
            _set_text(_heading(document, "Authorization bypass"), "First, retitled")
            _set_text(_heading(document, "Second finding"), "Second, retitled")

        self._assert_pairs(_edited(two_findings_render(), retitle), ["Authorization bypass", "Second finding"], [
            'The finding heading "First, retitled" was paired with the finding row "Authorization bypass" by position. The finding uses the row\'s title.',
            'The finding heading "Second, retitled" was paired with the finding row "Second finding" by position. The finding uses the row\'s title.',
        ])

    def test_a_heading_titled_like_another_pairs_with_the_remaining_row_by_position(self) -> None:
        data = _edited(two_findings_render(), lambda document: _set_text(_heading(document, "Second finding"), "Authorization bypass"))

        self._assert_pairs(data, ["Authorization bypass", "Second finding"], [
            'The finding heading "Authorization bypass" was paired with the finding row "Second finding" by position. The finding uses the row\'s title.',
        ])

    def test_a_cvss_row_that_differs_from_its_finding_row_only_in_capitals_pairs_without_a_warning(self) -> None:
        data = _edited(asia_findings_render(), lambda document: _set_text(_table_starting(document, "Section").rows[1].cells[1].paragraphs[0], "AUTHORIZATION BYPASS"))

        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, warnings = self._parse(data, mode)
                self.assertEqual(
                    [(finding["title"], finding["cvss_score"]) for finding in payload["vulnerabilities"]],
                    [("Authorization bypass", "8.1"), ("Second finding", "3.1")],
                )
                self.assertEqual([warning for warning in warnings if "by position" in warning], [])

    def test_a_cvss_row_in_other_words_pairs_with_its_finding_row_by_position_and_warns(self) -> None:
        data = _edited(asia_findings_render(), lambda document: _set_text(_table_starting(document, "Section").rows[2].cells[1].paragraphs[0], "Other finding"))

        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, warnings = self._parse(data, mode)
                self.assertEqual(
                    [(finding["title"], finding["cvss_score"]) for finding in payload["vulnerabilities"]],
                    [("Authorization bypass", "8.1"), ("Second finding", "3.1")],
                )
                self.assertIn('The CVSS row "Other finding" was paired with the finding row "Second finding" by position.', warnings)

    def test_a_finding_row_with_no_cvss_row_imports_without_scores(self) -> None:
        def remove_second(document):
            row = _table_starting(document, "Section").rows[2]._tr
            row.getparent().remove(row)

        data = _edited(asia_findings_render(), remove_second)

        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, _warnings = self._parse(data, mode)
                self.assertEqual(
                    [(finding["title"], finding["cvss_score"]) for finding in payload["vulnerabilities"]],
                    [("Authorization bypass", "8.1"), ("Second finding", "")],
                )

    def test_a_finding_table_that_disagrees_with_its_finding_row_imports_with_the_rows_values_and_warns(self) -> None:
        cases = (
            (3, "Low", "severity", "low", 'Severity is "High" in the finding table and "Low" in the Findings Summary table.'),
            (4, "009", "display_id", "009", 'ID is "001" in the finding table and "009" in the Findings Summary table.'),
            (
                5, "Open (Resolved on Non-Prod)", "status", "open_resolved_on_non_prod",
                'Status is "Open (Previously Discovered)" in the finding table and "Open (Resolved on Non-Prod)" in the Findings Summary table.',
            ),
        )
        for column, value, field, expected_value, sentence in cases:
            def change(document, column=column, value=value):
                _set_text(_table_starting(document, "Findings").rows[1].cells[column].paragraphs[0], value)

            data = _edited(two_findings_render(), change)
            for mode in ("retest", "editable"):
                with self.subTest(column=column, mode=mode):
                    payload, warnings = self._parse(data, mode)
                    self.assertEqual(payload["vulnerabilities"][0][field], expected_value)
                    self.assertIn(f'"Authorization bypass": {sentence} The import uses the Findings Summary table.', warnings)

    def test_a_finding_table_value_that_differs_only_in_capitals_gives_no_warning(self) -> None:
        def shout(document):
            row = next(row for row in _table_after(document, "Authorization bypass").rows if row.cells[0].text.strip() == "Severity")
            _set_text(row.cells[1].paragraphs[0], "HIGH")

        data = _edited(two_findings_render(), shout)

        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                _payload, warnings = self._parse(data, mode)
                self.assertEqual([warning for warning in warnings if "in the finding table" in warning], [])

    def test_a_resolved_finding_that_a_retest_drops_gives_no_finding_table_warning(self) -> None:
        def resolve(document):
            _set_text(_table_starting(document, "Findings").rows[1].cells[5].paragraphs[0], "Resolved")

        _payload, _evidence, summary = parse_report_docx(_edited(two_findings_render(), resolve), mode="retest")

        self.assertEqual(summary["dropped_resolved"], ["Authorization bypass"])
        self.assertEqual([warning for warning in summary.get("warnings", []) if "in the finding table" in warning], [])

    def test_a_resolved_finding_that_a_retest_drops_keeps_its_position_warning(self) -> None:
        def edit(document):
            _set_text(_heading(document, "Authorization bypass"), "Retitled")
            _set_text(_table_starting(document, "Findings").rows[1].cells[5].paragraphs[0], "Resolved")

        _payload, _evidence, summary = parse_report_docx(_edited(two_findings_render(), edit), mode="retest")

        self.assertEqual(summary["dropped_resolved"], ["Authorization bypass"])
        self.assertIn(
            'The finding heading "Retitled" was paired with the finding row "Authorization bypass" by position. The finding uses the row\'s title.',
            summary["warnings"],
        )

    def test_a_finding_table_status_in_other_capitals_is_still_compared_with_the_finding_row(self) -> None:
        def set_status(document, text):
            row = next(row for row in _table_after(document, "Authorization bypass").rows if row.cells[0].text.strip() == "Status")
            _set_text(row.cells[1].paragraphs[0], text)

        different = _edited(two_findings_render(), lambda document: set_status(document, "CLOSED"))
        same = _edited(two_findings_render(), lambda document: set_status(document, "open (previously discovered)"))
        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                _payload, warnings = self._parse(different, mode)
                self.assertIn(
                    '"Authorization bypass": Status is "CLOSED" in the finding table and "Open (Previously Discovered)" in the Findings Summary table. '
                    "The import uses the Findings Summary table.",
                    warnings,
                )
                _payload, warnings = self._parse(same, mode)
                self.assertEqual([warning for warning in warnings if "in the finding table" in warning], [])

    def test_an_unrecognised_status_in_either_place_adds_no_finding_table_warning(self) -> None:
        def table_status(document):
            return next(row for row in _table_after(document, "Authorization bypass").rows if row.cells[0].text.strip() == "Status")

        def unrecognised_in_the_row(document):
            _set_text(_table_starting(document, "Findings").rows[1].cells[5].paragraphs[0], "Pending")
            _set_text(table_status(document).cells[1].paragraphs[0], "Closed")

        def unrecognised_in_the_table(document):
            _set_text(table_status(document).cells[1].paragraphs[0], "Pending")

        for place, change in (("the finding row", unrecognised_in_the_row), ("the finding table", unrecognised_in_the_table)):
            data = _edited(two_findings_render(), change)
            for mode in ("retest", "editable"):
                with self.subTest(place=place, mode=mode):
                    _payload, warnings = self._parse(data, mode)
                    self.assertEqual([warning for warning in warnings if "in the finding table" in warning], [])

    def test_an_empty_finding_row_value_reads_empty_in_the_warning(self) -> None:
        def clear_id(document):
            _set_text(_table_starting(document, "Findings").rows[1].cells[4].paragraphs[0], "")

        data = _edited(two_findings_render(), clear_id)

        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, warnings = self._parse(data, mode)
                self.assertIsNone(payload["vulnerabilities"][0]["display_id"])
                self.assertIn(
                    '"Authorization bypass": ID is "001" in the finding table and empty in the Findings Summary table. '
                    "The import uses the Findings Summary table.",
                    warnings,
                )

    def test_a_finding_table_that_disagrees_twice_gives_one_warning_with_both_values(self) -> None:
        def change(document):
            cells = _table_starting(document, "Findings").rows[1].cells
            _set_text(cells[3].paragraphs[0], "Low")
            _set_text(cells[4].paragraphs[0], "009")

        data = _edited(two_findings_render(), change)

        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                _payload, warnings = self._parse(data, mode)
                self.assertEqual(
                    [warning for warning in warnings if "in the finding table" in warning],
                    [
                        '"Authorization bypass": Severity is "High" in the finding table and "Low" in the Findings Summary table. '
                        'ID is "001" in the finding table and "009" in the Findings Summary table. '
                        "The import uses the Findings Summary table.",
                    ],
                )

    def test_a_heading_pairs_with_the_finding_row_of_exactly_its_title_before_one_that_differs_in_capitals(self) -> None:
        def swap_capitals(document):
            rows = _table_starting(document, "Findings").rows
            _set_text(rows[1].cells[0].paragraphs[0], "AUTHORIZATION BYPASS")
            _set_text(rows[2].cells[0].paragraphs[0], "Authorization bypass")
            _set_text(_heading(document, "Second finding"), "AUTHORIZATION BYPASS")

        data = _edited(two_findings_render(), swap_capitals)

        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                payload, warnings = self._parse(data, mode)
                self.assertEqual(
                    [(finding["title"], finding["severity"], finding["display_id"]) for finding in payload["vulnerabilities"]],
                    [("Authorization bypass", "medium", "002"), ("AUTHORIZATION BYPASS", "high", "001")],
                )
                self.assertEqual([warning for warning in warnings if "by position" in warning], [])

    def test_a_cvss_row_with_no_title_pairs_by_position_and_is_named_as_untitled(self) -> None:
        data = _edited(asia_findings_render(), lambda document: _set_text(_table_starting(document, "Section").rows[2].cells[1].paragraphs[0], ""))

        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                _payload, warnings = self._parse(data, mode)
                self.assertIn('The CVSS row "(no title)" was paired with the finding row "Second finding" by position.', warnings)


class RefusalMessageTests(unittest.TestCase):
    """A refused Word report says which finding and which part of the report broke which rule, and what
    the importer expected and found, in the words the report itself uses (see GLOSSARY.md)."""

    def assertRefused(self, data: bytes, *expected: str) -> None:
        self.assertTrue(expected, "a refusal test must say what the message has to name")
        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                with self.assertRaises(ReportImportError) as raised:
                    parse_report_docx(data, mode=mode)
                for fragment in expected:
                    self.assertIn(fragment, str(raised.exception))

    def test_a_finding_heading_in_another_style_is_named_with_the_style_it_has(self) -> None:
        def restyle(document):
            _heading(document, "Second finding")._p.find(qn("w:pPr") + "/" + qn("w:pStyle")).set(qn("w:val"), "Heading2")

        self.assertRefused(
            _edited(two_findings_render(), restyle),
            'The finding row "Second finding" has no finding heading left to pair with.',
            "The importer pairs a finding heading with a finding row by title, ignoring capitals and extra spaces, "
            "and then by position, but only when the same number of each is left over.",
            'Finding headings found: "Authorization bypass".',
            "A finding heading is a Report Heading 2 paragraph after the table and before the Appendix.",
            'Headings in other styles after the table: Heading2 ("Second finding").',
        )

    def test_a_finding_heading_with_no_finding_row_is_named(self) -> None:
        def add_heading(document):
            original = _heading(document, "Second finding")
            original._p.addnext(copy.deepcopy(original._p))
            twin = [paragraph for paragraph in document.paragraphs if " ".join(paragraph.text.split()) == "Second finding"][1]
            _set_text(twin, "Extra finding")

        self.assertRefused(
            _edited(two_findings_render(), add_heading),
            'The finding heading "Extra finding" has no finding row left to pair with.',
            'Finding headings found: "Authorization bypass", "Second finding", "Extra finding".',
        )

    def test_headings_with_no_finding_row_are_named_together(self) -> None:
        def drop_rows(document):
            table = _table_starting(document, "Findings")
            for _ in range(2):
                row = table.rows[2]._tr
                row.getparent().remove(row)

        self.assertRefused(
            _edited(lettered_findings_render("BC"), drop_rows),
            'The finding headings "Finding B", "Finding C" have no finding rows left to pair with.',
        )

    def test_leftovers_that_differ_in_number_are_refused_without_pairing_any_by_position(self) -> None:
        def damage(document):
            heading = _heading(document, "Finding C")._p
            heading.getparent().remove(heading)
            _set_text(_heading(document, "Finding B"), "Retitled")

        self.assertRefused(
            _edited(lettered_findings_render("BC"), damage),
            'The finding rows "Finding B", "Finding C" have no finding headings left to pair with.',
            'The finding heading "Retitled" has no finding row left to pair with.',
        )

    def test_a_finding_row_that_lost_a_cell_is_counted_as_skipped(self) -> None:
        def damage(document):
            row = _table_starting(document, "Findings").rows[2]
            row._tr.remove(row.cells[-1]._tc)

        self.assertRefused(
            _edited(two_findings_render(), damage),
            'The finding heading "Second finding" has no finding row left to pair with.',
            "The Findings Summary table also has 1 row with no title or fewer than six cells, which the importer skips.",
        )

    def test_a_report_without_a_findings_summary_table_says_what_it_looked_for(self) -> None:
        def rename_header(document):
            _set_text(_table_starting(document, "Findings").rows[0].cells[0].paragraphs[0], "Issues")

        self.assertRefused(
            _edited(two_findings_render(), rename_header),
            "The Word report has no Findings Summary table.",
            'whose first header cell is "Findings"',
        )

    def test_cvss_rows_that_no_finding_row_matches_are_refused_when_the_leftovers_differ_in_number(self) -> None:
        def damage(document):
            table = _table_starting(document, "Section")
            first = table.rows[1]._tr
            first.getparent().remove(first)
            _set_text(table.rows[1].cells[1].paragraphs[0], "Other finding")

        self.assertRefused(
            _edited(asia_findings_render(), damage),
            'The CVSS table has a row for "Other finding" that no finding row matches.',
            "The importer pairs a CVSS row with a finding row by title, ignoring capitals and extra spaces, "
            "and then by position, but only when the same number of each is left over.",
        )

    def test_several_cvss_rows_that_no_finding_row_matches_are_named_together(self) -> None:
        def add_rows(document):
            table = _table_starting(document, "Section")
            for title in ("Extra one", "Extra two"):
                table._tbl.append(copy.deepcopy(table.rows[2]._tr))
                _set_text(table.rows[-1].cells[1].paragraphs[0], title)

        self.assertRefused(
            _edited(asia_findings_render(), add_rows),
            'The CVSS table has rows for "Extra one", "Extra two" that no finding row matches.',
        )

    def test_a_missing_section_lists_the_sections_expected_and_found(self) -> None:
        self.assertRefused(
            _edited(two_findings_render(), lambda document: _set_text(_heading(document, "Description:"), "Summary:")),
            '"Authorization bypass" is Open (Previously Discovered), so it needs these sections in this order: '
            "Description, Recommended Remediation, Previous Proof of Concept, Proof of Concept, In Conclusion.",
            "Found: Recommended Remediation, Previous Proof of Concept, Proof of Concept, In Conclusion.",
            'A section starts at a paragraph that holds only its label, such as "Description:".',
        )

    def test_a_section_that_appears_twice_is_named(self) -> None:
        self.assertRefused(
            _edited(two_findings_render(), lambda document: _set_text(_heading(document, "Recommended Remediation:"), "Description:")),
            '"Authorization bypass" has the Description section more than once.',
        )

    def test_a_finding_without_its_table_is_named(self) -> None:
        def remove_table(document):
            table = _table_after(document, "Authorization bypass")._tbl
            table.getparent().remove(table)

        self.assertRefused(
            _edited(two_findings_render(), remove_table),
            '"Authorization bypass" has no finding table.',
            "with the labels Severity, ID, Status and Location in its first column.",
        )

    def test_a_finding_with_two_tables_is_named(self) -> None:
        def repeat_table(document):
            table = _table_after(document, "Authorization bypass")._tbl
            table.addnext(copy.deepcopy(table))

        self.assertRefused(
            _edited(two_findings_render(), repeat_table),
            '"Authorization bypass" has more than one finding table.',
            "A finding table is a table under the finding heading, before its first section,",
        )

    def test_a_location_heading_the_importer_does_not_know_is_named(self) -> None:
        def rename_location_heading(document):
            row = next(row for row in _table_after(document, "Authorization bypass").rows if row.cells[0].text.strip() == "Location")
            _set_text(row.cells[1].paragraphs[0], "Somewhere else")

        self.assertRefused(
            _edited(two_findings_render(), rename_location_heading),
            '"Authorization bypass" has an unsupported location heading "Somewhere else".',
            'The first line of its Location row must start with "Production Environment" or end with "Environment:".',
        )

    def test_a_cover_value_given_two_ways_lists_both(self) -> None:
        def give_two_values(document):
            paragraphs = list(document.element.body.iter(qn("w:p")))
            labels = [
                index for index, paragraph in enumerate(paragraphs)
                if "".join(node.text or "" for node in paragraph.iter(qn("w:t"))).strip() == "Delivered To"
            ]
            for index, value in zip(labels, ("Alice", "Bob")):
                nodes = list(paragraphs[index + 1].iter(qn("w:t")))
                nodes[0].text = value
                for node in nodes[1:]:
                    node.text = ""

        with self.assertRaises(ReportImportError) as raised:
            parse_report_docx(_edited(base_render(), give_two_values), mode="editable")

        self.assertIn('The document contains conflicting values for "Delivered To": "Alice", "Bob".', str(raised.exception))

    def test_an_unknown_report_type_in_the_cover_title_lists_the_known_ones(self) -> None:
        def rename_type(document):
            title = next(paragraph for paragraph in document.paragraphs[:20] if "Annual Pentest" in paragraph.text)
            _set_text(title, "JH \u2013 Northstar Banking \u2013 Mystery Pentest 2026")

        with self.assertRaises(ReportImportError) as raised:
            parse_report_docx(_edited(base_render(), rename_type), mode="editable")

        self.assertIn(
            'Unknown report type "Mystery Pentest" in the cover title. It must be one of: Annual Pentest, Retest, Deployment Pentest, New Test.',
            str(raised.exception),
        )

    def test_a_repeated_title_with_one_heading_left_names_the_finding_row(self) -> None:
        def share_a_title(document):
            _set_text(_heading(document, "Second finding"), "Authorization bypass")
            _set_text(_table_starting(document, "Findings").rows[2].cells[0].paragraphs[0], "Authorization bypass")
            twin = [paragraph for paragraph in document.paragraphs if " ".join(paragraph.text.split()) == "Authorization bypass"][1]
            twin._p.find(qn("w:pPr") + "/" + qn("w:pStyle")).set(qn("w:val"), "Heading2")

        self.assertRefused(
            _edited(two_findings_render(), share_a_title),
            'The finding row "Authorization bypass" has no finding heading left to pair with.',
        )

    def test_a_long_list_of_finding_rows_with_no_heading_is_cut_after_eight(self) -> None:
        def drop_headings(document):
            for letter in "BCDEFGHIJ":
                heading = _heading(document, f"Finding {letter}")._p
                heading.getparent().remove(heading)

        self.assertRefused(
            _edited(lettered_findings_render("BCDEFGHIJ"), drop_headings),
            'The finding rows "Finding B", "Finding C", "Finding D", "Finding E", "Finding F", "Finding G", "Finding H", "Finding I" and 1 more '
            "have no finding headings left to pair with.",
        )


class FirstChangeTests(unittest.TestCase):
    """The editable import's last checks say what the app would change by comparing two dumps of one report."""

    PLAIN = {"bold": False, "italic": False, "underline": False}

    def _report(self, fragments: list[dict], contents: list[dict] | None = None, **changes) -> dict:
        finding = {
            "title": "Authorization bypass", "status": "open_new", "severity": "high", "scope": {"target_ids": ["a"]},
            "contents": [{"type": "description", "fragments": fragments}] if contents is None else contents,
        }
        return {"engagement": {"app_name": "Northstar"}, "vulnerabilities": [{**finding, **changes}]}

    def test_each_kind_of_difference_is_named(self) -> None:
        def paragraph(*runs):
            return {"type": "paragraph", "runs": [{"text": text, **self.PLAIN, **flags} for text, flags in runs]}

        def numbered(*items):
            return {"type": "numbered_list", "items": [{"runs": [{"text": text, **self.PLAIN}]} for text in items]}

        head = 'Importing would change the Description section of "Authorization bypass". '
        picture = {"type": "image", "environment": "production", "evidence_id": None, "caption": "", "width_mm": None}
        cases = (
            (
                "a part of another kind",
                self._report([paragraph(("Same", {}))]), self._report([picture]),
                head + "A paragraph would become a picture.",
            ),
            (
                "a different number of parts",
                self._report([paragraph(("Same", {}))]), self._report([paragraph(("Same", {})), picture]),
                head + "It has 1 part and the app would store 2 parts.",
            ),
            (
                "formatting only",
                self._report([paragraph(("Open now", {}))]), self._report([paragraph(("Open", {}), (" now", {"bold": True}))]),
                head + 'The formatting of its paragraph would change: " now" would change from plain to bold.',
            ),
            (
                "how a list is divided",
                self._report([numbered("ab")]), self._report([numbered("a", "b")]),
                head + "The structure of its numbered list would change.",
            ),
            (
                "a finding value",
                self._report([]), self._report([], severity="low"),
                'Importing would change "Authorization bypass": its severity would change from "high" to "low".',
            ),
            (
                "locations of the same number",
                self._report([]), self._report([], scope={"target_ids": ["b"]}),
                'Importing would change the locations of "Authorization bypass".',
            ),
            (
                "locations that would be dropped",
                self._report([], scope={"target_ids": ["a", "b"]}), self._report([]),
                'Importing would change the locations of "Authorization bypass": the report lists 2 scope targets for it and the app would keep 1.',
            ),
            (
                "a section that would go",
                self._report([paragraph(("Same", {}))]), self._report([], contents=[]),
                'Importing would change the sections of "Authorization bypass".',
            ),
            (
                "something outside the findings",
                self._report([]), {**self._report([]), "engagement": {"app_name": "Other"}},
                "Importing would change the report's engagement.",
            ),
            (
                "a changed conclusion sentence",
                self._report([], contents=[{"type": "in_conclusion", "fragments": [paragraph(('The finding "A" is still Open.', {}))]}], status="open_previously_discovered"),
                self._report([], contents=[{"type": "in_conclusion", "fragments": [paragraph(('The finding "B" is still Open.', {}))]}], status="open_previously_discovered"),
                'Importing would change the In Conclusion section of "Authorization bypass". '
                'The report has: "The finding "A" is still Open.". The app would store: "The finding "B" is still Open.". '
                "The app writes the first sentence from the finding's title and status. "
                'The title is "Authorization bypass", from the finding row, and the status is Open (Previously Discovered).',
            ),
            (
                "nothing at all",
                self._report([]), self._report([]),
                "Importing would change the report.",
            ),
        )
        for label, before, after, expected in cases:
            with self.subTest(label):
                self.assertEqual(first_change(before, after), expected)


class ImportRouteTests(unittest.TestCase):
    """A .docx is itself a ZIP, so the sharpest risk here is the new branch stealing the bundle's."""

    def test_supported_word_edits_survive_import_and_a_later_save(self) -> None:
        def edit(document):
            paragraph = _heading(document, "A justified paragraph.")
            paragraph._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = 0
            paragraph.runs[0].bold = False
            paragraph.runs[0].italic = False
            paragraph.runs[0].underline = False
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), "endpoint")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = " Sensitive endpoint"
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)
            table = _table_starting(document, "Header")
            table._tbl.remove(table.rows[1]._tr)

        response = self._import("report.docx", _edited(self._docx(), edit), "editable")
        self.assertEqual(response.status_code, 200, response.text)
        imported = main.workspace.load(response.json()["report_id"])
        finding = imported.vulnerabilities[0]
        paragraph = finding.contents[0].fragments[0]
        self.assertEqual(paragraph.type, "paragraph")
        self.assertEqual("".join(run.text for run in paragraph.runs), "A justified paragraph. Sensitive endpoint")
        self.assertTrue(all(not run.bold and not run.italic and not run.underline for run in paragraph.runs))
        proof = next(content for content in finding.contents if content.type == "proof_of_concept")
        self.assertEqual(next(fragment for fragment in proof.fragments if fragment.type == "table").rows, [])
        before = finding.model_dump(mode="json")
        payload = imported.model_dump(mode="json", by_alias=True)
        payload["scope_text"] = main.scope_text_from_targets(imported.scope_targets)
        payload["engagement"]["app_owner"] = "Owner after import"

        saved = self.client.put(f"/reports/{imported.report_id}", json=payload)

        self.assertEqual(saved.status_code, 200, saved.text)
        after = main.workspace.load(imported.report_id)
        self.assertEqual(after.engagement.app_owner, "Owner after import")
        self.assertEqual(after.vulnerabilities[0].model_dump(mode="json"), before)

    def test_an_extra_table_before_description_refuses_without_creating_a_draft(self) -> None:
        def insert_table(document):
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Account"
            table.cell(0, 1).text = "Sensitive evidence"
            _heading(document, "Description:")._p.addprevious(table._tbl)

        response = self._import("report.docx", _edited(self._docx(), insert_table), "editable")

        self.assertEqual(response.status_code, 422, response.text)
        self.assertIn('"Authorization bypass"', response.json()["detail"])
        self.assertIn("table", response.json()["detail"])
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_a_retest_does_not_reopen_a_closed_finding_row_written_in_capitals(self) -> None:
        def close(document):
            _set_text(_table_starting(document, "Findings").rows[1].cells[5].paragraphs[0], "  CLOSED  ")

        response = self._import("report.docx", _edited(two_findings_render(), close), "retest")

        self.assertEqual(response.status_code, 200, response.text)
        report = main.workspace.load(response.json()["report_id"])
        self.assertEqual([finding.title for finding in report.vulnerabilities], ["Second finding"])
        self.assertEqual(response.json()["summary"]["dropped_resolved"], ["Authorization bypass"])

    def setUp(self) -> None:
        self.root = use_temp_workspace(self, "QA Tester")
        self.client = TestClient(main.app)

    @staticmethod
    def _docx() -> bytes:
        return base_render()

    def _import(self, name: str, data: bytes, mode: str | None = None, import_id: str | None = None):
        form = {"docx_mode": mode} if mode is not None else None
        return self.client.post(
            "/reports/import",
            data=form,
            files={"file": (name, data, "application/octet-stream")},
            headers={"X-VulnReport-Import-Id": import_id} if import_id else None,
        )

    def test_a_generated_report_imports_as_a_retest_draft(self) -> None:
        response = self._import("report.docx", self._docx())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["source"], "docx")
        self.assertEqual(response.json()["summary"]["statuses_rewritten"], ["Authorization bypass"])

        report = main.workspace.load(response.json()["report_id"])
        finding = report.vulnerabilities[0]
        sections = {content.type: [fragment.type for fragment in content.fragments] for content in finding.contents}
        self.assertEqual(sections["previous_proof_of_concept"][-1], "image", "the carried evidence came across")
        self.assertEqual(sections["proof_of_concept"], ["numbered_list", "image"], "the retest gets somewhere to write and to paste")
        self.assertIn(
            "Authorization bypass: Production evidence image required",
            generation_issues(report),
            "last year's screenshot must not stand in for this year's proof",
        )
        stored = main.workspace.find_path(report.report_id).parent / "evidence"
        self.assertEqual(len(list(stored.glob("*.png"))), 1)

    def test_a_retest_inherits_every_status_except_open_new(self) -> None:
        """A retest re-tests the distinction last year's status recorded, so only Open (New) is
        meaningless on one. Before this, anything outside a two-name set was dropped and reported as
        resolved - which would have eaten a finding for carrying a status added after it was written."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = FragmentRecognitionTests._report(folder)
            carried = report.vulnerabilities[0]
            carried.status = "open_resolved_on_non_prod"
            carried.contents.extend([
                Content(type="previous_proof_of_concept", fragments=[
                    ListFragment(frag_id="f_prior", type="numbered_list", items=[ListItem(runs=[Run(text="Prior evidence.")])]),
                ]),
                Content(type="in_conclusion", fragments=[
                    ParagraphFragment(frag_id="f_prior_conclusion", type="paragraph", runs=[Run(text="Open in production.")]),
                ]),
            ])

            payload, _, summary = parse_report_docx(render_report_docx(report, TEMPLATE, folder), mode="retest")

            finding = payload["vulnerabilities"][0]
            self.assertEqual(finding["status"], "open_resolved_on_non_prod", "the retest lost the status it inherited")
            self.assertEqual(summary["dropped_resolved"], [], "an open finding was dropped and called resolved")
            self.assertEqual(summary["statuses_rewritten"], [], "nothing changed, so nothing may be reported as changed")

    def test_every_segment_the_model_allows_parses_back_out_of_a_title(self) -> None:
        """The title parser holds the segment set a second time, as a literal tuple, and the two
        copies have no other guard. Driving this from the Literal is what catches a segment added to
        one and not the other -- a miss loses segment, app name and report type together, and raises
        nothing, because the unknown-report-type error sits inside the branch a failed match never
        enters."""
        resources = Path(__file__).resolve().parent.parent / "resources"
        for segment in get_args(Segment):
            with self.subTest(segment=segment), tempfile.TemporaryDirectory() as temporary_directory:
                folder = Path(temporary_directory)
                report = FragmentRecognitionTests._report(folder)
                report.engagement.segment = segment
                if segment == "Asia":
                    report.vulnerabilities[0].cvss_score = "9.8"
                    report.vulnerabilities[0].cvss_vector = "CVSS:3.1/AV:N/AC:L/PR:N/UI:N/S:U/C:H/I:H/A:H"

                document = render_report_docx(report, main_template_path(report, resources), folder)
                payload, _, _ = parse_report_docx(document, mode="editable")

                engagement = payload["engagement"]
                self.assertEqual(engagement["segment"], segment)
                self.assertEqual(engagement["app_name"], "Northstar Banking")
                self.assertEqual(engagement["report_type"], "annual_pentest")

    def test_gdt_titles_round_trip_every_report_type_label(self) -> None:
        """The title parser receives report type only as a display label, so each label must map
        back to its stored value once the GDT segment has passed the cover-title gate."""
        for report_type in ("annual_pentest", "retest", "deployment_pentest", "new_test"):
            with self.subTest(report_type=report_type), tempfile.TemporaryDirectory() as temporary_directory:
                folder = Path(temporary_directory)
                report = FragmentRecognitionTests._report(folder)
                report.engagement.segment = "GDT"
                report.engagement.report_type = report_type

                payload, _, _ = parse_report_docx(render_report_docx(report, TEMPLATE, folder), mode="editable")

                engagement = payload["engagement"]
                self.assertEqual(engagement["segment"], "GDT")
                self.assertEqual(engagement["report_type"], report_type)

    def test_a_retest_drops_closed_findings_and_keeps_the_open_ones(self) -> None:
        """Closed borrows exactly one rule from Resolved -- the retest drop -- and a dropped finding
        takes its evidence and typed work with it, so the summary is its only trace."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = FragmentRecognitionTests._report(folder)
            kept = report.vulnerabilities[0]
            kept.status = "open_resolved_on_non_prod"
            kept.contents.extend([
                Content(type="previous_proof_of_concept", fragments=[
                    ListFragment(frag_id="f_prior", type="numbered_list", items=[ListItem(runs=[Run(text="Prior evidence.")])]),
                ]),
                Content(type="in_conclusion", fragments=[
                    ParagraphFragment(frag_id="f_prior_conclusion", type="paragraph", runs=[Run(text="Open in production.")]),
                ]),
            ])
            shut = copy.deepcopy(kept)
            shut.uid, shut.display_id, shut.title, shut.status = "v_shut", "002", "Closed last year", "closed"
            report.vulnerabilities = [kept, shut]

            document = render_report_docx(report, TEMPLATE, folder)
            self.assertIn("Closed", "\n".join(item.text for table in Document(BytesIO(document)).tables for row in table.rows for item in row.cells))

            payload, _, summary = parse_report_docx(document, mode="retest")

            self.assertEqual([finding["title"] for finding in payload["vulnerabilities"]], ["Authorization bypass"])
            self.assertEqual(summary["dropped_resolved"], ["Closed last year"])

    def test_an_editable_import_keeps_a_closed_finding_as_closed(self) -> None:
        """Editable mode is the recovery path for anything a retest drops, so the status has to
        survive it rather than downgrading to previously discovered."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = FragmentRecognitionTests._report(folder)
            finding = report.vulnerabilities[0]
            finding.status = "closed"
            finding.contents.extend([
                Content(type="previous_proof_of_concept", fragments=[
                    ListFragment(frag_id="f_prior", type="numbered_list", items=[ListItem(runs=[Run(text="Prior evidence.")])]),
                ]),
                Content(type="in_conclusion", fragments=[
                    ParagraphFragment(frag_id="f_prior_conclusion", type="paragraph", runs=[Run(text="Shut without a fix.")]),
                ]),
            ])

            payload, _, summary = parse_report_docx(render_report_docx(report, TEMPLATE, folder), mode="editable")

            self.assertEqual(payload["vulnerabilities"][0]["status"], "closed")
            # A recognised label warns about nothing; an unrecognised one downgrades and says so.
            self.assertEqual([warning for warning in summary["warnings"] if "imported as" in warning], [])

    def test_prompt_classifies_a_docx_without_creating_a_report(self) -> None:
        response = self._import("misleading.zip", self._docx(), "prompt")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"source": "docx", "mode_required": True})
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_route_imports_editable_then_first_put_preserves_visible_semantics(self) -> None:
        response = self._import("report.docx", self._docx(), "editable")
        self.assertEqual(response.status_code, 200, response.text)
        result = response.json()
        self.assertEqual(result["mode"], "editable")
        self.assertTrue(any("rendered Word copies" in warning for warning in result["summary"]["warnings"]))

        imported = main.workspace.load(result["report_id"])
        self.assertEqual(imported.vulnerabilities[0].status, "open_new")
        target_ids = [target.target_id for target in imported.scope_targets]
        before = main._editable_user_projection(imported)
        previous_revision = imported.saved_at
        payload = imported.model_dump(mode="json", by_alias=True)
        payload["scope_text"] = main.scope_text_from_targets(imported.scope_targets)
        saved = self.client.put(f"/reports/{imported.report_id}", json=payload)

        self.assertEqual(saved.status_code, 200, saved.text)
        after = main.workspace.load(imported.report_id)
        self.assertEqual([target.target_id for target in after.scope_targets], target_ids)
        self.assertEqual(main._editable_user_projection(after), before)
        self.assertEqual(after.app_id, "Northstar_Banking")
        self.assertGreater(after.saved_at, previous_revision)

    def test_repeated_import_request_id_returns_the_first_result_without_a_second_report(self) -> None:
        import_id = f"test-{uuid.uuid4()}"
        first = self._import("report.docx", self._docx(), "editable", import_id)
        second = self._import("report.docx", self._docx(), "editable", import_id)

        self.assertEqual(first.status_code, 200, first.text)
        self.assertEqual(second.status_code, 200, second.text)
        self.assertEqual(second.json(), first.json())
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_one_import_request_id_keeps_prompt_classification_separate_from_the_final_import(self) -> None:
        import_id = f"test-{uuid.uuid4()}"
        prompt = self._import("report.docx", self._docx(), "prompt", import_id)
        editable = self._import("report.docx", self._docx(), "editable", import_id)

        self.assertEqual(prompt.status_code, 200, prompt.text)
        self.assertEqual(prompt.json(), {"source": "docx", "mode_required": True})
        self.assertEqual(editable.status_code, 200, editable.text)
        self.assertEqual(editable.json()["mode"], "editable")
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_failed_import_request_id_can_be_retried_with_valid_bytes(self) -> None:
        import_id = f"test-{uuid.uuid4()}"
        refused = self._import("report.docx", b"not a report", "editable", import_id)
        retried = self._import("report.docx", self._docx(), "editable", import_id)

        self.assertEqual(refused.status_code, 422)
        self.assertEqual(retried.status_code, 200, retried.text)
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_overlapping_import_retries_share_one_in_flight_request(self) -> None:
        import_id = f"test-{uuid.uuid4()}"
        entered = Event()
        pending_claimed = Event()
        release = Event()
        calls = []
        original_import = main._import_report_once
        original_claim = main._claim_import

        async def held_import(file, mode):
            calls.append(mode)
            entered.set()
            await main.run_in_threadpool(release.wait)
            return await original_import(file, mode)

        def observed_claim(request_key):
            claim = original_claim(request_key)
            if claim[0] == "pending":
                pending_claimed.set()
            return claim

        try:
            with patch.object(main, "_import_report_once", side_effect=held_import), patch.object(main, "_claim_import", side_effect=observed_claim):
                with ThreadPoolExecutor(max_workers=2) as pool:
                    first = pool.submit(self._import, "report.docx", self._docx(), "editable", import_id)
                    self.assertTrue(entered.wait(2), "the first request never reached the importer")
                    second = pool.submit(self._import, "report.docx", self._docx(), "editable", import_id)
                    self.assertTrue(pending_claimed.wait(2), "the retry did not join the in-flight request")
                    release.set()
                    responses = [first.result(), second.result()]
        finally:
            release.set()

        self.assertEqual(calls, ["editable"])
        self.assertEqual([response.status_code for response in responses], [200, 200])
        self.assertEqual(responses[0].json(), responses[1].json())
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_overlapping_failed_imports_both_refuse_and_leave_the_request_id_retryable(self) -> None:
        import_id = f"test-{uuid.uuid4()}"
        entered = Event()
        pending_claimed = Event()
        release = Event()
        calls = []
        original_import = main._import_report_once
        original_claim = main._claim_import

        async def held_import(file, mode):
            calls.append(mode)
            entered.set()
            await main.run_in_threadpool(release.wait)
            return await original_import(file, mode)

        def observed_claim(request_key):
            claim = original_claim(request_key)
            if claim[0] == "pending":
                pending_claimed.set()
            return claim

        try:
            with patch.object(main, "_import_report_once", side_effect=held_import), patch.object(main, "_claim_import", side_effect=observed_claim):
                with ThreadPoolExecutor(max_workers=2) as pool:
                    first = pool.submit(self._import, "report.docx", b"not a report", "editable", import_id)
                    self.assertTrue(entered.wait(2), "the first request never reached the importer")
                    second = pool.submit(self._import, "report.docx", b"not a report", "editable", import_id)
                    self.assertTrue(pending_claimed.wait(2), "the retry did not join the in-flight request")
                    release.set()
                    responses = [first.result(), second.result()]
        finally:
            release.set()

        self.assertEqual(calls, ["editable"])
        self.assertEqual([response.status_code for response in responses], [422, 422])
        retried = self._import("report.docx", self._docx(), "editable", import_id)
        self.assertEqual(retried.status_code, 200, retried.text)
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_overlapping_final_modes_share_one_in_flight_import(self) -> None:
        import_id = f"test-{uuid.uuid4()}"
        entered = Event()
        pending_claimed = Event()
        release = Event()
        calls = []
        original_import = main._import_report_once
        original_claim = main._claim_import

        async def held_import(file, mode):
            calls.append(mode)
            entered.set()
            await main.run_in_threadpool(release.wait)
            return await original_import(file, mode)

        def observed_claim(request_key):
            claim = original_claim(request_key)
            if claim[0] == "pending":
                pending_claimed.set()
            return claim

        try:
            with patch.object(main, "_import_report_once", side_effect=held_import), patch.object(main, "_claim_import", side_effect=observed_claim):
                with ThreadPoolExecutor(max_workers=2) as pool:
                    editable = pool.submit(self._import, "report.docx", self._docx(), "editable", import_id)
                    self.assertTrue(entered.wait(2), "the editable request never reached the importer")
                    retest = pool.submit(self._import, "report.docx", self._docx(), "retest", import_id)
                    joined = pending_claimed.wait(1)
                    release.set()
                    responses = [editable.result(), retest.result()]
        finally:
            release.set()

        self.assertTrue(joined, "the second final mode did not join the in-flight import")
        self.assertEqual(calls, ["editable"])
        self.assertEqual([response.status_code for response in responses], [200, 200])
        self.assertEqual(responses[0].json(), responses[1].json())
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_legacy_omitted_mode_and_explicit_retest_share_one_in_flight_import(self) -> None:
        import_id = f"test-{uuid.uuid4()}"
        entered = Event()
        pending_claimed = Event()
        release = Event()
        calls = []
        original_import = main._import_report_once
        original_claim = main._claim_import

        async def held_import(file, mode):
            calls.append(mode)
            entered.set()
            await main.run_in_threadpool(release.wait)
            return await original_import(file, mode)

        def observed_claim(request_key):
            claim = original_claim(request_key)
            if claim[0] == "pending":
                pending_claimed.set()
            return claim

        try:
            with patch.object(main, "_import_report_once", side_effect=held_import), patch.object(main, "_claim_import", side_effect=observed_claim):
                with ThreadPoolExecutor(max_workers=2) as pool:
                    legacy = pool.submit(self._import, "report.docx", self._docx(), None, import_id)
                    self.assertTrue(entered.wait(2), "the legacy request never reached the importer")
                    explicit = pool.submit(self._import, "report.docx", self._docx(), "retest", import_id)
                    joined = pending_claimed.wait(1)
                    release.set()
                    responses = [legacy.result(), explicit.result()]
        finally:
            release.set()

        self.assertTrue(joined, "explicit retest did not join the legacy in-flight retest")
        self.assertEqual(calls, [None])
        self.assertEqual([response.status_code for response in responses], [200, 200])
        self.assertEqual(responses[0].json(), responses[1].json())
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_retry_reimports_when_the_report_in_its_receipt_was_deleted(self) -> None:
        import_id = f"test-{uuid.uuid4()}"
        first = self._import("report.docx", self._docx(), "editable", import_id)
        self.assertEqual(first.status_code, 200, first.text)
        main.workspace.delete(first.json()["report_id"])

        retried = self._import("report.docx", self._docx(), "editable", import_id)

        self.assertEqual(retried.status_code, 200, retried.text)
        self.assertNotEqual(retried.json()["report_id"], first.json()["report_id"])
        main.workspace.load(retried.json()["report_id"])
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_malformed_persistent_import_receipts_do_not_block_an_import(self) -> None:
        main._import_receipts_path().write_text("{not json", encoding="utf-8")
        with main._import_receipts_lock:
            main._import_receipts.clear()
            main._imports_in_flight.clear()

        response = self._import("report.docx", self._docx(), "editable", f"test-{uuid.uuid4()}")

        self.assertEqual(response.status_code, 200, response.text)
        main.workspace.load(response.json()["report_id"])
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_wrong_persistent_receipt_shape_does_not_block_an_import(self) -> None:
        main._import_receipts_path().write_text('{"receipts": 12}', encoding="utf-8")
        with main._import_receipts_lock:
            main._import_receipts.clear()
            main._imports_in_flight.clear()

        response = self._import("report.docx", self._docx(), "editable", f"test-{uuid.uuid4()}")

        self.assertEqual(response.status_code, 200, response.text)
        main.workspace.load(response.json()["report_id"])
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_persisted_receipt_for_a_deleted_report_is_discarded_after_restart(self) -> None:
        import_id = f"test-{uuid.uuid4()}"
        first = self._import("report.docx", self._docx(), "editable", import_id)
        self.assertEqual(first.status_code, 200, first.text)
        main.workspace.delete(first.json()["report_id"])
        with main._import_receipts_lock:
            main._import_receipts.clear()
            main._imports_in_flight.clear()

        retried = self._import("report.docx", self._docx(), "editable", import_id)

        self.assertEqual(retried.status_code, 200, retried.text)
        self.assertNotEqual(retried.json()["report_id"], first.json()["report_id"])
        main.workspace.load(retried.json()["report_id"])
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_newer_imports_do_not_evict_the_receipt_for_a_live_report(self) -> None:
        import_ids = [f"test-{uuid.uuid4()}" for _ in range(3)]
        with patch.object(main, "IMPORT_RECEIPT_LIMIT", 2):
            imported = [self._import("report.docx", self._docx(), "editable", import_id) for import_id in import_ids]
            self.assertEqual([response.status_code for response in imported], [200, 200, 200])
            with main._import_receipts_lock:
                main._import_receipts.clear()
                main._imports_in_flight.clear()
            retried = self._import("report.docx", self._docx(), "editable", import_ids[0])

        self.assertEqual(retried.status_code, 200, retried.text)
        self.assertEqual(retried.json(), imported[0].json())
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 3)

    def test_explicit_docx_mode_is_rejected_for_a_bundle(self) -> None:
        created = self.client.get("/new", follow_redirects=False)
        report_id = created.headers["location"].split("/")[2]
        exported = self.client.get(f"/reports/{report_id}/export")

        response = self._import("report.docx", exported.content, "editable")

        self.assertEqual(response.status_code, 422)
        self.assertIn("not a report DOCX", response.json()["detail"])

    def test_prompted_bundle_imports_without_a_mode_round_trip(self) -> None:
        created = self.client.get("/new", follow_redirects=False)
        report_id = created.headers["location"].split("/")[2]
        exported = self.client.get(f"/reports/{report_id}/export")

        response = self._import("misleading.docx", exported.content, "prompt")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["source"], "bundle")
        self.assertNotIn("mode_required", response.json())

    def test_unknown_docx_mode_is_rejected_without_writing(self) -> None:
        before = set(self.root.glob("apps/**/draft.json"))
        response = self._import("report.docx", self._docx(), "replace")
        self.assertEqual(response.status_code, 422)
        self.assertIn("Unknown DOCX import mode", response.json()["detail"])
        self.assertEqual(set(self.root.glob("apps/**/draft.json")), before)

    def test_an_unrecognised_status_is_assumed_rather_than_refusing_the_document(self) -> None:
        """A report written by an older version, or edited by hand, must still import. Only the
        Status in the Findings Summary table is changed here, so the finding table still reads the
        original label. The importer assumes a status and names it in a warning."""
        document = Document(BytesIO(self._docx()))
        summary = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Findings")
        title = summary.rows[1].cells[0].text.strip()
        summary.rows[1].cells[5].text = "Pending"
        output = BytesIO()
        document.save(output)

        response = self._import("report.docx", output.getvalue(), "editable")

        self.assertEqual(response.status_code, 200)
        draft = json.loads(next(self.root.glob("apps/**/draft.json")).read_text())
        assumed = next(finding for finding in draft["vulnerabilities"] if finding["title"] == title)
        self.assertEqual(assumed["status"], "open_previously_discovered")
        self.assertEqual(
            [content["type"] for content in assumed["contents"]],
            ["description", "recommended_remediation", "previous_proof_of_concept", "proof_of_concept", "in_conclusion"],
            "provisioning rebuilds to the status's shape, so what the importer hands over must match it",
        )
        warnings = response.json()["summary"]["warnings"]
        self.assertTrue(
            any(title in warning and "Pending" in warning for warning in warnings),
            f"the assumption was made silently: {warnings}",
        )

    def test_truncated_location_row_is_rejected_without_writing(self) -> None:
        document = Document(BytesIO(self._docx()))
        location_row = next(
            row
            for table in document.tables
            for row in table.rows
            if row.cells and row.cells[0].text.strip() == "Location"
        )
        location_row._tr.remove(location_row.cells[-1]._tc)
        output = BytesIO()
        document.save(output)

        response = self._import("report.docx", output.getvalue(), "editable")

        self.assertEqual(response.status_code, 422)
        self.assertIn('"Authorization bypass" has an invalid Location row', response.json()["detail"])
        self.assertIn("the row has 1 of the 2 cells it needs", response.json()["detail"])
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_truncated_component_row_is_rejected_without_writing(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = FragmentRecognitionTests._report(folder)
            report.engagement.tested_channels = ["thick_client"]
            report.scope_targets = [ScopeTarget(
                target_id="t_component",
                environment="production",
                channel="thick_client",
                value="Acme.exe",
                description="Desktop client",
            )]
            report.vulnerabilities[0].scope.target_ids = ["t_component"]
            document = Document(BytesIO(render_report_docx(
                report,
                main_template_path(report, Path(__file__).resolve().parent.parent / "resources"),
                folder,
            )))
        component = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Component")
        component.rows[1]._tr.remove(component.rows[1].cells[-1]._tc)
        output = BytesIO()
        document.save(output)

        response = self._import("report.docx", output.getvalue(), "editable")

        self.assertEqual(response.status_code, 422)
        self.assertIn("invalid Component row", response.json()["detail"])
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_component_table_with_both_channel_markers_is_rejected_without_writing(self) -> None:
        """Mutual exclusion between Mobile and Thick Client is a Setup-page UI rule only; nothing on
        the load path enforces it. A hand-edited document naming both markers must still be refused
        at import, since the parser cannot otherwise decide which channel the one Component table
        belongs to."""
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = FragmentRecognitionTests._report(folder)
            report.engagement.tested_channels = ["thick_client"]
            report.scope_targets = [ScopeTarget(
                target_id="t_component",
                environment="production",
                channel="thick_client",
                value="Acme.exe",
                description="Desktop client",
            )]
            report.vulnerabilities[0].scope.target_ids = ["t_component"]
            document = Document(BytesIO(render_report_docx(
                report,
                main_template_path(report, Path(__file__).resolve().parent.parent / "resources"),
                folder,
            )))
        marker_paragraph = next(p for p in document.paragraphs if "Thick Client Application Binary" in p.text)
        duplicate = copy.deepcopy(marker_paragraph._p)
        for node in duplicate.findall(".//" + qn("w:t")):
            if node.text == "Thick Client ":
                node.text = "Mobile "
        marker_paragraph._p.addnext(duplicate)
        output = BytesIO()
        document.save(output)

        response = self._import("report.docx", output.getvalue(), "editable")

        self.assertEqual(response.status_code, 422)
        self.assertIn("does not identify exactly one supported app type", response.json()["detail"])
        self.assertIn('mentions 2', response.json()["detail"])
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_truncated_asia_section_row_is_rejected_without_writing(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            folder = Path(temporary_directory)
            report = FragmentRecognitionTests._report(folder)
            report.engagement.segment = "Asia"
            finding = report.vulnerabilities[0]
            finding.cvss_score = "8.1"
            finding.cvss_vector = "CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:U/C:H/I:H/A:N"
            document = Document(BytesIO(render_report_docx(
                report,
                main_template_path(report, Path(__file__).resolve().parent.parent / "resources"),
                folder,
            )))
        section = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Section")
        section.rows[1]._tr.remove(section.rows[1].cells[-1]._tc)
        output = BytesIO()
        document.save(output)

        response = self._import("report.docx", output.getvalue(), "editable")

        self.assertEqual(response.status_code, 422)
        self.assertIn("invalid Section row", response.json()["detail"])
        self.assertIn('the row for "Authorization bypass" has 4 cells and needs 5', response.json()["detail"])
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_empty_limitations_row_is_rejected_without_writing(self) -> None:
        document = Document(BytesIO(self._docx()))
        limitations = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Limitations")
        limitations.rows[1]._tr.remove(limitations.rows[1].cells[0]._tc)
        output = BytesIO()
        document.save(output)

        response = self._import("report.docx", output.getvalue(), "editable")

        self.assertEqual(response.status_code, 422)
        self.assertIn("invalid Limitations row", response.json()["detail"])
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_a_word_report_that_fails_names_a_word_report(self) -> None:
        document = Document(BytesIO(self._docx()))
        limitations = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Limitations")
        limitations.rows[1]._tr.remove(limitations.rows[1].cells[0]._tc)
        output = BytesIO()
        document.save(output)

        response = self._import("report.zip", output.getvalue(), "editable")

        self.assertEqual(response.status_code, 422)
        self.assertRegex(response.json()["detail"], r"^This Word report can't be imported: .*invalid Limitations row")

    def test_a_refusal_through_the_route_names_what_did_not_pair_up(self) -> None:
        def restyle(document):
            _heading(document, "Second finding")._p.find(qn("w:pPr") + "/" + qn("w:pStyle")).set(qn("w:val"), "Heading2")

        response = self._import("report.docx", _edited(two_findings_render(), restyle), "editable")

        self.assertEqual(response.status_code, 422)
        self.assertRegex(
            response.json()["detail"],
            r'^This Word report can\'t be imported: The finding row "Second finding" has no finding heading left to pair with\.',
        )
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_a_heading_in_other_words_imports_through_the_route_with_the_position_warning(self) -> None:
        def retitle(document):
            _set_text(_heading(document, "Second finding"), "Totally different words")

        data = _edited(two_findings_render(), retitle)
        warning = (
            'The finding heading "Totally different words" was paired with the finding row "Second finding" by position. '
            "The finding uses the row's title."
        )
        for mode in ("retest", "editable"):
            with self.subTest(mode=mode):
                response = self._import("report.docx", data, mode)

                self.assertEqual(response.status_code, 200, response.text)
                self.assertIn(warning, response.json()["summary"]["warnings"])
                report = main.workspace.load(response.json()["report_id"])
                self.assertEqual([finding.title for finding in report.vulnerabilities], ["Authorization bypass", "Second finding"])

    def test_an_export_that_fails_names_an_export(self) -> None:
        response = self._import("report.docx", b"not valid JSON", "prompt")

        self.assertEqual(response.status_code, 422)
        self.assertRegex(response.json()["detail"], r"^Select a valid VulnReport export: ")

    def test_empty_scope_row_is_rejected_without_writing(self) -> None:
        document = Document(BytesIO(self._docx()))
        scope = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "URL(s) in Scope")
        value_row = next(row for row in scope.rows[1:] if row.cells[0].text.strip() == "https://prod.example.test")
        value_row._tr.remove(value_row.cells[0]._tc)
        output = BytesIO()
        document.save(output)

        response = self._import("report.docx", output.getvalue(), "editable")

        self.assertEqual(response.status_code, 422)
        self.assertIn("invalid URL(s) in Scope row", response.json()["detail"])
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_empty_table_header_row_does_not_break_lookup(self) -> None:
        document = Document(BytesIO(self._docx()))
        limitations = next(table for table in document.tables if table.rows[0].cells[0].text.strip() == "Limitations")
        header_row = limitations.rows[0]
        for cell in tuple(header_row.cells):
            header_row._tr.remove(cell._tc)
        output = BytesIO()
        document.save(output)

        response = self._import("report.docx", output.getvalue(), "editable")

        # A table whose header row is gone is simply not found; everything else still imports.
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["summary"]["retained"], 1)
        self.assertEqual(len(list(self.root.glob("apps/**/draft.json"))), 1)

    def test_empty_finding_detail_row_is_skipped_without_crashing(self) -> None:
        document = Document(BytesIO(self._docx()))
        detail = next(
            table
            for table in document.tables
            if any(row.cells and row.cells[0].text.strip() == "Location" for row in table.rows)
        )
        empty_row = detail.add_row()
        for cell in tuple(empty_row.cells):
            empty_row._tr.remove(cell._tc)
        output = BytesIO()
        document.save(output)

        response = self._import("report.docx", output.getvalue(), "editable")

        self.assertEqual(response.status_code, 200, response.text)

    def test_docx_parser_limits_surface_as_payload_too_large(self) -> None:
        with patch.object(main, "MAX_BUNDLE_FILES", 1):
            response = self._import("report.docx", self._docx(), "editable")
        self.assertEqual(response.status_code, 413)

    def test_editable_preflight_rejects_scope_reconciliation_loss(self) -> None:
        payload, _evidence, _summary = parse_report_docx(self._docx(), mode="editable")
        duplicate = copy.deepcopy(payload["scope_targets"][0])
        duplicate["target_id"] = "tgt_duplicate"
        payload["scope_targets"].append(duplicate)
        payload["vulnerabilities"][0]["scope"]["target_ids"].append("tgt_duplicate")

        with self.assertRaises(ValueError) as raised:
            main.finalize_editable_import(payload)

        self.assertIn('Importing would change the locations of "Authorization bypass": the report lists 2 scope targets for it and the app would keep 1.', str(raised.exception))

    def test_a_finding_whose_scope_would_be_dropped_is_named(self) -> None:
        payload, _evidence, _summary = parse_report_docx(self._docx(), mode="editable")
        payload["engagement"]["tested_environments"] = ["non_production"]

        with self.assertRaises(ValueError) as raised:
            main.finalize_editable_import(payload)

        self.assertIn(
            'Importing would leave "Authorization bypass" with no location. '
            "The report does not list the environment or app type of the scope targets it uses as tested.",
            str(raised.exception),
        )

    def test_a_conclusion_that_continues_after_the_apps_sentence_imports_as_an_editable_copy(self) -> None:
        """Only how the text is split into runs differs once the server rewrites the sentence, so the
        import must not take that for a change to what was imported."""
        runs = [Run(text='The finding "Authorization bypass" is still '), Run(text="Open", bold=True), Run(text=". Retested on 5 October.")]
        payload, _evidence, _summary = parse_report_docx(conclusion_render(runs), mode="editable")

        imported, _normalizations = main.finalize_editable_import(payload)

        conclusion = next(content for content in imported["vulnerabilities"][0]["contents"] if content["type"] == "in_conclusion")
        text = "".join(run["text"] for fragment in conclusion["fragments"] for run in fragment["runs"])
        self.assertEqual(text, 'The finding "Authorization bypass" is still Open. Retested on 5 October.')

    def test_a_title_with_extra_spaces_still_imports_as_an_editable_copy_when_the_conclusion_names_it(self) -> None:
        """The conclusion sentence carries the title as typed, so a title that is stored with fewer spaces
        makes provisioning rewrite the sentence and the guard take that for an edit."""
        for label, title in (("a double space", "Cross  site scripting"), ("a trailing space", "Cross site scripting ")):
            with self.subTest(title=label):
                runs = [Run(text=f'The finding "{title}" is still '), Run(text="Open", bold=True), Run(text=". Retested on 5 October.")]
                payload, _evidence, _summary = parse_report_docx(conclusion_render(runs, title=title), mode="editable")

                imported, _normalizations = main.finalize_editable_import(payload)

                self.assertEqual(imported["vulnerabilities"][0]["title"], title)

    def test_a_conclusion_sentence_whose_status_word_lost_its_bold_is_refused_with_the_word_named(self) -> None:
        runs = [Run(text='The finding "Authorization bypass" is still Open. Retested on 5 October.')]
        payload, _evidence, _summary = parse_report_docx(conclusion_render(runs), mode="editable")

        with self.assertRaises(ValueError) as raised:
            main.finalize_editable_import(payload)

        self.assertIn(
            'Importing would change the In Conclusion section of "Authorization bypass". '
            'The formatting of its paragraph would change: "Open" would change from plain to bold.',
            str(raised.exception),
        )

    def test_a_conclusion_sentence_that_disagrees_with_the_status_is_refused_by_name(self) -> None:
        runs = [Run(text='The finding "Authorization bypass" is still '), Run(text="Open", bold=True), Run(text=". Retested on 5 October.")]
        payload, _evidence, _summary = parse_report_docx(conclusion_render(runs, status="closed"), mode="editable")

        with self.assertRaises(ValueError) as raised:
            main.finalize_editable_import(payload)

        message = str(raised.exception)
        for expected in (
            'Importing would change the In Conclusion section of "Authorization bypass".',
            'The report has: "The finding "Authorization bypass" is still Open. Retested on 5 October."',
            'The app would store: "The finding "Authorization bypass" is now CLOSED. Retested on 5 October."',
            "The app writes the first sentence from the finding's title and status. "
            'The title is "Authorization bypass", from the finding row, and the status is Closed.',
        ):
            self.assertIn(expected, message)

    def test_a_finding_row_title_that_differs_from_the_conclusion_sentence_is_refused_with_the_title_named(self) -> None:
        runs = [Run(text='The finding "Authorization bypass" is still '), Run(text="Open", bold=True), Run(text=". Retested on 5 October.")]

        def shout(document):
            _set_text(_table_starting(document, "Findings").rows[1].cells[0].paragraphs[0], "AUTHORIZATION BYPASS")

        payload, _evidence, _summary = parse_report_docx(_edited(conclusion_render(runs), shout), mode="editable")
        self.assertEqual(payload["vulnerabilities"][0]["title"], "AUTHORIZATION BYPASS")

        with self.assertRaises(ValueError) as raised:
            main.finalize_editable_import(payload)

        message = str(raised.exception)
        for expected in (
            'Importing would change the In Conclusion section of "AUTHORIZATION BYPASS".',
            'The app would store: "The finding "AUTHORIZATION BYPASS" is still Open. Retested on 5 October."',
            "The app writes the first sentence from the finding's title and status. "
            'The title is "AUTHORIZATION BYPASS", from the finding row, and the status is Open (Previously Discovered).',
        ):
            self.assertIn(expected, message)

    def test_a_picture_in_the_description_is_refused_with_the_finding_and_section_named(self) -> None:
        def paste_picture(document):
            picture = next(paragraph for paragraph in document.paragraphs if paragraph._p.findall(".//" + qn("a:blip")))
            _heading(document, "Description:")._p.addnext(copy.deepcopy(picture._p))

        payload, _evidence, _summary = parse_report_docx(_edited(base_render(), paste_picture), mode="editable")

        with self.assertRaises(ValueError) as raised:
            main.finalize_editable_import(payload)

        self.assertIn(
            'Importing would change the Description section of "Authorization bypass". Its picture would change: environment from empty to "production".',
            str(raised.exception),
        )

    def test_editable_preflight_names_invalid_fields_before_scope_reconciliation_loss(self) -> None:
        """Acceptance checks fields before the import compares locations, so failing both names the fields."""
        payload, _evidence, _summary = parse_report_docx(self._docx(), mode="editable")
        duplicate = copy.deepcopy(payload["scope_targets"][0])
        duplicate["target_id"] = "tgt_duplicate"
        payload["scope_targets"].append(duplicate)
        payload["vulnerabilities"][0]["scope"]["target_ids"].append("tgt_duplicate")
        payload["vulnerabilities"][0]["cvss_score"] = "9,8"

        with self.assertRaisesRegex(ValueError, "^Imported report contains fields that cannot be saved: .*CVSS Score"):
            main.finalize_editable_import(payload)

    def test_an_editable_import_keeps_a_setup_value_that_breaks_a_rule_as_a_setup_issue(self) -> None:
        """Only a changed Setup value is refused, and nothing has changed yet on import."""
        payload, _evidence, _summary = parse_report_docx(self._docx(), mode="editable")
        payload["engagement"]["app_name"] = "Bad/App"

        imported, _normalizations = main.finalize_editable_import(payload)

        self.assertEqual(imported["engagement"]["app_name"], "Bad/App")
        self.assertIn(character_issue("app_name", "Bad/App"), setup_issues(Report.model_validate(imported)))

    def test_editable_import_rolls_back_when_an_evidence_write_fails(self) -> None:
        with patch("app.workspace.atomic_write_bytes", side_effect=OSError("disk full")):
            response = self._import("report.docx", self._docx(), "editable")

        self.assertEqual(response.status_code, 422)
        self.assertEqual(list(self.root.glob("apps/**/draft.json")), [])

    def test_a_bundle_still_takes_the_bundle_branch(self) -> None:
        created = self.client.get("/new", follow_redirects=False)
        report_id = created.headers["location"].split("/")[2]
        exported = self.client.get(f"/reports/{report_id}/export")
        response = self._import("report.zip", exported.content)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["source"], "bundle")

    def test_a_truncated_document_is_refused_rather_than_crashing(self) -> None:
        response = self._import("report.docx", self._docx()[: 4_000])
        self.assertEqual(response.status_code, 422)

    def test_the_wrong_kind_of_file_is_named_rather_than_decoded(self) -> None:
        """Every non-ZIP falls through to the JSON branch, so without this a tester who picks a PDF
        is told "Expecting value: line 1 column 1"."""
        cases = {
            b"%PDF-1.4\ntrailer<</Root 1 0 R>>\n%%EOF\n": "this is a PDF",
            b"\x89PNG\r\n\x1a\n" + b"\x00" * 20: "not a VulnReport ZIP",
            b"": "the file is empty",
        }
        for data, expected in cases.items():
            with self.subTest(expected=expected):
                response = self._import("report.docx", data)
                self.assertEqual(response.status_code, 422)
                self.assertIn(expected, response.json()["error"]["message"])


if __name__ == "__main__":
    unittest.main()
